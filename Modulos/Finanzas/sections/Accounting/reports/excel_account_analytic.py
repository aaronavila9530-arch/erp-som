from datetime import datetime
from tkinter import filedialog
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _period(row: Dict[str, Any]) -> str:
    period = str(row.get("period") or "").strip()
    if len(period) == 7 and period[4] == "-":
        return period
    value = row.get("entry_date") or row.get("created_at")
    if isinstance(value, datetime):
        return f"{value.year}-{value.month:02d}"
    try:
        dt = datetime.fromisoformat(str(value).replace(" ", "T"))
        return f"{dt.year}-{dt.month:02d}"
    except Exception:
        return ""


def _date_sort_value(row: Dict[str, Any]) -> str:
    return str(row.get("entry_date") or row.get("created_at") or "")


def _num(value: Any) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _account_type(row: Dict[str, Any]) -> str:
    raw = str(row.get("account_type") or "").strip().upper()
    code = str(row.get("account_code") or "").strip()
    if raw in ("PASIVO", "LIABILITY"):
        return "PASIVO"
    if raw in ("PATRIMONIO", "EQUITY"):
        return "PATRIMONIO"
    if raw in ("INGRESO", "REVENUE", "INCOME"):
        return "INGRESO"
    if raw in ("COSTO", "COST"):
        return "COSTO"
    if raw in ("GASTO", "EXPENSE"):
        return "GASTO"
    if raw in ("ACTIVO", "ASSET"):
        return "ACTIVO"
    if code.startswith("2"):
        return "PASIVO"
    if code.startswith("3"):
        return "PATRIMONIO"
    if code.startswith("4"):
        return "INGRESO"
    if code.startswith("6"):
        return "COSTO"
    if code.startswith("5"):
        return "GASTO"
    return "ACTIVO"


def _normal_delta(row: Dict[str, Any], credit_nature: bool) -> float:
    debit = _num(row.get("debit"))
    credit = _num(row.get("credit"))
    return credit - debit if credit_nature else debit - credit


def build_account_analytic(
    all_rows: List[Dict[str, Any]],
    period_from: str,
    period_to: str,
    account_filter: str,
) -> Dict[str, Any]:
    if not isinstance(all_rows, list):
        raise ValueError("Analítico de cuenta esperaba una lista de líneas")
    text = str(account_filter or "").strip().lower()
    if not text:
        raise ValueError("Seleccione o escriba una cuenta para el analítico")

    def matches(row: Dict[str, Any]) -> bool:
        code = str(row.get("account_code") or "").strip().lower()
        name = str(row.get("account_name") or "").strip().lower()
        label = f"{code} - {name}"
        return code == text or code.startswith(text) or text in name or text in label

    matched = [row for row in all_rows if isinstance(row, dict) and matches(row)]
    if not matched:
        raise ValueError("No hay movimientos para la cuenta seleccionada")

    movement_rows = [
        row for row in matched
        if period_from <= _period(row) <= period_to
    ]
    opening_rows = [
        row for row in matched
        if _period(row) and _period(row) < period_from
    ]
    if not movement_rows and not opening_rows:
        raise ValueError("No hay saldo ni movimientos para el rango seleccionado")

    sample = movement_rows[0] if movement_rows else opening_rows[0]
    account_type = _account_type(sample)
    credit_nature = account_type in ("PASIVO", "PATRIMONIO", "INGRESO")
    opening_balance = sum(_normal_delta(row, credit_nature) for row in opening_rows)

    running = opening_balance
    detail = []
    total_debit = 0.0
    total_credit = 0.0
    for row in sorted(movement_rows, key=lambda r: (_date_sort_value(r), r.get("entry_id") or 0, r.get("line_id") or 0)):
        debit = _num(row.get("debit"))
        credit = _num(row.get("credit"))
        running += _normal_delta(row, credit_nature)
        total_debit += debit
        total_credit += credit
        detail.append({
            "date": row.get("entry_date") or row.get("created_at"),
            "period": _period(row),
            "entry_id": row.get("entry_id"),
            "origin": row.get("origin"),
            "description": row.get("line_description") or row.get("entry_description"),
            "debit": round(debit, 2),
            "credit": round(credit, 2),
            "balance": round(running, 2),
        })

    return {
        "account_code": sample.get("account_code"),
        "account_name": sample.get("account_name"),
        "account_type": account_type,
        "period_label": period_from if period_from == period_to else f"{period_from} a {period_to}",
        "opening_balance": round(opening_balance, 2),
        "total_debit": round(total_debit, 2),
        "total_credit": round(total_credit, 2),
        "closing_balance": round(running, 2),
        "rows": detail,
    }


def export_account_analytic_excel(analytic_or_rows: Any, period_from: str | None = None, period_to: str | None = None, account_filter: str | None = None):
    if isinstance(analytic_or_rows, dict):
        analytic = analytic_or_rows
    else:
        analytic = build_account_analytic(analytic_or_rows, period_from or "", period_to or "", account_filter or "")

    file_path = filedialog.asksaveasfilename(
        title="Guardar Analítico de Cuenta",
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")]
    )
    if not file_path:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Analítico Cuenta"
    widths = [14, 12, 12, 16, 52, 16, 16, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    ws.merge_cells("A1:H1")
    ws["A1"] = "ANALÍTICO DE CUENTA"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = f"{analytic.get('account_code')} - {analytic.get('account_name')} | {analytic.get('period_label')}"
    ws["A2"].alignment = Alignment(horizontal="center")

    summary = [
        ("Tipo", analytic.get("account_type")),
        ("Saldo inicial", analytic.get("opening_balance", 0)),
        ("Debe periodo", analytic.get("total_debit", 0)),
        ("Haber periodo", analytic.get("total_credit", 0)),
        ("Saldo final", analytic.get("closing_balance", 0)),
    ]
    row = 4
    for label, value in summary:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        if isinstance(value, (int, float)):
            ws.cell(row=row, column=2).number_format = "#,##0.00"
        row += 1

    row += 1
    headers = ["Fecha", "Periodo", "Asiento", "Origen", "Detalle", "Debe", "Haber", "Saldo"]
    ws.append(headers)
    header_row = row
    fill = PatternFill("solid", fgColor="003A75")
    for col in range(1, 9):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    row += 1

    for item in analytic.get("rows") or []:
        ws.append([
            item.get("date"),
            item.get("period"),
            item.get("entry_id"),
            item.get("origin"),
            item.get("description"),
            item.get("debit", 0),
            item.get("credit", 0),
            item.get("balance", 0),
        ])
        for col in range(6, 9):
            ws.cell(row=row, column=col).number_format = "#,##0.00"
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="right")
        row += 1

    ws.append(["", "", "", "", "TOTALES", analytic.get("total_debit", 0), analytic.get("total_credit", 0), analytic.get("closing_balance", 0)])
    for col in range(1, 9):
        ws.cell(row=row, column=col).font = Font(bold=True)
        if col >= 6:
            ws.cell(row=row, column=col).number_format = "#,##0.00"

    for sheet_row in ws.iter_rows(min_row=1, max_row=row, min_col=1, max_col=8):
        for cell in sheet_row:
            if cell.value is not None:
                cell.border = BORDER

    wb.save(file_path)
    return file_path
