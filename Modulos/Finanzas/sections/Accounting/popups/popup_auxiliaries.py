import tkinter as tk
from datetime import date
from pathlib import Path
from tkinter import ttk, messagebox, simpledialog, filedialog

from api_client import (
    apply_accounting_auxiliary_transaction_api,
    bulk_update_accounting_auxiliary_documents_api,
    create_accounting_auxiliary_document_api,
    create_accounting_auxiliary_entity_api,
    get_accounting_accounts_api,
    get_accounting_auxiliary_aging_api,
    get_accounting_auxiliary_documents_api,
    get_accounting_auxiliary_entities_api,
    get_accounting_auxiliary_reconciliation_api,
    get_accounting_auxiliary_reconciliation_details_api,
    get_accounting_auxiliary_settings_api,
    sync_accounting_auxiliaries_api,
    update_accounting_auxiliary_document_api,
    update_accounting_auxiliary_entity_api,
    update_accounting_auxiliary_setting_api,
)
from session_context import get_user


TYPES = ("CUSTOMER", "SUPPLIER", "BANK", "EMPLOYEE", "TAX", "RETENTION", "ASSET", "ADVANCE", "LOAN")


class PopupAccountingAuxiliaries(tk.Toplevel):
    def __init__(self, parent, period=None):
        super().__init__(parent)
        self.title("Auxiliares contables")
        self.geometry("1220x720")
        self.transient(parent)
        self.grab_set()
        self.period = period
        self.entity_type = tk.StringVar(value="CUSTOMER")
        self.search = tk.StringVar()
        self.mapping = tk.StringVar()
        self.include_closed = tk.BooleanVar(value=False)
        self.accounts = []
        self.account_map = {}
        self.settings = {}
        self._build_ui()
        self._load_accounts()
        self._refresh_all()

    def _build_ui(self):
        toolbar = ttk.Frame(self, padding=10)
        toolbar.pack(fill="x")
        ttk.Label(toolbar, text="Tipo de auxiliar").pack(side="left")
        cmb = ttk.Combobox(toolbar, textvariable=self.entity_type, values=TYPES, state="readonly", width=14)
        cmb.pack(side="left", padx=6)
        cmb.bind("<<ComboboxSelected>>", lambda _e: self._refresh_all())
        ttk.Label(toolbar, text="Buscar").pack(side="left", padx=(15, 2))
        search = ttk.Entry(toolbar, textvariable=self.search, width=24)
        search.pack(side="left")
        search.bind("<Return>", lambda _e: self._load_entities())
        ttk.Button(toolbar, text="Buscar", command=self._load_entities).pack(side="left", padx=4)
        ttk.Checkbutton(toolbar, text="Incluir cerrados", variable=self.include_closed, command=self._refresh_all).pack(side="left", padx=10)
        ttk.Button(toolbar, text="Exportar Excel", command=self._export_current_type).pack(side="left", padx=4)
        ttk.Button(toolbar, text="Sincronizar fuentes", command=self._sync).pack(side="right", padx=4)
        ttk.Button(toolbar, text="Nuevo auxiliar", command=self._new_entity).pack(side="right", padx=4)

        mapping = ttk.LabelFrame(self, text="Cuenta de control del mayor", padding=8)
        mapping.pack(fill="x", padx=10, pady=(0, 8))
        self.cmb_mapping = ttk.Combobox(mapping, textvariable=self.mapping, state="readonly", width=55)
        self.cmb_mapping.pack(side="left", padx=5)
        ttk.Button(mapping, text="Guardar mapeo", command=self._save_mapping).pack(side="left", padx=5)
        ttk.Label(mapping, text="La conciliación compara el saldo abierto con esta cuenta.").pack(side="left", padx=15)

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.entities_tab = ttk.Frame(notebook)
        self.recon_tab = ttk.Frame(notebook)
        self.detail_recon_tab = ttk.Frame(notebook)
        self.aging_tab = ttk.Frame(notebook)
        notebook.add(self.entities_tab, text="Auxiliares y documentos")
        notebook.add(self.detail_recon_tab, text="Detalle real vs mayor")
        notebook.add(self.recon_tab, text="Conciliación con mayor")
        notebook.add(self.aging_tab, text="Antigüedad de saldos")
        self._build_entities_tab()
        self._build_recon_tab()
        self._build_detail_recon_tab()
        self._build_aging_tab()

    def _tree(self, parent, columns, headers, widths=None):
        tree = ttk.Treeview(parent, columns=columns, show="headings", selectmode="extended")
        for index, (col, header) in enumerate(zip(columns, headers)):
            tree.heading(col, text=header)
            tree.column(col, width=(widths or {}).get(col, 120))
        scroll = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scroll.set)
        tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        return tree

    def _build_entities_tab(self):
        upper = ttk.LabelFrame(self.entities_tab, text="Maestro y saldos", padding=5)
        upper.pack(fill="both", expand=True, pady=(5, 2))
        cols = ("id", "code", "name", "currency", "account", "docs", "balance")
        self.entities_tree = self._tree(upper, cols,
            ("ID", "Código", "Nombre", "Moneda", "Cuenta control", "Documentos", "Saldo abierto"),
            {"id": 60, "code": 110, "name": 280, "account": 120, "docs": 90, "balance": 130})
        self.entities_tree.bind("<<TreeviewSelect>>", lambda _e: self._load_documents())

        actions = ttk.Frame(self.entities_tab)
        actions.pack(fill="x", pady=4)
        ttk.Button(actions, text="Aplicar pago / movimiento", command=self._apply_transaction).pack(side="right", padx=5)
        ttk.Button(actions, text="Cerrar seleccionados", command=lambda: self._bulk_document_status("CLOSED")).pack(side="right", padx=5)
        ttk.Button(actions, text="Reabrir seleccionados", command=lambda: self._bulk_document_status("OPEN")).pack(side="right", padx=5)
        ttk.Button(actions, text="Editar documento", command=self._edit_document).pack(side="right", padx=5)
        ttk.Button(actions, text="Agregar documento manual", command=self._new_document).pack(side="right")
        ttk.Button(actions, text="Editar auxiliar", command=self._edit_entity).pack(side="right", padx=5)
        lower = ttk.LabelFrame(self.entities_tab, text="Documentos del auxiliar seleccionado", padding=5)
        lower.pack(fill="both", expand=True, pady=(2, 5))
        dcols = ("id", "number", "type", "issue", "due", "currency", "original", "open", "status", "overdue")
        self.documents_tree = self._tree(lower, dcols,
            ("ID", "Documento", "Tipo", "Emisión", "Vence", "Moneda", "Original", "Saldo", "Estado", "Días vencido"),
            {"id": 55, "number": 170, "type": 110, "issue": 90, "due": 90, "original": 110, "open": 110})

    def _build_recon_tab(self):
        frame = ttk.Frame(self.recon_tab, padding=5)
        frame.pack(fill="both", expand=True)
        cols = ("type", "account", "aux", "ledger", "difference", "status")
        self.recon_tree = self._tree(frame, cols,
            ("Auxiliar", "Cuenta control", "Saldo auxiliar", "Saldo mayor", "Diferencia", "Estado"),
            {"type": 130, "account": 130, "aux": 150, "ledger": 150, "difference": 150, "status": 120})
        self.recon_tree.tag_configure("OK", foreground="#167c3a")
        self.recon_tree.tag_configure("DIFFERENCE", foreground="#b42318")
        self.recon_tree.tag_configure("UNMAPPED", foreground="#946200")
        self.recon_tree.tag_configure("FX_REQUIRED", foreground="#6b4eff")

    def _build_detail_recon_tab(self):
        frame = ttk.Frame(self.detail_recon_tab, padding=5)
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text="Esta vista compara cada documento abierto contra su asiento o cuenta de control. Si dice diferencia, revise moneda, pago aplicado, cuenta contable o documento fuente.",
            wraplength=1200,
        ).pack(anchor="w", pady=(0, 6))
        cols = ("type", "entity", "document", "doc_type", "currency", "account",
                "aux", "ledger", "difference", "scope", "status")
        self.detail_recon_tree = self._tree(frame, cols,
            ("Tipo", "Auxiliar", "Documento", "Clase", "Moneda", "Cuenta",
             "Saldo auxiliar", "Saldo mayor", "Diferencia", "Alcance", "Estado"),
            {"type": 95, "entity": 240, "document": 170, "doc_type": 130, "currency": 75,
             "account": 125, "aux": 130, "ledger": 130, "difference": 130, "scope": 145, "status": 110})
        self.detail_recon_tree.tag_configure("OK", foreground="#167c3a")
        self.detail_recon_tree.tag_configure("DIFFERENCE", foreground="#b42318")
        self.detail_recon_tree.tag_configure("UNMAPPED", foreground="#946200")

    def _build_aging_tab(self):
        frame = ttk.Frame(self.aging_tab, padding=5)
        frame.pack(fill="both", expand=True)
        cols = ("code", "name", "currency", "current", "d30", "d60", "d90", "over90", "total")
        self.aging_tree = self._tree(frame, cols,
            ("Código", "Nombre", "Moneda", "Al día", "1-30", "31-60", "61-90", ">90", "Total"),
            {"name": 250, "current": 110, "d30": 100, "d60": 100, "d90": 100, "over90": 100, "total": 120})

    @staticmethod
    def _amount(value):
        if value is None:
            return "—"
        return f"{float(value or 0):,.2f}"

    def _load_accounts(self):
        self.accounts = get_accounting_accounts_api()
        self.account_map = {f"{a['account_code']} - {a['account_name']}": a["account_code"] for a in self.accounts}
        self.cmb_mapping["values"] = [""] + list(self.account_map)

    def _refresh_all(self):
        try:
            self.settings = {row["entity_type"]: row for row in get_accounting_auxiliary_settings_api()}
            setting = self.settings.get(self.entity_type.get()) or {}
            code = setting.get("control_account_code")
            label = next((key for key, value in self.account_map.items() if value == code), "")
            self.mapping.set(label)
            self._load_entities()
            self._load_reconciliation()
            self._load_detail_reconciliation()
            self._load_aging()
        except Exception as exc:
            messagebox.showerror("Auxiliares", str(exc))

    def _load_entities(self):
        rows = get_accounting_auxiliary_entities_api(
            self.entity_type.get(),
            self.search.get().strip() or None,
            self.include_closed.get(),
        )
        self.entities_tree.delete(*self.entities_tree.get_children())
        self.documents_tree.delete(*self.documents_tree.get_children())
        for row in rows:
            self.entities_tree.insert("", "end", values=(row["id"], row["entity_code"], row["entity_name"],
                row["currency_code"], row.get("effective_control_account") or "SIN MAPEO",
                row.get("open_document_count", row["document_count"]), self._amount(row["open_balance"])))

    def _load_documents(self):
        selected = self.entities_tree.selection()
        if not selected: return
        entity_id = self.entities_tree.item(selected[0], "values")[0]
        self.documents_tree.delete(*self.documents_tree.get_children())
        for row in get_accounting_auxiliary_documents_api(entity_id, self.include_closed.get()):
            self.documents_tree.insert("", "end", values=(row["id"], row["document_number"], row["document_type"], row.get("issue_date") or "",
                row.get("due_date") or "", row["currency_code"], self._amount(row["original_amount"]),
                self._amount(row["open_amount"]), row["status"], row.get("days_overdue") or 0))

    def _load_reconciliation(self):
        self.recon_tree.delete(*self.recon_tree.get_children())
        for row in get_accounting_auxiliary_reconciliation_api(self.period):
            self.recon_tree.insert("", "end", values=(row["entity_type"], row.get("control_account_code") or "SIN MAPEO",
                self._amount(row["auxiliary_balance"]), self._amount(row["ledger_balance"]),
                self._amount(row["difference"]), row["status"]), tags=(row["status"],))

    def _load_detail_reconciliation(self):
        self.detail_recon_tree.delete(*self.detail_recon_tree.get_children())
        rows = get_accounting_auxiliary_reconciliation_details_api(
            self.entity_type.get(), self.period, self.include_closed.get()
        )
        for row in rows:
            entity = f"{row.get('entity_code') or ''} - {row.get('entity_name') or ''}".strip(" -")
            self.detail_recon_tree.insert("", "end", values=(
                row.get("entity_type") or "",
                entity,
                row.get("document_number") or "",
                row.get("document_type") or "",
                row.get("currency_code") or "",
                row.get("control_account_code") or "SIN MAPEO",
                self._amount(row.get("auxiliary_balance")),
                self._amount(row.get("ledger_balance")),
                self._amount(row.get("difference")),
                self._scope_label(row.get("ledger_scope")),
                self._status_label(row.get("status")),
            ), tags=(row.get("status") or "",))

    @staticmethod
    def _scope_label(scope):
        return {
            "DOCUMENT_SOURCE": "Documento fuente",
            "ACCOUNTING_LINE": "Linea contable",
            "CONTROL_ACCOUNT": "Cuenta control",
            "NO_ACCOUNT": "Sin cuenta mapeada",
        }.get(scope or "", scope or "")

    @staticmethod
    def _status_label(status):
        return {
            "OK": "OK",
            "DIFFERENCE": "Diferencia: revisar",
            "UNMAPPED": "Sin mapeo contable",
            "FX_REQUIRED": "Requiere TC/revaluacion",
        }.get(status or "", status or "")

    def _load_aging(self):
        self.aging_tree.delete(*self.aging_tree.get_children())
        for row in get_accounting_auxiliary_aging_api(self.entity_type.get(), date.today().isoformat()):
            self.aging_tree.insert("", "end", values=(row["entity_code"], row["entity_name"], row["currency_code"],
                self._amount(row["current"]), self._amount(row["days_1_30"]), self._amount(row["days_31_60"]),
                self._amount(row["days_61_90"]), self._amount(row["over_90"]), self._amount(row["total"])))

    def _sync(self):
        try:
            result = sync_accounting_auxiliaries_api()
            self._refresh_all()
            messagebox.showinfo("Auxiliares", f"Sincronización completada: {result.get('synced')}")
        except Exception as exc: messagebox.showerror("Auxiliares", str(exc))

    def _save_mapping(self):
        try:
            code = self.account_map.get(self.mapping.get(), "")
            update_accounting_auxiliary_setting_api(self.entity_type.get(), code, get_user() or "unknown")
            self._refresh_all()
        except Exception as exc: messagebox.showerror("Mapeo contable", str(exc))

    def _new_entity(self):
        code = simpledialog.askstring("Nuevo auxiliar", "Código:", parent=self)
        if not code: return
        name = simpledialog.askstring("Nuevo auxiliar", "Nombre o descripción:", parent=self)
        if not name: return
        try:
            create_accounting_auxiliary_entity_api({"entity_type": self.entity_type.get(), "entity_code": code,
                "entity_name": name, "user": get_user() or "unknown"})
            self._load_entities()
        except Exception as exc: messagebox.showerror("Nuevo auxiliar", str(exc))

    def _edit_entity(self):
        selected = self.entities_tree.selection()
        if not selected:
            messagebox.showwarning("Auxiliar", "Seleccione un auxiliar."); return
        values = self.entities_tree.item(selected[0], "values")
        entity_id, code, name, currency, account = values[0], values[1], values[2], values[3], values[4]
        new_code = simpledialog.askstring("Editar auxiliar", "Código:", initialvalue=code, parent=self)
        if not new_code: return
        new_name = simpledialog.askstring("Editar auxiliar", "Nombre:", initialvalue=name, parent=self)
        if not new_name: return
        new_currency = simpledialog.askstring("Editar auxiliar", "Moneda:", initialvalue=currency, parent=self)
        new_account = simpledialog.askstring("Editar auxiliar", "Cuenta control:", initialvalue="" if account == "SIN MAPEO" else account, parent=self)
        try:
            update_accounting_auxiliary_entity_api(entity_id, {
                "entity_code": new_code,
                "entity_name": new_name,
                "currency_code": (new_currency or "CRC").upper(),
                "control_account_code": new_account or None,
            })
            self._refresh_all()
        except Exception as exc:
            messagebox.showerror("Editar auxiliar", str(exc))

    def _new_document(self):
        selected = self.entities_tree.selection()
        if not selected:
            messagebox.showwarning("Documento", "Seleccione un auxiliar."); return
        entity_id = self.entities_tree.item(selected[0], "values")[0]
        number = simpledialog.askstring("Documento", "Número o referencia:", parent=self)
        amount = simpledialog.askfloat("Documento", "Monto original:", parent=self, minvalue=0.01)
        if not number or amount is None: return
        due = simpledialog.askstring("Documento", "Fecha de vencimiento (YYYY-MM-DD), opcional:", parent=self)
        try:
            create_accounting_auxiliary_document_api(entity_id, {"document_type": "OTHER", "document_number": number,
                "issue_date": date.today().isoformat(), "due_date": due or None, "original_amount": str(amount),
                "open_amount": str(amount), "currency_code": "CRC"})
            self._load_documents(); self._load_entities(); self._load_reconciliation(); self._load_detail_reconciliation(); self._load_aging()
        except Exception as exc: messagebox.showerror("Documento", str(exc))

    def _edit_document(self):
        selected = self.documents_tree.selection()
        if not selected:
            messagebox.showwarning("Documento", "Seleccione un documento."); return
        if len(selected) > 1:
            messagebox.showwarning("Documento", "Seleccione solo un documento para editar detalle."); return
        values = self.documents_tree.item(selected[0], "values")
        document_id = values[0]
        number = simpledialog.askstring("Editar documento", "Número:", initialvalue=values[1], parent=self)
        if not number: return
        doc_type = simpledialog.askstring("Editar documento", "Tipo:", initialvalue=values[2], parent=self)
        currency = simpledialog.askstring("Editar documento", "Moneda:", initialvalue=values[5], parent=self)
        original = simpledialog.askstring("Editar documento", "Monto original:", initialvalue=str(values[6]).replace(",", ""), parent=self)
        open_amount = simpledialog.askstring("Editar documento", "Saldo abierto:", initialvalue=str(values[7]).replace(",", ""), parent=self)
        status = simpledialog.askstring("Editar documento", "Estado (OPEN/CLOSED/VOID):", initialvalue=values[8], parent=self)
        try:
            update_accounting_auxiliary_document_api(document_id, {
                "document_number": number,
                "document_type": (doc_type or "OTHER").upper(),
                "currency_code": (currency or "CRC").upper(),
                "original_amount": original,
                "open_amount": open_amount,
                "status": (status or "OPEN").upper(),
            })
            self._load_documents(); self._load_entities(); self._load_reconciliation(); self._load_detail_reconciliation(); self._load_aging()
        except Exception as exc:
            messagebox.showerror("Editar documento", str(exc))

    def _bulk_document_status(self, status):
        selected = self.documents_tree.selection()
        if not selected:
            messagebox.showwarning("Corrección masiva", "Seleccione uno o varios documentos."); return
        ids = [self.documents_tree.item(iid, "values")[0] for iid in selected]
        label = "cerrar" if status == "CLOSED" else "reabrir"
        if not messagebox.askyesno("Corrección masiva", f"¿Desea {label} {len(ids)} documento(s)?", parent=self):
            return
        try:
            payload = {"status": status}
            if status == "OPEN":
                amount = simpledialog.askstring("Reabrir documentos", "Saldo abierto para los documentos seleccionados:", parent=self)
                if amount is None: return
                payload["open_amount"] = amount
            result = bulk_update_accounting_auxiliary_documents_api(ids, payload)
            self._load_documents(); self._load_entities(); self._load_reconciliation(); self._load_detail_reconciliation(); self._load_aging()
            messagebox.showinfo("Corrección masiva", f"Actualizados: {result.get('updated')}")
        except Exception as exc:
            messagebox.showerror("Corrección masiva", str(exc))

    def _apply_transaction(self):
        selected = self.documents_tree.selection()
        if not selected:
            messagebox.showwarning("Movimiento", "Seleccione un documento."); return
        values = self.documents_tree.item(selected[0], "values")
        document_id = values[0]
        amount = simpledialog.askfloat("Movimiento", "Monto a reducir del saldo:", parent=self, minvalue=0.01)
        if amount is None: return
        reference = simpledialog.askstring("Movimiento", "Referencia del pago o ajuste:", parent=self)
        try:
            apply_accounting_auxiliary_transaction_api(document_id, {"transaction_type": "PAYMENT", "effect": "REDUCE",
                "amount": str(amount), "reference": reference, "user": get_user() or "unknown"})
            self._load_documents(); self._load_entities(); self._load_reconciliation(); self._load_detail_reconciliation(); self._load_aging()
        except Exception as exc: messagebox.showerror("Movimiento", str(exc))

    def _export_current_type(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except Exception as exc:
            messagebox.showerror("Exportar Excel", f"No se pudo cargar openpyxl:\n{exc}", parent=self)
            return
        typ = self.entity_type.get()
        path = filedialog.asksaveasfilename(
            parent=self,
            title="Exportar auxiliar contable",
            defaultextension=".xlsx",
            initialfile=f"auxiliar_contable_{typ.lower()}_{date.today().isoformat()}.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        wb = Workbook()
        wb.remove(wb.active)

        def add_tree_sheet(title, tree):
            ws = wb.create_sheet(title[:31])
            headers = [tree.heading(col)["text"] for col in tree["columns"]]
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for iid in tree.get_children():
                ws.append(list(tree.item(iid, "values")))
            for idx, header in enumerate(headers, start=1):
                width = max(len(str(header)) + 2, 12)
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    width = max(width, min(len(str(row[0].value or "")) + 2, 45))
                ws.column_dimensions[get_column_letter(idx)].width = width
            ws.freeze_panes = "A2"

        def add_rows_sheet(title, headers, rows):
            ws = wb.create_sheet(title[:31])
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for row in rows:
                ws.append(row)
            for idx, header in enumerate(headers, start=1):
                width = max(len(str(header)) + 2, 12)
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    width = max(width, min(len(str(row[0].value or "")) + 2, 45))
                ws.column_dimensions[get_column_letter(idx)].width = width
            ws.freeze_panes = "A2"

        doc_rows = []
        for iid in self.entities_tree.get_children():
            entity_values = self.entities_tree.item(iid, "values")
            entity_id = entity_values[0]
            entity_label = f"{entity_values[1]} - {entity_values[2]}"
            for doc in get_accounting_auxiliary_documents_api(entity_id, self.include_closed.get()):
                doc_rows.append([
                    typ,
                    entity_label,
                    doc.get("id"),
                    doc.get("document_number"),
                    doc.get("document_type"),
                    doc.get("issue_date"),
                    doc.get("due_date"),
                    doc.get("currency_code"),
                    doc.get("original_amount"),
                    doc.get("open_amount"),
                    doc.get("status"),
                    doc.get("source_table"),
                    doc.get("source_id"),
                ])

        add_tree_sheet("Maestro", self.entities_tree)
        add_rows_sheet("Documentos tipo", [
            "Tipo", "Auxiliar", "ID", "Documento", "Clase", "Emision", "Vence",
            "Moneda", "Original", "Saldo abierto", "Estado", "Fuente", "Fuente ID",
        ], doc_rows)
        add_tree_sheet("Documentos seleccion", self.documents_tree)
        add_tree_sheet("Detalle vs mayor", self.detail_recon_tree)
        add_tree_sheet("Conciliacion", self.recon_tree)
        add_tree_sheet("Antiguedad", self.aging_tree)
        try:
            wb.save(path)
            messagebox.showinfo("Exportar Excel", f"Archivo generado:\n{Path(path)}", parent=self)
        except Exception as exc:
            messagebox.showerror("Exportar Excel", str(exc), parent=self)
