"""Explicit MV LAURA import. Dry-run by default; refuses to overwrite other rows."""
import sys, json, re, hashlib, subprocess
from pathlib import Path
from datetime import datetime, time, timedelta
from uuid import uuid5, NAMESPACE_URL
from urllib.request import Request, urlopen
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend_api'))
from openpyxl import load_workbook
from services.tally_import import driver_records, text
from routers.tally import TallyRow, plate_key

TALLY=Path(r'C:\Users\aaron\Desktop\Tally Control MV LAURA 16 10 AM.xlsx')
DRIVERS=Path(r'C:\Users\aaron\Downloads\AGROPECUARIA EL SURCO-MV LAURA EMB.13-26 MAIZ&DDGS- SURCO DOS PINOS TERNERINA G PIG (1).xls')
BASE='https://api-som-fastapi-production-e66d.up.railway.app'
HEADERS={'X-User':'admin','X-Role':'master','X-Company-Code':'MCI-CR'}
def api(path, method='GET', data=None, headers=None):
    raw=data if isinstance(data,bytes) else json.dumps(data).encode() if data is not None else None
    command=['curl.exe','-sS','--fail-with-body','--max-time','120','-X',method]
    for key,value in {**HEADERS,'Content-Type':'application/json',**(headers or {})}.items():command.extend(['-H',f'{key}: {value}'])
    if raw is not None:command.extend(['--data-binary','@-'])
    command.append(BASE+path)
    response=subprocess.run(command,input=raw,capture_output=True)
    if response.returncode:raise RuntimeError(response.stdout.decode(errors='replace')+response.stderr.decode(errors='replace'))
    return json.loads(response.stdout)

def temporal(value, kind, notes, location):
    if value in (None,''):return ''
    if kind=='date' and isinstance(value,datetime):return value.date().isoformat()
    if kind!='date' and isinstance(value,time):return value.isoformat()
    raw=text(value)
    if kind=='date':
        m=re.fullmatch(r'(\d{1,2})/(\d{2})/?(\d{4})',raw)
        if m and 2000<=int(m[3])<=2100:
            try:
                result=datetime(int(m[3]),int(m[2]),int(m[1])).date().isoformat()
                notes.append(f'{location} original: {raw}')
                return result
            except ValueError:pass
    if kind!='date':
        clean=raw.replace(':;',':').replace(':_',':').replace(';',':').replace('.',':').strip()
        m=re.fullmatch(r'(\d{1,2}):(\d{2})(?::(\d{2}))?',clean)
        if m and int(m[1])<24 and int(m[2])<60 and int(m[3] or 0)<60:
            if clean!=raw:notes.append(f'{location} original: {raw}')
            return f'{int(m[1]):02}:{m[2]}:{m[3] or "00"}'
    notes.append(f'REVISAR {location} original: {raw}')
    return ''

def prepare():
    drivers=driver_records(DRIVERS.read_bytes(), 'Placas Aut')
    wb=load_workbook(TALLY,read_only=True,data_only=True)
    sheets={};issues=[]
    digest=hashlib.sha256(TALLY.read_bytes()).hexdigest()
    for sheet in wb:
        m=re.fullmatch(r'Bodega (\d+)',sheet.title)
        if not m:continue
        hold=int(m[1]);rows=[]
        for n,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            v=list(values)+[None]*14
            if not any(x is not None and text(x) for x in v[2:14]):continue
            notes=[text(v[13])] if v[13] else []
            row={'id':str(uuid5(NAMESPACE_URL,f'laura:{digest}:{hold}:{n}')),'number':'',
                 'date':temporal(v[2],'date',notes,'Fecha'), 'entry':temporal(v[3],'time',notes,'Entrada'),
                 'exit':temporal(v[4],'time',notes,'Salida')}
            for field,i in [('spc',5),('company',6),('ticket',7),('guide',8),('seal',9),('plate',10),('stowage',12)]:
                row[field]=text(v[i])
            row['plate']=plate_key(row['plate'])
            if row['seal'].casefold() not in ('','si','sí','no'):
                notes.append('REVISAR Guia Sello original: '+row['seal']);row['seal']=''
            row['notes']=' | '.join(notes)
            row=TallyRow(**row).model_dump(mode='json');rows.append(row)
            if any('original:' in note for note in notes):issues.append({'hold':hold,'row':n,'notes':row['notes']})
        sheets[hold]=rows
    wb.close()
    return drivers,sheets,issues

def main():
    drivers,sheets,issues=prepare()
    boot=api('/tally/bootstrap?q=LAURA')
    projects=[p for p in boot['projects'] if p['name']=='MV LAURA' and p['service_id']==519]
    assert len(projects)==1,'Proyecto ambiguo'
    project=projects[0];pid=project['id'];assert set(project['holds'])==set(sheets)
    before={str(h):api(f'/tally/projects/{pid}/sheets/{h}') for h in sheets}
    report={'project':project,'drivers':len(drivers),'rows':{h:len(r) for h,r in sheets.items()},'issues':issues,
            'unmatched_plates':sorted({r['plate'] for rows in sheets.values() for r in rows if r['plate'] and r['plate'] not in drivers}),
            'existing_rows':{h:len(d['rows']) for h,d in before.items()}}
    print(json.dumps(report,ensure_ascii=True,indent=2))
    out=Path(__file__).parent/'laura-import';out.mkdir(exist_ok=True)
    (out/'prepared.json').write_text(json.dumps(sheets,ensure_ascii=True,indent=2),encoding='utf-8')
    (out/'report.json').write_text(json.dumps(report,ensure_ascii=True,indent=2),encoding='utf-8')
    if '--apply' not in sys.argv:return
    for h,rows in sheets.items():
        old=before[str(h)]['rows']
        if old and {r['id'] for r in old}!={r['id'] for r in rows}:raise RuntimeError(f'Bodega {h} contiene datos ajenos; no sobrescribir')
    snapshot=out/'before.json'
    if not snapshot.exists():snapshot.write_text(json.dumps(before,ensure_ascii=True,indent=2),encoding='utf-8')
    boundary='som-laura-import-boundary'
    content=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="drivers.xls"\r\nContent-Type: application/vnd.ms-excel\r\n\r\n'.encode()+DRIVERS.read_bytes()+f'\r\n--{boundary}--\r\n'.encode())
    print('catalog',api(f'/tally/projects/{pid}/catalog/import?sheet_name=Placas%20Aut','POST',content,{'Content-Type':f'multipart/form-data; boundary={boundary}'}))
    for company in sorted({r['company'] for rows in sheets.values() for r in rows if r['company']}):api('/tally/companies','POST',{'name':company})
    for h,rows in sheets.items():
        old=before[str(h)]
        if old['rows']:continue
        print('saved',h,api(f'/tally/projects/{pid}/sheets/{h}','PUT',{'revision':old['revision'],'rows':rows}))
    for h,rows in sheets.items():
        actual=api(f'/tally/projects/{pid}/sheets/{h}')['rows']
        expected=[{**r,'driver':drivers.get(r['plate'],'')} for r in rows]
        assert actual==expected,f'Verification failed hold {h}'
    print('VERIFIED all rows and project driver assignments')

if __name__=='__main__':main()
