"""GET-only post-deployment checks; no production business data writes."""
from io import BytesIO
from pathlib import Path
import requests
import sys
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

base='https://api-som-fastapi-production-e66d.up.railway.app'
client=requests.Session()
client.headers.update({'X-User':'admin','X-Role':'master','X-Company-Code':'MCI-CR'})
response=client.get(base+'/informes/logra/bootstrap',timeout=30)
response.raise_for_status()
boot=response.json()
assert boot['scope']=='shared'
assert sum(len(f[s]) for f in boot['forms'] for s in ('critical_questions','detailed_questions'))==390
print('Bootstrap 7 forms / 390 questions',boot['permissions'])
response=client.get(base+'/som',timeout=30);response.raise_for_status()
assert 'data-report-view="logra"' in response.text
for asset in ('logra.js','logra.css'):
    response=client.get(base+'/informes/logra/assets/'+asset,timeout=30);response.raise_for_status();assert len(response.content)>1000
ids=None
for company in ('MCI-CR','MSL-CR'):
    client.headers['X-Company-Code']=company
    response=client.get(base+'/informes/logra?kind=all',timeout=30);response.raise_for_status()
    listing=response.json()
    current_ids=[r['id'] for r in listing['data']]
    if ids is not None:assert set(ids)==set(current_ids)
    ids=current_ids
    print(company,'shared reports',listing['total'])
rows=[r for r in listing['data'] if r['answer_count']]
assert rows
rid=rows[0]['id']
response=client.get(f'{base}/informes/logra/{rid}',timeout=30);response.raise_for_status()
detail=response.json();assert len(detail['revision'])==64
print('Detail answers',len(detail['answers']))
if '--local-pdf' in sys.argv:
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend_api'))
    from services.logra_web_export import build
    target=Path(__file__).resolve().parent/'logra-qa/local-final.pdf'
    target.write_bytes(build(detail,'pdf'))
    print(target)
    raise SystemExit(0)
response=client.get(f'{base}/informes/logra/{rid}/history',timeout=30);response.raise_for_status()
print('History entries',len(response.json()['data']))
folder=Path(__file__).resolve().parent/'logra-qa';folder.mkdir(exist_ok=True)
for kind in ('pdf','docx','xlsx'):
    response=client.get(f'{base}/informes/logra/{rid}/export/{kind}',timeout=90);response.raise_for_status()
    (folder/f'live-export.{kind}').write_bytes(response.content)
    if kind=='pdf':print('PDF pages',len(PdfReader(BytesIO(response.content)).pages))
    elif kind=='docx':print('Word paragraphs',len(Document(BytesIO(response.content)).paragraphs))
    else:
        book=load_workbook(BytesIO(response.content));print('Excel sheets',book.sheetnames);book.close()
client.headers.update({'X-User':'surveyor01','X-Role':'user'})
response=client.get(base+'/informes/logra/bootstrap',timeout=30);response.raise_for_status()
print('User permissions',response.json()['permissions'])
unauth=requests.get(base+'/informes/logra/bootstrap',headers={'X-Tally-Token':'x'*43},timeout=30)
assert unauth.status_code==401
print('PASS live LOGRA GET-only checks, shared scope and Tally token isolation')
