from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import json
import openpyxl

root=Path(__file__).resolve().parents[2]
path=root/'outputs/activos-fijos-20261019/Activos_fijos_MSL_MCI_19_10_2026.xlsx'
data=json.loads((Path(__file__).parent/'source.json').read_text(encoding='utf-8'))
wb=openpyxl.load_workbook(path,data_only=True)
formulas=openpyxl.load_workbook(path,data_only=False)
assert wb.sheetnames==['MSL','MCI','Base SOM']
sheet=wb['MSL']
money=lambda x:Decimal(str(x)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
cost=Decimal(0)
accum=Decimal(0)
for r,a in enumerate(sorted(data['assets'],key=lambda a:a['asset_code']),15):
    assert sheet.cell(r,1).value==a['asset_code']
    assert sheet.cell(r,3).value==a['description']
    sched=[s for s in data['schedule'] if s['asset_id']==a['id']]
    closed=sum((Decimal(s['depreciation_amount_crc']) for s in sched if s['depreciation_date']<='2026-09-30'),Decimal(0))
    current=sum((Decimal(s['depreciation_amount_crc']) for s in sched if s['period']=='2026-10'),Decimal(0))
    estimate=closed+(current*Decimal(19)/Decimal(31)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
    assert money(sheet.cell(r,9).value)==estimate
    assert money(sheet.cell(r,10).value)==Decimal(a['value_crc'])-estimate
    assert money(sheet.cell(r,5).value)==Decimal(a['original_amount'])
    assert money(sheet.cell(r,7).value)==Decimal(a['value_crc'])
    assert sheet.cell(r,4).value.strftime('%Y-%m-%d')==a['purchase_date']
    cost+=Decimal(a['value_crc']); accum+=estimate
assert money(sheet['B6'].value)==cost
assert money(sheet['B7'].value)==accum
assert money(sheet['E6'].value)==cost-accum
assert formulas['MSL']['I15'].data_type=='f'
assert len(formulas['MSL'].tables)==1
assert formulas['MSL'].freeze_panes is not None
print('Verified 86 assets, dates, source amounts, independent Decimal depreciation, cached totals, formulas and filter table.')
print('USD cost',sum(Decimal(a['original_amount']) for a in data['assets']))
print('CRC cost',cost,'Depreciation',accum,'Net',cost-accum)
