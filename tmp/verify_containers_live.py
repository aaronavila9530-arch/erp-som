"""Post-deploy GET-only smoke test. Does not create/edit/approve reports."""
from io import BytesIO
from pathlib import Path
import requests
from openpyxl import load_workbook
from pypdf import PdfReader

base = "https://api-som-fastapi-production-e66d.up.railway.app"
client = requests.Session()
client.headers.update({"X-User":"admin", "X-Role":"master", "X-Company-Code":"MSL-CR"})
home = client.get(base + "/som", timeout=30)
home.raise_for_status()
assert '/informes/containers/assets/containers.js' in home.text, 'Containers deployment not live yet'
for username, role in (("admin","master"),("surveyor01","user")):
    client.headers.update({"X-User":username,"X-Role":role})
    response = client.get(base + "/informes/containers/bootstrap", timeout=30)
    response.raise_for_status()
    print(username, "permissions", response.json()["permissions"])
client.headers.update({"X-User":"admin","X-Role":"master"})
for company in ("MCI-CR","MSL-CR"):
    client.headers["X-Company-Code"] = company
    response = client.get(base + "/informes/containers", timeout=30)
    response.raise_for_status()
    payload = response.json()
    print(company, "reports", payload["total"])
    services = client.get(base + "/informes/containers/services", timeout=30)
    services.raise_for_status()
    print(company, "services page", len(services.json()["data"]))
    if company == 'MSL-CR' and payload['data']:
        rid = payload['data'][0]['id']
        detail = client.get(f'{base}/informes/containers/{rid}', timeout=30)
        detail.raise_for_status()
        print('Detail version', len(detail.json()['version']))
        folder = Path(__file__).resolve().parent / 'containers-qa'
        folder.mkdir(exist_ok=True)
        for kind in ('excel','pdf','presentation','unified'):
            response = client.get(f'{base}/informes/containers/{rid}/export/{kind}', timeout=180)
            if not response.ok:
                print('Export error',kind,response.status_code,response.text[:200])
            response.raise_for_status()
            suffix = 'xlsx' if kind == 'excel' else 'pdf'
            (folder / f'live-{kind}.{suffix}').write_bytes(response.content)
            if kind == 'excel':
                book = load_workbook(BytesIO(response.content))
                print(kind, 'sheets',len(book.sheetnames))
                assert not book.active['A26'].alignment.wrap_text, 'Final export layout is not deployed yet'
                assert book.active['A25'].alignment.wrap_text, 'Full-text export is not deployed yet'
                assert book.active['AB53'].value is None, 'Scope answer is not in its box yet'
                book.close()
            else:
                pdf = PdfReader(BytesIO(response.content))
                print(kind, 'pages',len(pdf.pages))
                if kind in ('presentation','unified'):
                    assert 'CONTAINER INSPECTION' in pdf.pages[0].extract_text()
        for path in ('/container-reports/statuses',f'/container-presentation/{rid}'):
            legacy = client.get(base+path,timeout=30)
            legacy.raise_for_status()
            assert isinstance(legacy.json(),(dict,list))
            print('Desktop legacy GET',path.rsplit('/',1)[0],legacy.status_code)
print('PASS GET-only live Containers smoke')
