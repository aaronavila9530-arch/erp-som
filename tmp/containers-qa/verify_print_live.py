from pathlib import Path
from urllib.request import Request, urlopen
from io import BytesIO
import json
from openpyxl import load_workbook
from pypdf import PdfReader
import pypdfium2 as pdfium

BASE = 'https://api-som-fastapi-production-e66d.up.railway.app'
folder = Path(__file__).resolve().parent
with urlopen(BASE + '/container-reports/17/excel', timeout=90) as response:
    data = response.read()
book = load_workbook(BytesIO(data))
sheet = book.active
assert sheet.page_setup.scale is None, 'Deployment is not serving the new layout yet'
assert sheet.page_setup.orientation == 'landscape'
assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
assert 'A26:AJ26' in sheet.merged_cells
book.close()
(folder / 'live.xlsx').write_bytes(data)
with urlopen(Request(BASE + '/container-reports/17/generate-pdf', method='POST'), timeout=180) as response:
    pdf_bytes = response.read()
(folder / 'live-print.pdf').write_bytes(pdf_bytes)
pdf = PdfReader(BytesIO(pdf_bytes))
for page in pdf.pages:
    assert page.mediabox.width > page.mediabox.height
text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
assert 'HLBU9375471' in text
assert '2249-1710-2026' in text
doc = pdfium.PdfDocument(pdf_bytes)
for i in range(len(doc)):
    doc[i].render(scale=1.5).to_pil().save(folder / f'live-print-{i + 1}.png')
print(json.dumps({'pages': len(pdf.pages), 'landscape': True, 'one_page_wide': True}))
