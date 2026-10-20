"""Container exports keep data and fixed-width landscape print geometry."""
from io import BytesIO
from pathlib import Path
import sys
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from services.container_report_excel_service import generate_container_report_excel
from services.container_web_export import build


class ContainerPrintTests(unittest.TestCase):
    def report(self):
        return dict(id=17, report_no='CONTAINER-17', linked_report_number='TEST-17',
                    appointment='Appointment', shippers='Shipper name ' * 10,
                    seals='123456 / 789012 / ABCDEF', scope_items=24.067,
                    cause_detail=('First paragraph with detailed observations. ' * 15 + '\n\nEnd of cause.'),
                    remarks='Observation of the inspection. ' * 45,
                    conclusion='Conclusion.\n\n' * 20,
                    picture_link='https://example.com/' + 'x' * 200)

    def check(self, sheet):
        self.assertEqual(sheet.page_setup.orientation, 'landscape')
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertIsNone(sheet.page_setup.scale)
        self.assertFalse(sheet.sheet_format.zeroHeight)
        self.assertEqual(sheet['D8'].value, 'Appointment')
        self.assertEqual(sheet['AF53'].value, 24.067)
        self.assertEqual(sheet['AF53'].number_format, 'General')
        self.assertEqual(sheet['A31'].value, self.report()['remarks'])
        self.assertEqual(sheet['A37'].value, self.report()['conclusion'])
        self.assertEqual(sheet['A42'].hyperlink.target, self.report()['picture_link'])
        for address in ('A26', 'I14', 'A31', 'A37', 'A42', 'K49'):
            self.assertTrue(sheet[address].alignment.wrap_text)
            self.assertFalse(sheet[address].alignment.shrink_to_fit)
        self.assertTrue(all((d.height or 0) < 409 for d in sheet.row_dimensions.values()))
        self.assertNotIn('Q', sheet.column_dimensions)
        self.assertEqual(sheet.oddFooter.right.text, 'Page &P of &N')

    def test_desktop_export(self):
        path = Path(generate_container_report_excel(self.report()))
        try:
            book = load_workbook(path)
            self.check(book.active)
            book.close()
        finally:
            path.unlink(missing_ok=True)

    def test_web_export(self):
        book = load_workbook(BytesIO(build(self.report(), 'Customer', 'excel')))
        self.check(book.active)
        book.close()


if __name__ == '__main__':
    unittest.main()
