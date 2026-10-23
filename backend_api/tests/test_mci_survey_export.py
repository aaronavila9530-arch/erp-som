import sys
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml.ns import qn
from services.mci_survey_export import build_word
from services.certificate_web_export import build


class ExportTests(unittest.TestCase):
    def test_template_parts_and_disabled_sections(self):
        row = dict(company_code='MCI-CR', vessel='QA VESSEL', requested_by='QA CLIENT',
                   report_number='QA-REPORT', content=dict(sections={
                       'work': dict(enabled=True, blocks=[dict(text='VISIBLE NARRATIVE')]),
                       'conclusion': dict(enabled=False, blocks=[dict(text='HIDDEN NARRATIVE')])}))
        data = build_word(row)
        doc = Document(BytesIO(data))
        self.assertEqual(len(doc.sections), 2)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('VISIBLE NARRATIVE', text)
        self.assertNotIn('HIDDEN NARRATIVE', text)
        self.assertNotIn('LINA AKSOY', text)
        self.assertEqual(doc.tables[0].cell(0, 1).text, 'QA CLIENT')
        self.assertFalse(doc.sections[0].footer._element.xpath('.//w:drawing'))
        self.assertEqual(len(doc.sections[-1].footer._element.xpath('.//w:drawing')), 2)
        for section in doc.sections:
            tags = [e.tag for e in section._sectPr]
            self.assertLess(tags.index(qn('w:headerReference')), tags.index(qn('w:pgSz')))
            self.assertIn('QA-REPORT', ''.join(section.header._element.xpath('.//w:t/text()')))
        with ZipFile(BytesIO(data)) as archive:
            self.assertTrue(any(n.startswith('word/media/') for n in archive.namelist()))

    def test_msl_original_exporters(self):
        for kind in ('condition', 'cargo_condition'):
            with self.subTest(kind=kind):
                data = build(kind, dict(company_code='MSL-CR', report_number='QA-REPORT', vessel='QA VESSEL'), 'word')
                doc = Document(BytesIO(data))
                self.assertTrue(doc.sections)
                self.assertIn('QA VESSEL', '\n'.join(p.text for p in doc.paragraphs))
                with self.assertRaises(ValueError):
                    build(kind, dict(company_code='MCI-CR'), 'word')


if __name__ == '__main__':
    unittest.main()
