from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest.mock import MagicMock
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from services.mci_container_export import build_word
from services.mci_survey_schema import SurveySave
from routers.mci_surveys_web import persist, detail
from routers.container_web import version


class ContainerTests(unittest.TestCase):
    def row(self):
        return dict(company_code='MCI-CR',kind='container',report_number='QA-77',vessel='QA-CONTAINER',
                    requested_by='QA CLIENT',content={'hierarchy_version':2,'cargo':'QA cargo',
                    'sections':{'introduction':{'enabled':True},'purpose':{'enabled':True,'blocks':[{'text':'Unique purpose'}]},
                                'damage':{'enabled':True,'blocks':[{'text':'Damage details'}]},
                                'particulars':{'enabled':True},
                                'enclosure':{'enabled':True,'blocks':[{'text':'https://example.com/photos'}]}}})

    def test_retains_format_and_no_sample_facts(self):
        doc=Document(BytesIO(build_word(self.row())))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Unique purpose',text)
        self.assertIn('QA-CONTAINER',text)
        self.assertIn('https://example.com/photos',text)
        for forbidden in ('FIGAL','HAMU4659393','MSL Marine','{{'):
            self.assertNotIn(forbidden,text)
        self.assertEqual(round(doc.sections[0].page_width.inches,2),8.5)
        self.assertIn('QA-77',doc.sections[0].header._element.xml)
        self.assertIn('NUMPAGES',doc.sections[0].footer._element.xml)

    def test_hidden_parent_and_msl_rejected(self):
        row=self.row();row['content']['sections']['introduction']['enabled']=False
        self.assertNotIn('Unique purpose','\n'.join(p.text for p in Document(BytesIO(build_word(row))).paragraphs))
        row['company_code']='MSL-CR'
        with self.assertRaises(ValueError):build_word(row)

    def test_container_save_and_version_from_stored_record(self):
        conn=MagicMock();cur=conn.cursor.return_value.__enter__.return_value
        stored=dict(id=4,kind='container',service_id=538,content={},status='Draft')
        cur.fetchone.side_effect=[dict(num_informe='QA-77',buque_contenedor='QA-CONTAINER',cliente='QA CLIENT',puerto='La Union',pais='El Salvador'),stored]
        result=persist(SurveySave(kind='container',service_id=538,content=self.row()['content']),{'permissions':{'create':True}},conn)
        self.assertEqual(result['version'],version(stored))
        self.assertEqual(cur.execute.call_args.args[1][0],'container')

    def test_detail_token_not_derived_from_enriched_fields(self):
        row=self.row();row.update(id=4,service_id=538)
        conn=MagicMock();cur=conn.cursor.return_value.__enter__.return_value
        cur.fetchone.side_effect=[row,dict(puerto='La Union',pais='El Salvador'),dict(filename='original.docx')]
        result=detail(4,{},conn)
        self.assertEqual(result['version'],version(row))
        self.assertEqual(result['data']['source_filename'],'original.docx')


if __name__=='__main__':unittest.main()
