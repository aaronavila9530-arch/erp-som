import sys
from pathlib import Path
import unittest
from unittest.mock import MagicMock, patch
from uuid import uuid4
from pydantic import ValidationError
from fastapi import HTTPException
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from routers import tally


class TallyTests(unittest.TestCase):
    def setUp(self):
        self.conn = MagicMock()
        self.cur = self.conn.cursor.return_value.__enter__.return_value
        self.ctx = dict(company="MSL-CR", user="test", editable=True)

    def test_dynamic_holds(self):
        self.assertEqual(tally.ProjectInput(name="Ship", holds=[1, 3, 5]).holds, [1, 3, 5])
        for holds in ([], [1, 1], [0], [1000]):
            with self.assertRaises(ValidationError):
                tally.ProjectInput(name="Ship", holds=holds)

    def test_types_and_identifiers(self):
        row = tally.TallyRow(id=uuid4(), number="001", plate="001234", date="2026-10-05", entry="08:00", seal="si")
        self.assertEqual(row.number, "001")
        self.assertEqual(row.seal, "Si")
        self.assertEqual(tally.plate_key(" ab-012 "), "AB012")
        for data in ({"date":"2026-02-30"}, {"entry":"25:00"}, {"seal":"maybe"}, {"driver":"forged"}):
            with self.assertRaises(ValidationError):
                tally.TallyRow(id=uuid4(), **data)
        with self.assertRaises(ValidationError):
            tally.SheetInput(revision=0, rows=[row,row])

    def test_tenant_boundary(self):
        self.cur.fetchone.return_value=None
        with self.assertRaises(HTTPException) as exc:
            tally.project(self.cur, 1, self.ctx)
        self.assertEqual(exc.exception.status_code,404)
        self.assertEqual(self.cur.execute.call_args.args[1], (1,"MSL-CR"))

    def test_read_only_rejects_mutation(self):
        with self.assertRaises(HTTPException) as exc:
            tally.writer({"editable":False})
        self.assertEqual(exc.exception.status_code,403)

    def test_conflict_rejects_save(self):
        self.cur.fetchone.side_effect=[{"holds":[3]}, {"revision":2}]
        with patch.object(tally,"schema"), self.assertRaises(HTTPException) as exc:
            tally.save_sheet(1,3,tally.SheetInput(revision=1,rows=[]),self.ctx,self.conn)
        self.assertEqual(exc.exception.status_code,409)
        self.conn.commit.assert_not_called()

    def test_cannot_remove_populated_hold(self):
        self.cur.fetchone.side_effect=[{"holds":[1,3],"revision":0},{"hold":3}]
        with patch.object(tally,"schema"), self.assertRaises(HTTPException) as exc:
            tally.update_project(1,tally.ProjectInput(name="Ship",holds=[1]),self.ctx,self.conn)
        self.assertEqual(exc.exception.status_code,409)

    def test_save_resolves_driver_on_server_and_audits(self):
        self.cur.fetchone.side_effect=[{"holds":[3]},{"revision":0}]
        self.cur.fetchall.side_effect=[[{"plate":"001","driver":"Test Driver"}],[{"name":"Carrier"}]]
        row=tally.TallyRow(id=uuid4(),plate="001",company="Carrier")
        with patch.object(tally,"schema"):
            result=tally.save_sheet(1,3,tally.SheetInput(revision=0,rows=[row]),self.ctx,self.conn)
        self.assertEqual(result["revision"],1)
        update=[c for c in self.cur.execute.call_args_list if "UPDATE tally_sheets" in c.args[0]][0]
        self.assertEqual(update.args[1][0].adapted[0]["driver"],"Test Driver")
        self.assertTrue(any("INSERT INTO tally_audit" in c.args[0] for c in self.cur.execute.call_args_list))

    def test_excel_columns_dates_and_formula_safety(self):
        out=tally.export_workbook({"holds":[1,3,5]}, {3:[dict(number="001",plate="001234",date="2026-10-05",entry="08:30:00",notes='=HYPERLINK("bad")')]})
        wb=load_workbook(out)
        self.assertEqual(wb.sheetnames,["Bodega 1","Bodega 3","Bodega 5"])
        ws=wb["Bodega 3"]
        self.assertEqual([c.value for c in ws[1]],tally.HEADERS)
        self.assertEqual(ws["K2"].value,"001234")
        self.assertEqual(ws["N2"].data_type,"s")
        self.assertEqual(ws["C2"].number_format,"dd/mm/yyyy")
        self.assertEqual(ws.freeze_panes,"D2")


if __name__ == "__main__":
    unittest.main()
