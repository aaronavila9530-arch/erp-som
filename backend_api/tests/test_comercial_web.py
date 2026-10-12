import inspect
import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi import HTTPException
from fastapi.params import Param
from routers import comercial, comercial_clients_analytics as clients, comercial_ports_analytics as ports
from routers import comercial_servicios_analytics as services, comercial_web as web
from services.commercial_access import require_commercial_detail


def invoke(function, **overrides):
    args = {}
    for name, param in inspect.signature(function).parameters.items():
        default = param.default
        args[name] = default.default if isinstance(default, Param) else default
    args.update(overrides)
    return function(**args)


class CommercialTests(unittest.TestCase):
    def connection(self):
        conn = MagicMock()
        cur = conn.cursor.return_value
        cur.__enter__.return_value = cur
        cur.fetchall.return_value = []
        cur.fetchone.return_value = {key: 0 for key in ['total', 'clientes', 'paises', 'puertos', 'facturacion_neta', 'honorarios', 'costos_operativos', 'costo_tarjetas']}
        return conn, cur

    def test_board_scopes_company_and_does_not_hide_previous_finalized_years(self):
        conn, cur = self.connection()
        invoke(comercial.comercial_board, conn=conn, estados=['Confirmado','FINALIZADO'], x_company_code='MCI-CR')
        sql, params = cur.execute.call_args.args
        self.assertEqual(params['company'], 'MCI-CR')
        self.assertIn('company_code', sql)
        self.assertIn("LOWER(estado) NOT IN", sql)
        invoke(comercial.comercial_board, conn=conn, estados=['Confirmado'], year=2025, x_company_code='MSL-CR')
        _, params = cur.execute.call_args.args
        self.assertNotIn('y_start', params)
        self.assertEqual(params['year_start'], '2025-01-01')

    def test_clients_all_queries_are_tenant_scoped_and_payments_aggregated(self):
        conn, cur = self.connection()
        invoke(clients.comercial_client_view, conn=conn, year=2026, x_company_code='MCI-CR')
        for call in cur.execute.call_args_list:
            sql, params = call.args
            self.assertIn('company_code', sql)
            self.assertEqual(params['company'], 'MCI-CR')
        self.assertIn('LEFT JOIN LATERAL', cur.execute.call_args_list[0].args[0])
        invoke(clients.get_comercial_clientes, conn=conn, nombre='NORDEN', x_company_code='MCI-CR')
        self.assertEqual(cur.execute.call_args.args[1]['company'], 'MCI-CR')

    def test_ports_table_and_kpis_filter_same_port_and_company(self):
        for function in (ports.comercial_ports_analytics, ports.comercial_ports_kpis):
            conn, cur = self.connection()
            invoke(function, conn=conn, puerto='Caldera', x_company_code='MCI-CR')
            sql, params = cur.execute.call_args.args
            self.assertIn('s.puerto = %(puerto)s', sql)
            self.assertEqual(params['company'], 'MCI-CR')
            self.assertEqual(params['puerto'], 'Caldera')

    def test_coverage_scopes_company_and_closes_cursor(self):
        conn, cur = self.connection()
        invoke(ports.get_ports_coverage, conn=conn, x_company_code='MCI-CR')
        self.assertEqual(cur.execute.call_args.args[1]['company'], 'MCI-CR')
        cur.close.assert_called_once()

    def test_service_tables_kpis_and_pareto_scope_company_and_quarter(self):
        for function in (services.servicios_analytics_by_servicio, services.servicios_no_ofrecidos,
                         services.servicios_kpis_ejecutivos, services.costos_surveyor_pareto):
            conn, cur = self.connection()
            invoke(function, conn=conn, year_from=2025, year_to=2026, quarter='Q2', x_company_code='MCI-CR')
            sql, params = cur.execute.call_args.args
            self.assertIn('company_code', sql, function.__name__)
            self.assertIn('EXTRACT(QUARTER', sql, function.__name__)
            values = list(params.values()) if isinstance(params, dict) else list(params)
            self.assertIn('MCI-CR', values)
            self.assertIn(2, values)

    def test_invalid_quarter_is_validation_error(self):
        for quarter in ['Q0', 'Q5', 'x']:
            with self.assertRaises(HTTPException) as error:
                services._period_filters([], [], 'MSL-CR', quarter)
            self.assertEqual(error.exception.status_code, 422)

    def test_nonoffered_accepts_quarter_range(self):
        conn, cur = self.connection()
        invoke(services.servicios_no_ofrecidos, conn=conn, quarter_from='Q2', quarter_to='Q3', x_company_code='MSL-CR')
        sql, params = cur.execute.call_args.args
        self.assertIn('EXTRACT(QUARTER FROM s.fecha_inicio::date) >= %s', sql)
        self.assertEqual(params[-2:], (2, 3))

    def test_metadata_is_read_only_and_surveyor_has_only_board(self):
        conn, cur = self.connection()
        with patch.object(web, 'has_permission', return_value=True):
            meta = web.metadata('admin','surveyor01','MCI-CR',conn)
        self.assertEqual([k for k,v in meta['permissions'].items() if v], ['board'])
        sql, params = cur.execute.call_args.args
        self.assertTrue(sql.startswith('SELECT'))
        self.assertEqual(params, ('MCI-CR',))

    def test_exports_include_all_rows_and_neutralize_spreadsheet_formulas(self):
        from openpyxl import load_workbook
        payload=web.TableExport(columns=['cliente','precio'], labels=['Cliente','Precio'], rows=[{'cliente':'=HYPERLINK("x")','precio':-2}]+[{'cliente':f'Cliente {i}','precio':i} for i in range(1205)])
        response=web.export_table('xlsx',payload)
        sheet=load_workbook(BytesIO(response.body)).active
        self.assertEqual(sheet.max_row,1207)
        self.assertEqual(sheet['A2'].data_type,'s')
        self.assertEqual(sheet['B2'].value,-2)
        self.assertTrue(web.export_table('csv',payload).body.startswith(b'\xef\xbb\xbf'))
        self.assertTrue(web.export_table('pdf',web.TableExport(columns=['cliente'],labels=['Cliente'],rows=[{'cliente':'A < B'}])).body.startswith(b'%PDF'))

    def test_assets_allowlist(self):
        for name in ('comercial.js','comercial.css'):
            self.assertTrue(Path(web.asset(name).path).is_file())
        with self.assertRaises(HTTPException):
            web.asset('../main.py')

    def test_board_only_restriction_is_enforced_by_api(self):
        for username in ('surveyor01', 'SURVEYOR02', 'surveyor03'):
            with self.assertRaises(HTTPException) as error:
                require_commercial_detail(username)
            self.assertEqual(error.exception.status_code, 403)
        require_commercial_detail('admin')


if __name__ == '__main__':
    unittest.main()
