import unittest
from docx import Document
from docx.oxml.ns import qn
from backend_api.services.draft_survey_layout import build_presentation, harden_written_report


class StableDraftLayoutTests(unittest.TestCase):
    def test_cover_is_editable_without_floating_shapes(self):
        data = {'draft_report_number': '2238-0410-2026',
                'word_vessel': 'MV SANTA CAROLINA V2',
                'word_survey_requested_by': 'A long client name ' * 12}
        original = dict(data)
        doc = build_presentation(data, 'Caldera, October 09, 2026')
        self.assertEqual(data, original)
        self.assertFalse(doc._element.xpath('.//wp:anchor'))
        self.assertEqual(len(doc.tables), 1)
        for table in doc.tables:
            self.assertFalse(table.autofit)
            self.assertFalse(table._tbl.xpath('.//w:trHeight'))
        self.assertEqual(doc.sections[0].left_margin.inches, .8)

    def test_prose_not_mistaken_for_columns_and_times_preserved(self):
        doc = Document()
        prose = 'A long narrative  with repeated spaces must remain a paragraph.'
        doc.add_paragraph(prose)
        doc.add_paragraph('EXTRACT TIME SHEET')
        doc.add_paragraph('Final Draft Survey\tOct 09 2026 16:10 LT.')
        harden_written_report(doc, {'word_vessel': 'MV SANTA CAROLINA V2'})
        self.assertIn(prose, [p.text for p in doc.paragraphs])
        self.assertEqual(doc.tables[0].cell(0, 1).text, 'Oct 09 2026 16:10 LT.')
        self.assertEqual([n.get(qn('w:instr')) for n in
                          doc.sections[0].footer._element.iter(qn('w:fldSimple'))],
                         ['PAGE', 'NUMPAGES'])


if __name__ == '__main__':
    unittest.main()
