"""Containers phase 1: contracts, isolation and writes without production data."""
import ast
from datetime import datetime
from pathlib import Path
import sys
import unittest
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend_api"))
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from routers import container_web as web
from services.container_web_fields import FIELDS, BOOLEAN_FIELDS, DATE_FIELDS


class ContainersTests(unittest.TestCase):
    def setUp(self):
        self.ctx = {"company": "MCI-CR", "user": "tester", "permissions": dict(view=True, edit=True, create=True, approve=True)}
        self.conn = MagicMock()
        self.cur = self.conn.cursor.return_value.__enter__.return_value
        self.row = dict(id=1, company_code="MCI-CR", linked_report_number="TEST-01", report_no="BOX-01",
                        status="Pending for review", updated_at=datetime(2026, 10, 19), remarks="Original", user="creator")

    def test_desktop_payload_field_parity(self):
        tree = ast.parse((ROOT / "Modulos/Informes/container_report_form.py").read_text(encoding="utf-8-sig"))
        method = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "_build_payload")
        payload = next(n.value for n in method.body if isinstance(n, ast.Return))
        self.assertEqual(set(FIELDS), {k.value for k in payload.keys})

    def test_all_fields_round_trip_validation(self):
        fields = {k: True if k in BOOLEAN_FIELDS else "2026-10-19T09:01" if k in DATE_FIELDS else "Example"
                  for k in FIELDS}
        cleaned = web.clean_fields(fields)
        self.assertEqual(set(cleaned), set(FIELDS))
        self.assertEqual(cleaned["init_to"], "2026-10-19 09:01")
        self.assertIs(cleaned["cause_damage"], True)

    def test_protected_and_invalid_fields_rejected(self):
        for fields in ({"status":"Approved"}, {"user":"admin"}, {"company_code":"MSL-CR"},
                       {"doc_bl":"false"}, {"final_to":"invalid"}, {"remarks":"\0"}):
            with self.subTest(fields=fields), self.assertRaises(HTTPException) as error:
                web.clean_fields(fields)
            self.assertEqual(error.exception.status_code, 422)

    def test_read_and_lock_are_tenant_scoped(self):
        self.cur.fetchone.return_value = self.row
        self.assertEqual(web.read_report(self.cur, 1, self.ctx, True), self.row)
        query, params = self.cur.execute.call_args.args
        self.assertIn("s.companies <= 1", query)
        self.assertIn("FOR UPDATE OF cr", query)
        self.assertEqual(params, ("MCI-CR", 1))

    def test_foreign_report_not_found(self):
        self.cur.fetchone.return_value = None
        with self.assertRaises(HTTPException) as error:
            web.detail(8, self.ctx, self.conn)
        self.assertEqual(error.exception.status_code, 404)

    def test_create_and_service_update_are_one_transaction(self):
        self.cur.fetchall.return_value = [{"company_code":"MCI-CR"}]
        self.cur.fetchone.side_effect = [{"present":True},self.row]
        result = web.create(web.SaveInput(fields={"linked_report_number":"TEST-01", "remarks":"Text"}, submit=True), self.ctx, self.conn)
        self.assertEqual(result["version"], web.version(self.row))
        queries = self.cur.execute.call_args_list
        self.assertIn('"company_code"', queries[2].args[0])
        self.assertIn('"user"', queries[2].args[0])
        self.assertIn("Pending for review", queries[2].args[1])
        self.assertEqual(queries[3].args[1], ("TEST-01", "MCI-CR"))
        self.conn.commit.assert_called_once()

    def test_create_supports_legacy_table_without_company_column(self):
        self.cur.fetchall.return_value = [{"company_code":"MCI-CR"}]
        self.cur.fetchone.side_effect = [{"present":False},self.row]
        web.create(web.SaveInput(fields={"linked_report_number":"TEST-01"}),self.ctx,self.conn)
        insert = self.cur.execute.call_args_list[2].args[0]
        self.assertNotIn('"company_code"', insert)
        self.assertNotIn('ALTER TABLE', ' '.join(c.args[0] for c in self.cur.execute.call_args_list))

    def test_service_ambiguity_and_other_company_cannot_create(self):
        for rows in ([], [{"company_code":"MSL-CR"}], [{"company_code":"MCI-CR"}]*2):
            self.cur.fetchall.return_value = rows
            with self.assertRaises(HTTPException):
                web.create(web.SaveInput(fields={"linked_report_number":"TEST-01"}), self.ctx, self.conn)
        self.conn.commit.assert_not_called()
        self.assertEqual(self.conn.rollback.call_count, 3)

    def test_create_rolls_back_when_service_update_fails(self):
        self.cur.fetchall.return_value = [{"company_code":"MCI-CR"}]
        self.cur.fetchone.side_effect = [{"present":True},self.row]
        self.cur.execute.side_effect = [None, None, None, RuntimeError("failure")]
        with self.assertRaises(RuntimeError):
            web.create(web.SaveInput(fields={"linked_report_number":"TEST-01"}), self.ctx, self.conn)
        self.conn.rollback.assert_called_once()
        self.conn.commit.assert_not_called()

    def test_patch_preserves_untouched_columns_and_author(self):
        self.cur.fetchone.side_effect = [self.row, {**self.row,"remarks":"Changed"}]
        result = web.update(1, web.SaveInput(fields={"remarks":"Changed"},version=web.version(self.row)), self.ctx, self.conn)
        query, params = self.cur.execute.call_args.args
        self.assertNotIn('"user"', query)
        self.assertNotIn('"vessel"', query)
        self.assertEqual(params[0], "Changed")
        self.assertEqual(result["data"]["user"], "creator")
        self.conn.commit.assert_called_once()

    def test_stale_desktop_edit_is_rejected(self):
        self.cur.fetchone.return_value = {**self.row,"remarks":"Desktop edit"}
        with self.assertRaises(HTTPException) as error:
            web.update(1, web.SaveInput(fields={"remarks":"Web edit"},version=web.version(self.row)), self.ctx, self.conn)
        self.assertEqual(error.exception.status_code, 409)
        self.conn.commit.assert_not_called()
        self.conn.rollback.assert_called_once()

    def test_approved_report_cannot_be_edited(self):
        row = {**self.row,"status":"Approved"}; self.cur.fetchone.return_value = row
        with self.assertRaises(HTTPException) as error:
            web.update(1, web.SaveInput(fields={"remarks":"Edit"},version=web.version(row)), self.ctx, self.conn)
        self.assertEqual(error.exception.status_code, 409)
        self.conn.commit.assert_not_called()

    def test_user_cannot_approve_or_save_without_permission(self):
        ctx = {**self.ctx,"permissions":dict(view=True,edit=False,create=False,approve=False)}
        for action in (lambda: web.create(web.SaveInput(fields={}),ctx,self.conn),
                       lambda: web.update(1,web.SaveInput(fields={}),ctx,self.conn),
                       lambda: web.change_status(1,web.StatusInput(version="x",status="Approved"),ctx,self.conn)):
            with self.assertRaises(HTTPException) as error: action()
            self.assertEqual(error.exception.status_code, 403)
        self.conn.cursor.assert_not_called()

    def test_status_does_not_depend_on_pdf_generation(self):
        self.cur.fetchone.side_effect = [self.row, {**self.row,"status":"Approved"}]
        result = web.change_status(1,web.StatusInput(version=web.version(self.row),status="Approved"),self.ctx,self.conn)
        self.assertEqual(result["data"]["status"], "Approved")
        self.conn.commit.assert_called_once()

    def test_roles_approval_and_denied_module(self):
        for role in ("user","admin","master"):
            with patch.object(web, "_has_visual_permission", return_value=True):
                result = web.access({"usuario":"tester","rol":role},"MCI-CR")
            self.assertEqual(result["permissions"]["approve"], role != "user")
        with patch.object(web, "_has_visual_permission", return_value=False), self.assertRaises(HTTPException):
            web.access({"usuario":"tester","rol":"user"},"MCI-CR")

    def test_excel_uses_original_template_without_database_changes(self):
        from openpyxl import load_workbook
        from services.container_report_excel_service import generate_container_report_excel
        row = {k: False if k in BOOLEAN_FIELDS else "" for k in FIELDS}
        row.update(report_no="QA-CONTAINER-ONLY", linked_report_number="QA-REPORT")
        path = Path(generate_container_report_excel(row))
        try:
            book = load_workbook(path)
            values = [c.value for sheet in book for row in sheet for c in row]
            self.assertIn("QA-CONTAINER-ONLY", values)
            book.close()
        finally:
            path.unlink(missing_ok=True)

    def test_kiosk_token_alone_is_not_a_containers_identity(self):
        app = FastAPI(); app.include_router(web.router)
        app.dependency_overrides[web.get_db] = lambda: self.conn
        with TestClient(app) as client:
            response = client.get("/informes/containers/bootstrap",headers={"X-Tally-Token":"x"*43})
        self.assertEqual(response.status_code, 401)

    def test_web_export_cleans_files_and_merges_in_order(self):
        from pypdf import PdfWriter, PdfReader
        from io import BytesIO
        from services import container_web_export as exporter
        directories = []
        def convert(source, folder):
            directories.append(folder)
            target = folder / (source.stem + '.pdf')
            writer = PdfWriter()
            writer.add_blank_page(width=100 if source.stem == 'presentation' else 200, height=200)
            writer.write(str(target))
            writer.close()
            return target
        with patch.object(exporter, 'convert', side_effect=convert):
            for kind in ('pdf', 'presentation', 'unified'):
                data = exporter.build(self.row, 'QA Client', kind)
                pdf = PdfReader(BytesIO(data))
                self.assertEqual(len(pdf.pages), 2 if kind == 'unified' else 1)
                if kind == 'unified':
                    self.assertIn('CONTAINER INSPECTION', pdf.pages[0].extract_text())
                    self.assertEqual(int(pdf.pages[1].mediabox.width), 200)
        self.assertTrue(all(not folder.exists() for folder in directories))

    def test_web_workbook_restores_template_widths_and_static_labels(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from services.container_web_export import build
        book = load_workbook(BytesIO(build({**self.row, 'appointment':'Appointment detail'}, 'QA Client', 'excel')))
        sheet = book.active
        self.assertNotIn('Q',sheet.column_dimensions)
        self.assertTrue(sheet['A26'].alignment.wrap_text)
        self.assertIn('A26:AJ26', sheet.merged_cells)
        self.assertEqual(sheet.page_setup.orientation, 'landscape')
        self.assertIsNone(sheet.page_setup.scale)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertEqual(sheet['D8'].value,'Appointment detail')
        self.assertIsNone(sheet['C8'].value)
        book.close()

    def test_export_failure_does_not_change_report(self):
        self.cur.fetchone.side_effect = [self.row, {'customer':'QA Client'}]
        with patch('services.container_web_export.build',side_effect=RuntimeError('Unavailable')):
            with self.assertRaises(HTTPException) as error:
                web.export(1,'pdf',self.ctx,self.conn)
        self.assertEqual(error.exception.status_code, 503)
        self.conn.commit.assert_not_called()
        self.assertFalse(any('UPDATE ' in call.args[0] for call in self.cur.execute.call_args_list))

    def test_web_export_preserves_long_text_and_legacy_date(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from services.container_web_export import build
        long_text = 'Full inspection observations with details. ' * 35
        row = {**self.row, 'goods_condition': long_text, 'scope_items': long_text,
               'init_inspection_datetime': 'February 03, 2026'}
        book = load_workbook(BytesIO(build(row, 'QA Client', 'excel')))
        self.assertEqual(book.active['A25'].value, long_text)
        self.assertTrue(book.active['A25'].alignment.wrap_text)
        self.assertGreater(book.active.row_dimensions[25].height, 40)
        self.assertEqual(book.active['AF53'].value, long_text)
        self.assertIsNone(book.active['AB53'].value)
        book.close()
        pdf = PdfReader(BytesIO(build(row, 'QA Client', 'presentation')))
        self.assertIn('February 03, 2026', pdf.pages[0].extract_text())


if __name__ == "__main__":
    unittest.main()
