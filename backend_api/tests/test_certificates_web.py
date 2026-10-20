"""Certificates web isolation and compatibility; no production writes."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import hashlib
import sys
import unittest
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend_api"))
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from routers import certificates_web as web
from services.certificate_web_fields import TYPES, EXCEL_FIELDS
from services import certificate_web_export as export


class CertificatesTests(unittest.TestCase):
    def setUp(self):
        self.ctx = dict(company="MCI-CR", user="qa", permissions=dict(view=True,create=True,edit=True,approve=True))
        self.conn = MagicMock()
        self.cur = self.conn.cursor.return_value.__enter__.return_value
        self.row = dict(id=1,report_number="QA",report_no="QA",status="Pending for review",
                        updated_at=datetime(2026,10,20),remarks="Original",legacy_column="Keep me")

    def test_five_types_and_all_fields_unique(self):
        self.assertEqual(set(TYPES), {"weight","holds","sampling","sealing","lashing"})
        for spec in TYPES.values():
            fields = [f for section in spec["sections"] for f in section["fields"]]
            self.assertEqual(len(fields), len(spec["fields"]))
            self.assertIn(spec["number"], spec["fields"])

    def test_all_fields_validated_without_losing_leading_zeros(self):
        for spec in TYPES.values():
            values = {k: "20/10/2026" if f["type"]=="date" else "09:01" if f["type"]=="time" else "00123"
                      for k,f in spec["fields"].items()}
            result = web.clean_fields(spec, values)
            self.assertEqual(set(result), set(values))
            for k,f in spec["fields"].items():
                self.assertEqual(result[k], "2026-10-20" if f["type"]=="date" else "09:01" if f["type"]=="time" else "00123")

    def test_bad_dates_times_and_protected_fields_rejected(self):
        for kind, fields in [("weight",{"date":"31/02/2026"}),("holds",{"accepted_time":"26:01"}),
                             ("sampling",{"status":"Approved"}),("lashing",{"company_code":"MSL-CR"}),
                             ("sealing",{"remarks":"x\0y"}),("weight",{"quantity":12})]:
            with self.subTest(kind=kind,fields=fields), self.assertRaises(HTTPException) as err:
                web.clean_fields(TYPES[kind],fields)
            self.assertEqual(err.exception.status_code,422)

    def test_all_reads_and_locks_are_company_scoped(self):
        self.cur.fetchone.return_value = self.row
        for kind,spec in TYPES.items():
            web.read_record(self.cur,spec,1,self.ctx,True)
            sql, params = self.cur.execute.call_args.args
            self.assertIn("s.companies=1",sql)
            self.assertIn("FOR UPDATE OF c",sql)
            self.assertEqual(params,("MCI-CR",1))

    def test_missing_or_foreign_record_is_404(self):
        self.cur.fetchone.return_value = None
        for kind in TYPES:
            with self.assertRaises(HTTPException) as err:
                web.detail(kind,1,self.ctx,self.conn)
            self.assertEqual(err.exception.status_code,404)

    def test_all_creates_use_service_link_and_same_desktop_table(self):
        self.cur.fetchall.return_value = [{"company_code":"MCI-CR"}]
        self.cur.fetchone.return_value = self.row
        for kind,spec in TYPES.items():
            result = web.create(kind,web.SaveInput(fields={spec["number"]:"QA"},submit=True),self.ctx,self.conn)
            sql, params = self.cur.execute.call_args.args
            self.assertIn(f"INSERT INTO {spec['table']}",sql)
            self.assertIn("Pending for review",params)
            self.assertNotIn('"company_code"',sql)
            self.assertEqual(result["version"],web.version(self.row))
        self.assertEqual(self.conn.commit.call_count,5)

    def test_ambiguous_or_foreign_service_rolls_back(self):
        for rows in ([],[{"company_code":"MSL-CR"}],[{"company_code":"MCI-CR"}]*2):
            self.cur.fetchall.return_value = rows
            with self.assertRaises(HTTPException):
                web.create("weight",web.SaveInput(fields={"report_number":"QA"}),self.ctx,self.conn)
        self.conn.commit.assert_not_called()
        self.assertEqual(self.conn.rollback.call_count,3)

    def test_partial_patch_keeps_legacy_fields(self):
        self.cur.fetchone.side_effect = [self.row,{**self.row,"remarks":"New"}]
        result = web.update("weight",1,web.SaveInput(fields={"remarks":"New"},version=web.version(self.row)),self.ctx,self.conn)
        sql, params = self.cur.execute.call_args.args
        self.assertNotIn("legacy_column",sql)
        self.assertNotIn('"vessel"',sql)
        self.assertEqual(result["data"]["legacy_column"],"Keep me")
        self.assertEqual(params[0],"New")
        self.conn.commit.assert_called_once()

    def test_stale_approved_and_relink_edits_blocked(self):
        for row,body in [
            (self.row,web.SaveInput(fields={},version="stale")),
            ({**self.row,"status":"Approved"},None),
            (self.row,web.SaveInput(fields={"report_number":"OTHER"},version=web.version(self.row))),
        ]:
            self.cur.fetchone.return_value = row
            with self.assertRaises(HTTPException):
                web.update("weight",1,body or web.SaveInput(fields={},version=web.version(row)),self.ctx,self.conn)
        self.conn.commit.assert_not_called()

    def test_permission_denials_before_database(self):
        ctx={**self.ctx,"permissions":{}}
        for action in (
            lambda:web.create("weight",web.SaveInput(fields={}),ctx,self.conn),
            lambda:web.update("weight",1,web.SaveInput(fields={}),ctx,self.conn),
            lambda:web.change_status("weight",1,web.StatusInput(version="x",status="Approved"),ctx,self.conn),
        ):
            with self.assertRaises(HTTPException) as err: action()
            self.assertEqual(err.exception.status_code,403)
        self.conn.cursor.assert_not_called()

    def test_status_is_version_checked_and_independent_of_exports(self):
        self.cur.fetchone.side_effect=[self.row,{**self.row,"status":"Approved"}]
        result=web.change_status("weight",1,web.StatusInput(version=web.version(self.row),status="Approved"),self.ctx,self.conn)
        self.assertEqual(result["data"]["status"],"Approved")
        self.conn.commit.assert_called_once()

    def test_kiosk_cannot_access_certificates(self):
        app=FastAPI();app.include_router(web.router)
        app.dependency_overrides[web.get_db]=lambda:self.conn
        with TestClient(app) as client:
            self.assertEqual(client.get("/informes/certificates/bootstrap",headers={"X-Tally-Token":"x"*43}).status_code,401)

    def test_exports_leave_templates_unchanged_and_keep_every_field(self):
        from openpyxl import load_workbook
        from docx import Document
        templates=list(export.TEMPLATES.glob("*certificate*"))
        before={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in templates if p.is_file()}
        for kind,spec in TYPES.items():
            data={k:f"VALUE_{k}" for k in spec["fields"]}
            content=export.build(kind,data,spec["document"])
            if spec["document"]=="excel":
                book=load_workbook(BytesIO(content))
                self.assertEqual(book.active.page_setup.fitToWidth,1)
                self.assertEqual(book.active.page_setup.fitToHeight,0)
                self.assertEqual(book["data"].sheet_state,"hidden")
                for i,key in enumerate(EXCEL_FIELDS[kind],1):
                    if key in data:self.assertEqual(book["data"][f"B{i}"].value,data[key])
                self.assertFalse(any(c.data_type=="f" for row in book.active for c in row))
                book.close()
            else:
                doc=Document(BytesIO(content))
                text="\n".join(p.text for p in doc.paragraphs)
                self.assertNotIn("{",text)
                self.assertIn("VALUE_date",text)
                if kind=="lashing":self.assertNotIn("THIRTENN",text)
        self.assertEqual(before,{p:hashlib.sha256(p.read_bytes()).hexdigest() for p in before})

    def test_excel_specific_template_corrections_and_formula_injection(self):
        from openpyxl import load_workbook
        cases={
            "holds":({"cargo_holds":"5","vessel_holds":"1-5","product":"DDGS","place_location":"Caldera","place_date":"2026-10-20"},
                     {"D29":"5","A27":"DDGS","B33":"Caldera","B35":"2026-10-20"}),
            "sampling":({f"hold_{i}_seal":f"000{i}" for i in range(1,11)},{"C20":"0001","C28":"0005","H20":"0006","H28":"00010"}),
            "sealing":({"hold_2_fwd_escape":"001","hold_2_fwd_aft_hatch":"002","remarks":"=1+1"},{"C21":"001","C22":"002","A33":"=1+1"}),
        }
        for kind,(data,expected) in cases.items():
            book=load_workbook(BytesIO(export.build(kind,data,"excel")))
            for cell,value in expected.items():
                self.assertEqual(book.active[cell].value,value)
                self.assertEqual(book.active[cell].data_type,"s")
            book.close()

    def test_pdf_failure_never_modifies_record(self):
        self.cur.fetchone.return_value=self.row
        with patch.object(export,"build",side_effect=RuntimeError("converter unavailable")),self.assertRaises(HTTPException) as err:
            web.export("weight",1,"pdf",self.ctx,self.conn)
        self.assertEqual(err.exception.status_code,503)
        self.conn.commit.assert_not_called()
        self.assertFalse(any("UPDATE " in call.args[0] for call in self.cur.execute.call_args_list))

    def test_invalid_kind_never_enters_sql(self):
        with self.assertRaises(HTTPException) as err:
            web.specification("weight; DELETE FROM servicios")
        self.assertEqual(err.exception.status_code,404)


if __name__=="__main__":
    unittest.main()
