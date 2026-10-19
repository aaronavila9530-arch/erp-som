import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend_api'))
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient
from routers import logra_web as web


class LograTests(unittest.TestCase):
    def setUp(self):
        self.ctx = {'user':'qa', 'permissions':dict(view=True,create=True,edit=True)}
        self.conn = MagicMock()
        self.cur = self.conn.cursor.return_value.__enter__.return_value
        self.slug = web.CATALOG[0]['slug']
        self.answer = dict(form_slug=self.slug,section='critical_questions',item_key='1',bullets=['Original'],question_text='Original question',form_title='Original title')
        self.payload = {'report':dict(id=1,form_slug=self.slug,title='ONG - QA',version=2,created_by='desktop',agenda_items=[{'topic':'Legacy','custom':'keep'}],agenda_notes='Keep notes'),
                        'answers':[self.answer], 'attachments':[]}
        self.payload['revision'] = web.digest(self.payload)

    def test_catalog_exactly_matches_desktop(self):
        import runpy
        source = runpy.run_path(str(ROOT/'Modulos/Informes/logra_questionnaires_data.py'))['ONG_QUESTIONNAIRES']
        self.assertEqual(web.CATALOG,source)
        self.assertEqual(len(web.QUESTIONS),390)

    def test_no_schema_or_desktop_endpoint_calls(self):
        source = (ROOT/'backend_api/routers/logra_web.py').read_text()
        for term in ('ALTER TABLE', 'CREATE TABLE', '_ensure_schema(', 'save_logra_report('):
            self.assertNotIn(term,source)

    def test_denied_permissions(self):
        ctx = {'user':'reader','permissions':dict(view=True,create=False,edit=False)}
        for action in (lambda:web.create(web.NewReport(form_slug=self.slug),ctx,self.conn),
                       lambda:web.update(1,web.Save(revision='x'),ctx,self.conn)):
            with self.assertRaises(HTTPException) as error: action()
            self.assertEqual(error.exception.status_code,403)
        self.conn.cursor.assert_not_called()

    def test_access_requires_module_permission(self):
        with patch.object(web,'_has_visual_permission',return_value=False),self.assertRaises(HTTPException):
            web.access({'usuario':'qa','rol':'user'})

    def test_shared_catalog_not_tenant_migration(self):
        self.assertEqual(web.bootstrap(self.ctx)['scope'],'shared')
        self.assertNotIn('company_code',(ROOT/'backend_api/routers/logra_web.py').read_text())

    def test_token_changes_for_desktop_agenda_edits(self):
        changed = {**self.payload,'report':{**self.payload['report'],'agenda_notes':'Changed without version'}}
        self.assertNotEqual(web.digest(changed),self.payload['revision'])

    def test_stale_update_never_writes(self):
        with patch.object(web,'read',return_value=self.payload),self.assertRaises(HTTPException) as error:
            web.update(1,web.Save(revision='old'),self.ctx,self.conn)
        self.assertEqual(error.exception.status_code,409)
        self.conn.commit.assert_not_called()
        self.conn.rollback.assert_called_once()
        self.cur.execute.assert_not_called()

    def test_question_identity_and_empty_answer_preserved(self):
        body=web.Save(revision='x',answers=[web.Answer(form_slug=self.slug,section='critical_questions',item_key='1',bullets=[])])
        clean=web.clean_answers(body,self.payload)
        self.assertEqual(clean[0]['question_text'],'Original question')
        self.assertEqual(clean[0]['bullets'],[])
        self.assertEqual(clean[0]['form_title'],'Original title')

    def test_wrong_form_duplicate_and_unknown_question_rejected(self):
        variants=[{**self.answer,'form_slug':'other'},{**self.answer,'item_key':'unknown'}]
        for variant in variants:
            a={k:v for k,v in variant.items() if k in web.Answer.model_fields}
            with self.assertRaises(HTTPException):web.clean_answers(web.Save(revision='x',answers=[a]),self.payload)
        a={k:v for k,v in self.answer.items() if k in web.Answer.model_fields}
        with self.assertRaises(HTTPException):web.clean_answers(web.Save(revision='x',answers=[a,a]),self.payload)

    def test_partial_update_preserves_agenda_metadata_and_author(self):
        self.cur.fetchone.return_value={'version':3}
        body=web.Save(revision=self.payload['revision'],answers=[web.Answer(form_slug=self.slug,section='critical_questions',item_key='1',bullets=['New'])])
        with patch.object(web,'read',return_value=self.payload),patch.object(web,'_save_revision') as revision:
            web.update(1,body,self.ctx,self.conn)
            revision.assert_called_once_with(self.cur,1,3,self.slug,'qa')
        sql=' '.join(c.args[0] for c in self.cur.execute.call_args_list)
        self.assertNotIn('agenda_items=',sql)
        self.assertNotIn('created_by=',sql)
        self.assertNotIn('DELETE',sql)
        self.assertNotIn('status=',sql)
        self.conn.commit.assert_called_once()

    def test_failure_rolls_back(self):
        self.cur.execute.side_effect=RuntimeError('Database unavailable')
        with patch.object(web,'read',return_value=self.payload),self.assertRaises(RuntimeError):
            web.update(1,web.Save(revision=self.payload['revision']),self.ctx,self.conn)
        self.conn.rollback.assert_called_once()
        self.conn.commit.assert_not_called()

    def test_agenda_validation_and_unknown_legacy_fields(self):
        row={'date_iso':'2026-10-19','start_time':'09:01','end_time':'10:00','custom':'Keep'}
        self.assertEqual(web.clean_agenda([row]),[row])
        for changes in ({'date_iso':'2026-02-30'},{'start_time':'25:30'},{'end_time':'08:30'}):
            with self.assertRaises(HTTPException):web.clean_agenda([{**row,**changes}])

    def test_upload_unknown_question_does_not_write_file(self):
        with patch.object(web,'read',return_value=self.payload),self.assertRaises(HTTPException):
            web.upload(1,self.payload['revision'],'../../outside','critical_questions','1',UploadFile(io.BytesIO(b'abc'),filename='qa.pdf'),self.ctx,self.conn)
        self.conn.commit.assert_not_called()

    def test_upload_rolls_back_file_on_database_failure(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(web,'STORAGE_ROOT',Path(folder)),patch.object(web,'read',return_value=self.payload):
            self.cur.execute.side_effect=RuntimeError('Failed insert')
            with self.assertRaises(RuntimeError):
                web.upload(1,self.payload['revision'],self.slug,'critical_questions','1',UploadFile(io.BytesIO(b'abc'),filename='qa.pdf'),self.ctx,self.conn)
            self.assertFalse(list(Path(folder).rglob('*.pdf')))
        self.conn.rollback.assert_called_once()

    def test_download_prevents_path_escape(self):
        self.cur.fetchone.return_value={'stored_path':str(ROOT/'backend_api/main.py'),'original_filename':'main.py'}
        with self.assertRaises(HTTPException):web.download(1,self.ctx,self.conn)

    def test_ai_failure_preserves_original(self):
        with patch('ai.maritime_ai.improve_logra_questionnaire_text',side_effect=RuntimeError('Unavailable')),self.assertRaises(HTTPException) as error:
            web.improve(web.Improve(form_slug=self.slug,section='critical_questions',item_key='1',text='Original'),self.ctx)
        self.assertEqual(error.exception.status_code,503)

    def test_exports_escape_markup_and_excel_formulas(self):
        from services.logra_web_export import build
        from openpyxl import load_workbook
        from docx import Document
        from pypdf import PdfReader
        payload=json.loads(json.dumps(self.payload))
        payload['answers'][0]['bullets']=['<unsafe> & details', '=HYPERLINK("bad")']
        pdf=PdfReader(io.BytesIO(build(payload,'pdf')))
        self.assertIn('<unsafe> & details',''.join(p.extract_text() for p in pdf.pages))
        doc=Document(io.BytesIO(build(payload,'docx')))
        self.assertTrue(any('<unsafe>' in p.text for p in doc.paragraphs))
        book=load_workbook(io.BytesIO(build(payload,'xlsx')))
        self.assertNotEqual(book.active['D2'].data_type,'f')
        book.close()

    def test_http_authentication_and_strict_input(self):
        app=FastAPI();app.include_router(web.router)
        app.dependency_overrides[web.get_db]=lambda:self.conn
        with TestClient(app) as client:
            self.assertEqual(client.get('/informes/logra/bootstrap',headers={'X-Tally-Token':'x'*43}).status_code,401)
            app.dependency_overrides[web.access]=lambda:self.ctx
            response=client.patch('/informes/logra/1',json={'revision':'x','created_by':'other'})
            self.assertEqual(response.status_code,422)


if __name__ == '__main__':unittest.main()
