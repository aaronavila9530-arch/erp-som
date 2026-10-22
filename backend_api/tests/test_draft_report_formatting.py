from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn

from backend_api.services.draft_survey_report_formatting import (
    report_date, report_place_line, replace_placeholders, refresh_draft_difference,
    normalize_draft_layout,
)


class DraftFormattingTests(unittest.TestCase):
    def test_long_vessel_footer_has_separate_label_and_live_page_fields(self):
        root = Path(__file__).resolve().parents[1]
        doc = Document(root / 'templates/draft_word_template.docx')
        for p in doc.sections[0].footer.paragraphs:
            replace_placeholders(p, {
                '{draft_report_number}': '2238-0410-2026',
                '{word_vessel}': 'MV SANTA CAROLINA V2',
            })
        normalize_draft_layout(doc)
        p = next(p for p in doc.sections[0].footer.paragraphs if 'Page' in p.text)
        self.assertIn('2238-0410-2026 MV SANTA CAROLINA V2 DRAFT SURVEY\nPage', p.text)
        self.assertNotIn('MV MV', p.text)
        self.assertNotIn('   ', p.text)
        self.assertEqual(p.paragraph_format.right_indent, 0)
        self.assertEqual([n.get(qn('w:instr')) for n in p._p.iter(qn('w:fldSimple'))],
                         ['PAGE', 'NUMPAGES'])
        self.assertIn('\u00a0of\u00a0', p.text)
        website = next(p for p in doc.sections[0].footer.paragraphs
                       if 'www.mslogisticsgroup.com' in p.text)
        self.assertFalse(website._p.xpath('.//w:drawing | .//w:pict'))

    def test_cover_website_is_flow_text_and_background_fits_page(self):
        root = Path(__file__).resolve().parents[1]
        doc = Document(root / 'templates/presentation_draft_survey.docx')
        normalize_draft_layout(doc, presentation=True)
        website = next(p for p in doc.paragraphs if 'www.mslogisticsgroup.com' in p.text)
        self.assertFalse(website._p.xpath('.//w:drawing | .//w:pict'))
        self.assertEqual(website.paragraph_format.line_spacing, 1.0)
        section = doc.sections[0]
        anchor = section.header._element.xpath('.//wp:anchor')[0]
        extent = anchor.find(qn('wp:extent'))
        for axis, size, limit in [('H', 'cx', section.page_width),
                                  ('V', 'cy', section.page_height)]:
            pos = anchor.find(qn('wp:position' + axis))
            self.assertEqual(pos.get('relativeFrom'), 'page')
            offset = int(pos.find(qn('wp:posOffset')).text)
            self.assertGreaterEqual(offset, 0)
            self.assertLessEqual(offset + int(extent.get(size)), limit)
        normalize_draft_layout(doc, presentation=True)
        self.assertEqual(sum(p.text.count('www.mslogisticsgroup.com') for p in doc.paragraphs), 1)

    def test_old_saved_sign_is_recomputed_without_mutating_input(self):
        data={'word_draft_figures':'51444,60','word_bl_figures':'51442',
              'word_difference':'-2.60','word_percentage':'-0.0051'}
        result=refresh_draft_difference(data)
        self.assertEqual(result['word_difference'],'2.60')
        self.assertEqual(result['word_percentage'],'0.0051')
        self.assertEqual(data['word_difference'],'-2.60')

    def test_final_date_without_hour(self):
        data = {'word_final_draft': 'Oct 19 2026 11:40',
                'word_commenced': 'Oct 13 2026 22:35',
                'word_port': 'Caldera', 'word_country': 'Costa Rica'}
        self.assertEqual(report_date(data), date(2026, 10, 19))
        self.assertEqual(report_place_line(data), 'Caldera - Costa Rica, October 19, 2026')

    def test_split_date_and_completed_fallback(self):
        self.assertEqual(report_date({'word_final_draft_date': '19/10/2026'}), date(2026,10,19))
        self.assertEqual(report_date({'word_completed': datetime(2026,10,19,11)}),date(2026,10,19))
        self.assertIsNone(report_date({'word_commenced':'2026-10-13'}))

    def test_split_and_repeated_placeholders(self):
        doc=Document(); p=doc.add_paragraph()
        for text in ('Before {na','me} / {name}', ' after {unknown}'):
            p.add_run(text)
        replace_placeholders(p,{'{name}':'MV LAURA'})
        self.assertEqual(p.text,'Before MV LAURA / MV LAURA after ')

    def test_real_footer_keeps_page_fields_untouched(self):
        root=Path(__file__).resolve().parents[1]
        doc=Document(root/'templates/draft_word_template.docx')
        footer=doc.sections[0].footer
        paragraphs=[p for p in footer.paragraphs if 'Page ' in p.text]
        self.assertEqual(len(paragraphs),1)
        p=paragraphs[0]
        before=[r._r.xml for r in p.runs if r._r.xpath('.//w:fldChar | .//w:instrText')]
        replace_placeholders(p,{'{draft_report_number}':'2246-1310-2026','{word_vessel}':'MV LAURA'})
        self.assertEqual(before,[r._r.xml for r in p.runs if r._r.xpath('.//w:fldChar | .//w:instrText')])
        self.assertRegex(p.text,r'Page 3 of 4\s*$')
        self.assertEqual([n.text.strip() for n in p._p.iter(qn('w:instrText'))],['PAGE','NUMPAGES'])

    def test_real_word_template_date_and_times(self):
        from backend_api.services.draft_survey_word_pdf_service import generate_draft_survey_word_pdf
        class StopConversion(Exception): pass
        with tempfile.TemporaryDirectory() as directory:
            payload={'draft_report_number':'QA-DATE', 'word_final_draft':'2026-10-19 11:40',
                     'word_initial_draft':'2026-10-13 20:00','word_port':'Caldera','word_country':'Costa Rica'}
            with patch('tempfile.gettempdir',return_value=directory),patch(
                'backend_api.services.draft_survey_word_pdf_service.subprocess.run',side_effect=StopConversion):
                with self.assertRaises(StopConversion):generate_draft_survey_word_pdf(payload)
            doc=Document(Path(directory)/'QA-DATE.docx')
            text='\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
            self.assertIn('Caldera - Costa Rica, October 19, 2026',text)
            self.assertNotIn('January 14, 2026',text)
            self.assertIn('Oct 13 2026 20:00',text)
            self.assertIn('Oct 19 2026 11:40',text)


if __name__=='__main__':unittest.main()
