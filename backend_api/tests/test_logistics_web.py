from datetime import date
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import MagicMock
from docx import Document
from fastapi import HTTPException
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend_api'))
from routers import certificates_web as web
from services.logistics_web_fields import SPEC
from services.logistics_web_export import word
from services.mci_survey_export import build_word
from services.mci_survey_schema import SurveyContent, SurveySave
from services.logistics_import import content_from_legacy

class LogisticsTests(unittest.TestCase):
    def test_msl_only_and_service_scope(self):
        self.assertEqual(web.specification('logistics',{'company':'MSL-CR'}),SPEC)
        with self.assertRaises(HTTPException):web.specification('logistics',{'company':'MCI-CR'})
        self.assertIn('s.companies=1',web.source(SPEC))

    def test_exact_desktop_fields(self):
        expected='cert_no customer port country report_date vessel_name flag_port_registry grt nrt imo_no build_year captain chief_officer arrival_date inspection_date supervision_completed_date process_text conclusion_text findings_documental_text findings_operational_text incidents_text'.split()
        self.assertEqual(set(SPEC['fields']),set(expected))
        self.assertEqual(web.clean_fields(SPEC,{'report_date':'29/10/2026'})['report_date'],'2026-10-29')
        with self.assertRaises(HTTPException):web.clean_fields(SPEC,{'report_date':'31/02/2026'})

    def test_create_duplicate_does_not_write(self):
        conn=MagicMock();cur=conn.cursor.return_value.__enter__.return_value
        cur.fetchall.return_value=[{'company_code':'MSL-CR'}]
        cur.fetchone.return_value={'id':7}
        with self.assertRaises(HTTPException) as err:
            web.create('logistics',web.SaveInput(fields={'cert_no':'QA'}),{'company':'MSL-CR','permissions':{'create':True}},conn)
        self.assertEqual(err.exception.status_code,409)
        conn.commit.assert_not_called()

    def test_service_identity_and_chronology(self):
        cur=MagicMock();cur.fetchone.return_value=dict(cliente='Client',buque_contenedor='Vessel',puerto='Port',pais='Country')
        values=dict(cert_no='QA',customer='Spoof',inspection_date='2026-10-29',supervision_completed_date='2026-10-28')
        with self.assertRaises(HTTPException):web.logistics_identity(cur,values,{'company':'MSL-CR'})
        values['supervision_completed_date']='2026-10-30'
        web.logistics_identity(cur,values,{'company':'MSL-CR'})
        self.assertEqual(values['customer'],'Client')
        self.assertEqual(values['company_code'],'MSL-CR')

    def test_templates_untouched_and_all_fields_render(self):
        templates=ROOT/'backend_api/templates'
        names=['vessel_truck_supervision_aligned_dates.docx','presentation_truck_supervision.docx']
        before=[sha256((templates/n).read_bytes()).hexdigest() for n in names]
        data={k:'UNIQUE_'+k for k in SPEC['fields']}
        for k,f in SPEC['fields'].items():
            if f['type']=='date':data[k]=date(2026,10,29)
        data['company_code']='MSL-CR'
        with tempfile.TemporaryDirectory() as folder:
            for cover in (False,True):
                path=Path(folder)/'report.docx';word(data,path,cover)
                doc=Document(path)
                text='\n'.join(p.text for p in doc.paragraphs)
                self.assertIn('UNIQUE_vessel_name',text)
                self.assertNotIn('{process_text}',text)
                self.assertNotIn('{findings_operational_text}',text)
                if not cover:self.assertIn('UNIQUE_findings_operational_text',text)
                if not cover:
                    conclusion=next(p for p in doc.paragraphs if p.text.strip().startswith('CONCLUS'))
                    self.assertTrue(conclusion.paragraph_format.keep_with_next)
                self.assertTrue(doc.inline_shapes or doc.part.rels)
            data['company_code']='MCI-CR'
            with self.assertRaises(ValueError):word(data,path)
        self.assertEqual(before,[sha256((templates/n).read_bytes()).hexdigest() for n in names])

    def test_mci_hierarchy_and_specifics(self):
        content={'hierarchy_version':2,'grt':'34567','arrival_date':'2026-10-29','sections':{
            'instructions':{'enabled':True}, 'chronology':{'enabled':True},
            'particulars':{'enabled':True}, 'process':{'enabled':True,'blocks':[{'text':'Truck control process'}]},
            'findings':{'enabled':False},'documentary':{'enabled':True,'blocks':[{'text':'HIDDEN DOCUMENTARY'}]},
            'incidents':{'enabled':True,'blocks':[{'text':'HIDDEN INCIDENT'}]}}}
        row=dict(company_code='MCI-CR',kind='logistics',vessel='QA Vessel',requested_by='QA Client',report_number='QA-123',content=content)
        doc=Document(BytesIO(build_word(row)));text='\n'.join(p.text for p in doc.paragraphs)
        for expected in ('Logistics Supervision','Truck control process','GRT: 34567','October 29, 2026'):
            self.assertIn(expected,text)
        self.assertNotIn('HIDDEN',text)
        self.assertNotIn('maiz',text)
        SurveySave(kind='logistics',service_id=1,content=content)
        with self.assertRaises(ValueError):SurveyContent(inspection_date='2026-10-29',supervision_completed_date='2026-10-28')

    def test_import_retains_narrative_without_assuming_cargo_or_signatory(self):
        original={'process_text':'Exact original\nSecond paragraph','findings_operational_text':'Checks', 'status':'Approved'}
        content=content_from_legacy(original)
        self.assertEqual(content.sections['process'].blocks[0].text,original['process_text'])
        self.assertEqual(content.report_status,'Draft')
        self.assertEqual(content.cargo,'')
        self.assertEqual(content.surveyor,'')
        self.assertEqual(original['status'],'Approved')

    def test_permission_and_approval_conflict(self):
        conn=MagicMock()
        with self.assertRaises(HTTPException) as err:
            web.change_status('logistics',1,web.StatusInput(version='x',status='Approved'),
                              {'company':'MSL-CR','permissions':{'view':True}},conn)
        self.assertEqual(err.exception.status_code,403)
        conn.cursor.assert_not_called()
        with self.assertRaises(HTTPException):
            web.logistics_assist(web.GrainAssistInput(text='QA'),{'permissions':{'view':True}})

if __name__=='__main__':unittest.main()
