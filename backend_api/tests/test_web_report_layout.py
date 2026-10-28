from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook, load_workbook
from services.web_docx_layout import fit_web_docx
from services.web_excel_layout import continue_long_cells


class WebReportLayoutTests(unittest.TestCase):
    def test_word_preserves_artwork_margins_fonts_and_heading_rules(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = 'Original header'
        doc.sections[0].footer.paragraphs[0].text = 'Original footer'
        header = doc.sections[0].header._element.xml
        footer = doc.sections[0].footer._element.xml
        margins = doc.sections[0]._sectPr.xml
        heading = doc.add_paragraph('SURVEY FINDINGS')
        heading.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Observation ' * 200
        run = table.cell(0, 0).paragraphs[0].runs[0]
        run.font.size = Pt(11)
        table.rows[0].height = Inches(.2)
        table.rows[0]._tr.trPr.find(qn('w:trHeight')).set(qn('w:hRule'), 'exact')
        table.rows[0]._tr.trPr.append(OxmlElement('w:cantSplit'))
        fit_web_docx(doc)
        self.assertEqual(run.font.size.pt, 11)
        self.assertTrue(heading.paragraph_format.keep_with_next)
        self.assertFalse(table.autofit)
        self.assertEqual(table.rows[0]._tr.trPr.find(qn('w:trHeight')).get(qn('w:hRule')), 'atLeast')
        self.assertIsNone(table.rows[0]._tr.trPr.find(qn('w:cantSplit')))
        self.assertEqual(header, doc.sections[0].header._element.xml)
        self.assertEqual(footer, doc.sections[0].footer._element.xml)
        self.assertEqual(margins, doc.sections[0]._sectPr.xml)

    def test_wide_table_is_bounded_and_idempotent(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for column in table.columns:
            column.width = Inches(6)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged'
        fit_web_docx(doc)
        width = sum(column.width for column in table.columns)
        section = doc.sections[0]
        self.assertLessEqual(width, section.page_width-section.left_margin-section.right_margin)
        xml = doc.element.xml
        fit_web_docx(doc)
        self.assertEqual(xml, doc.element.xml)

    def test_inline_image_fits_page_without_distortion(self):
        doc = Document()
        picture = doc.add_paragraph().add_run().add_picture(
            str(Path(__file__).resolve().parents[1] / 'assets/mci_logo.png'), width=Inches(12))
        ratio = picture.width / picture.height
        fit_web_docx(doc)
        section = doc.sections[0]
        self.assertLessEqual(picture.width, section.page_width-section.left_margin-section.right_margin)
        self.assertLessEqual(picture.height, section.page_height-section.top_margin-section.bottom_margin)
        self.assertAlmostEqual(picture.width/picture.height, ratio, places=5)

    def test_desktop_grain_keeps_legacy_layout_default(self):
        from services.grain_sampling_doc_service import generate_grain_sampling_doc
        with tempfile.TemporaryDirectory() as folder:
            with patch('services.grain_sampling_doc_service.apply_docx_autofit') as legacy:
                generate_grain_sampling_doc({}, output_path=str(Path(folder)/'desktop.docx'))
                legacy.assert_called_once()
            with patch('services.grain_sampling_doc_service.apply_docx_autofit') as legacy:
                generate_grain_sampling_doc({}, output_path=str(Path(folder)/'web.docx'), web_layout=True)
                legacy.assert_not_called()

    def test_excel_continuation_is_lossless_and_printable(self):
        book = Workbook()
        sheet = book.active
        sheet.column_dimensions['A'].width = 30
        value = '=NOT_A_FORMULA ' + 'Long observation with damage details.\n' * 200
        sheet['A1'] = value
        sheet['A1'].data_type = 's'
        extra = continue_long_cells(book, sheet, {'A1': 'Remarks'}, 'QA')
        self.assertIsNotNone(extra)
        chunks = [extra.cell(row, 1).value for row in range(4, extra.max_row+1)
                  if extra.cell(row, 1).value is not None]
        self.assertEqual(''.join(chunks), value)
        self.assertTrue(all(d.height <= 240 for r,d in extra.row_dimensions.items() if r >= 4))
        self.assertEqual(extra.page_setup.fitToWidth, 1)
        self.assertEqual(extra.page_setup.fitToHeight, 0)
        self.assertTrue(all(c.data_type != 'f' for row in extra for c in row))

    def test_normal_excel_has_no_continuation(self):
        book = Workbook()
        book.active['A1'] = 'Normal text'
        self.assertIsNone(continue_long_cells(book, book.active, {'A1':'Remarks'}))
        self.assertEqual(len(book.worksheets), 1)

    def test_long_container_cover_and_report(self):
        from services.container_web_export import build
        from pypdf import PdfReader
        value = 'Long narrative with inspection details. ' * 250
        data = dict(company_code='MCI-CR', report_no='QA', remarks=value, vessel='QA Vessel')
        book = load_workbook(BytesIO(build(data, 'Client', 'excel')))
        self.assertIn('Continuation', book.sheetnames)
        book.close()
        content = build(data, value, 'presentation')
        pdf = PdfReader(BytesIO(content))
        self.assertGreater(len(pdf.pages), 1)
        self.assertIn('inspection details', ''.join(p.extract_text() for p in pdf.pages))

    def test_long_certificate_continuation_visible(self):
        from services.certificate_web_export import build
        book = load_workbook(BytesIO(build('holds', {'remarks':'Inspection. '*1000}, 'excel')))
        self.assertEqual(book['Continuation'].sheet_state, 'visible')
        book.close()

    def test_vessel_footer_uses_fields_without_cached_literal_numbers(self):
        from services.certificate_web_export import build
        for kind in ('condition', 'cargo_condition'):
            with self.subTest(kind=kind):
                doc = Document(BytesIO(build(kind, {'company_code':'MSL-CR',
                    'report_number':'QA-123', 'vessel':'QA VESSEL', 'narrative_1':'Details. '*200}, 'word')))
                text = '\n'.join(p.text for p in doc.paragraphs)
                self.assertNotIn('{narrative_', text)
                self.assertNotIn('{service_start_date}', text)
                footer = doc.sections[0].footer
                self.assertNotIn('Page 1 of 5', '\n'.join(p.text for p in footer.paragraphs))
                fields = footer._element.xpath('.//w:fldSimple/@w:instr')
                self.assertIn('PAGE', fields)
                self.assertIn('NUMPAGES', fields)
                self.assertFalse(doc.element.body.xpath('.//wp:anchor'))


if __name__ == '__main__':
    unittest.main()
