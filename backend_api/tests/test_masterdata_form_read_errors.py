import sys
import unittest
import os
import tempfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend_api"))
from Modulos.MasterData import masterdata_forms as desktop
from routers import masterdata_forms as web


class FormReadTests(unittest.TestCase):
    def test_bilingual_client_headers_are_scoped(self):
        pairs = {"Nombresicea juridico / Legal name": "NombreJuridico",
                 "Correo electronico / Email": "Correo", "Telefono / Phone": "Telefono",
                 "Provincia / Province": "Provincia", "Prefijo telefonico / Phone prefix": "Prefijo"}
        for header, expected in pairs.items():
            self.assertEqual(web._field_from_header(header, web.UPLOAD_SPECS["cliente"]), expected)
            self.assertEqual(desktop._field_from_header(header, desktop.FORM_SPECS["cliente"]), expected)
        self.assertEqual(web._field_from_header("Correo electronico / Email", web.UPLOAD_SPECS["surveyor"]), "email")
        self.assertEqual(web._field_from_header("Unknown", web.UPLOAD_SPECS["cliente"]), "Unknown")

    def test_supplier_word_retains_contact_and_bank_fields(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "supplier.docx"
            doc = Document()
            doc.add_paragraph("Formulario Master Data - Proveedor")
            table = doc.add_table(rows=1, cols=2)
            values = {"Nombre / First name": "Demo", "Pais / Country": "Costa Rica",
                      "Correo electronico / Email": "demo@example.com", "Banco / Bank": "Demo Bank",
                      "Direccion exacta / Full address USA": "Example address"}
            for header, value in values.items():
                cells = table.add_row().cells
                cells[0].text, cells[1].text = header, value
            doc.save(path)
            for read in (web._read_docx_upload, desktop._read_docx):
                data = read(path)[0]["data"]
                self.assertEqual(data["Correo"], "demo@example.com")
                self.assertEqual(data["Banco"], "Demo Bank")
                self.assertEqual(data["DireccionExacta"], "Example address")

    @unittest.skipUnless(os.name == "nt", "Windows sharing semantics")
    def test_read_while_sync_handle_allows_reads(self):
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        create = kernel.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                           ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "synced.xlsx"
            path.write_bytes(b"saved content")
            handle = create(str(path), 0x10000 | 0x40000000, 7, None, 3, 0, None)
            self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
            try:
                with self.assertRaises(PermissionError):
                    path.read_bytes()
                self.assertEqual(desktop._read_shared_file(path), b"saved content")
            finally:
                kernel.CloseHandle(handle)

    def test_word_snapshot_import(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "surveyor.docx"
            doc = Document()
            doc.add_paragraph("Surveyor")
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Campo"
            table.rows[0].cells[1].text = "Valor"
            for field, value in [("nombre", "Demo"), ("apellidos", "Prueba"), ("nacionalidad", "Peru")]:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = field, value
            doc.save(path)
            result = desktop.import_masterdata_files([str(path)])
            self.assertEqual(result[0]["data"]["nombre"], "Demo")
            self.assertFalse(result[0]["error"])
            path.unlink()

    def test_snapshot_matches_file_and_releases_handle(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "form.xlsx"
            path.write_bytes(b"saved form content")
            with desktop._form_snapshot(path) as snapshot:
                self.assertEqual(snapshot.read(), b"saved form content")
            path.rename(path.with_suffix(".moved"))

    def test_transient_sharing_violation_retries(self):
        error = OSError("sharing violation")
        error.winerror = 32
        with patch.object(desktop, "_read_shared_file", side_effect=[error, b"ok"]) as reader, patch.object(desktop.time, "sleep"):
            with desktop._form_snapshot(Path("cloud.xlsx")) as snapshot:
                self.assertEqual(snapshot.read(), b"ok")
        self.assertEqual(reader.call_count, 2)

    def test_permission_denied_is_not_retried(self):
        with patch.object(desktop, "_read_shared_file", side_effect=PermissionError()) as reader:
            with self.assertRaises(PermissionError):
                desktop._form_snapshot(Path("denied.xlsx"))
        self.assertEqual(reader.call_count, 1)

    def test_excel_snapshot_import(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "surveyor.xlsx"
            wb = Workbook()
            wb.active.title = "Surveyor"
            wb.active.append(["nombre", "apellidos", "nacionalidad"])
            wb.active.append(["Demo", "Prueba", "Peru"])
            wb.save(path)
            wb.close()
            result = desktop.import_masterdata_files([str(path)])
            self.assertEqual(result[0]["data"]["nombre"], "Demo")
            self.assertFalse(result[0]["error"])
            path.unlink()

    def test_desktop_permission_failure_does_not_stop_batch(self):
        good = {"file": "good.xlsx", "entity": "surveyor", "data": {"nombre": "Demo"}, "error": ""}
        with patch.object(desktop, "_read_xlsx", side_effect=[PermissionError(13, "Permission denied"), [good]]):
            rows = desktop.import_masterdata_files(["blocked.xlsx", "good.xlsx"])
        self.assertIn("OneDrive", rows[0]["error"])
        self.assertEqual(rows[0]["file"], "blocked.xlsx")
        self.assertEqual(rows[1], good)

    def test_web_corrupt_file_does_not_stop_batch(self):
        good = {"file": "good.xlsx", "entity": "surveyor", "data": {}, "error": ""}
        with patch.object(web, "_read_xlsx_upload", side_effect=[ValueError("bad zip"), [good]]):
            rows = web._import_masterdata_files([Path("bad.xlsx"), Path("good.xlsx")])
        self.assertTrue(rows[0]["error"])
        self.assertEqual(rows[1], good)


if __name__ == "__main__":
    unittest.main()
