"""Project-specific tally sheets. Each hold has its own optimistic revision."""
from datetime import date as Date, time as Time
from io import BytesIO
from pathlib import Path
import re
from uuid import UUID

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field, field_validator
from psycopg2.extras import Json, RealDictCursor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from database import get_db
from security.auth import get_current_user
from routers.mobile_auth import _has_visual_permission
from services.tenanting import company_code

router = APIRouter(tags=["Tally Control"])
ASSETS = Path(__file__).resolve().parents[1] / "assets" / "tally"
FIELDS = ["number", "hold", "date", "entry", "exit", "spc", "company", "ticket", "guide", "seal", "plate", "driver", "stowage", "notes"]
HEADERS = ["No.", "Bodega", "Fecha", "Entrada", "Salida", "SPC", "Empresa", "Ficha", "Guia Surco", "Guia Sello", "Placa", "Chofer", "Consecutivo Estiba", "Observacion"]


def context(x_company_code: str = Header(default="MSL-CR"), user=Depends(get_current_user)):
    if not _has_visual_permission(user["usuario"], user["rol"], "informes", "view"):
        raise HTTPException(403, "Sin permiso para Informes")
    editable = any(_has_visual_permission(user["usuario"], user["rol"], "informes", a) for a in ("edit", "create"))
    return {"company": company_code(header_value=x_company_code), "user": user["usuario"], "editable": editable}


def writer(ctx=Depends(context)):
    if not ctx["editable"]:
        raise HTTPException(403, "Sin permiso para editar Informes")
    return ctx


def schema(conn):
    with conn.cursor() as cur:
        cur.execute("""CREATE TABLE IF NOT EXISTS tally_projects (
            id BIGSERIAL PRIMARY KEY, company_code TEXT NOT NULL, name TEXT NOT NULL,
            service_id BIGINT, holds JSONB NOT NULL, revision INTEGER NOT NULL DEFAULT 0,
            created_by TEXT NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
        cur.execute("""CREATE TABLE IF NOT EXISTS tally_sheets (
            project_id BIGINT NOT NULL REFERENCES tally_projects(id), hold INTEGER NOT NULL,
            rows JSONB NOT NULL DEFAULT '[]', revision INTEGER NOT NULL DEFAULT 0,
            updated_by TEXT, updated_at TIMESTAMPTZ, PRIMARY KEY(project_id,hold))""")
        cur.execute("""CREATE TABLE IF NOT EXISTS tally_drivers (
            company_code TEXT NOT NULL, plate TEXT NOT NULL, driver TEXT NOT NULL,
            PRIMARY KEY(company_code,plate))""")
        cur.execute("""CREATE TABLE IF NOT EXISTS tally_companies (
            company_code TEXT NOT NULL, name TEXT NOT NULL, PRIMARY KEY(company_code,name))""")
        cur.execute("""CREATE TABLE IF NOT EXISTS tally_audit (
            id BIGSERIAL PRIMARY KEY, project_id BIGINT NOT NULL REFERENCES tally_projects(id),
            hold INTEGER, action TEXT NOT NULL, revision INTEGER, performed_by TEXT NOT NULL,
            detail JSONB NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    conn.commit()


def project(cur, pid, ctx, lock=False):
    cur.execute("SELECT * FROM tally_projects WHERE id=%s AND company_code=%s" + (" FOR UPDATE" if lock else ""), (pid, ctx["company"]))
    row = cur.fetchone()
    if not row:
        raise HTTPException(404, "Proyecto no encontrado")
    return row


def audit(cur, pid, ctx, action, detail, hold=None, revision=None):
    cur.execute("INSERT INTO tally_audit(project_id,hold,action,revision,performed_by,detail) VALUES(%s,%s,%s,%s,%s,%s)",
                (pid, hold, action, revision, ctx["user"], Json(detail)))


class ProjectInput(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    service_id: int | None = None
    holds: list[int] = Field(min_length=1, max_length=50)
    revision: int = Field(default=0, ge=0)

    @field_validator("name")
    @classmethod
    def clean_name(cls, v):
        if not v.strip():
            raise ValueError("Nombre requerido")
        return v.strip()

    @field_validator("holds")
    @classmethod
    def valid_holds(cls, v):
        if any(n < 1 or n > 999 for n in v) or len(set(v)) != len(v):
            raise ValueError("Bodegas unicas entre 1 y 999")
        return v


class TallyRow(BaseModel):
    model_config = ConfigDict(extra="forbid", str_max_length=1000)
    id: UUID
    number: str = ""
    date: Date | None = None
    entry: Time | None = None
    exit: Time | None = None
    spc: str = ""
    company: str = ""
    ticket: str = ""
    guide: str = ""
    seal: str = ""
    plate: str = ""
    stowage: str = ""
    notes: str = ""

    @field_validator("date", "entry", "exit", mode="before")
    @classmethod
    def empty_temporal(cls, v):
        return None if v == "" else v

    @field_validator("seal")
    @classmethod
    def valid_seal(cls, v):
        normalized = v.strip().lower()
        if normalized not in ("", "si", "sí", "no"):
            raise ValueError("Guia Sello debe ser Si o No")
        return "Si" if normalized in ("si", "sí") else ("No" if normalized else "")


class SheetInput(BaseModel):
    revision: int = Field(ge=0)
    rows: list[TallyRow] = Field(max_length=5000)

    @field_validator("rows")
    @classmethod
    def unique_rows(cls, rows):
        if len({r.id for r in rows}) != len(rows):
            raise ValueError("Identificadores de filas repetidos")
        return rows


class DriverInput(BaseModel):
    plate: str = Field(min_length=1, max_length=30)
    driver: str = Field(min_length=1, max_length=150)


class CompanyInput(BaseModel):
    name: str = Field(min_length=1, max_length=100)


def plate_key(value):
    return re.sub(r"[\s-]+", "", str(value or "")).upper()


@router.get("/som/tally-assets/{name}", include_in_schema=False)
def asset(name: str):
    if name not in {"tally.js", "tally.css", "tabulator.min.js", "tabulator.min.css", "papaparse.min.js"}:
        raise HTTPException(404)
    return FileResponse(ASSETS / name, headers={"Cache-Control": "no-store" if name in {"tally.js", "tally.css"} else "public,max-age=86400"})


@router.get("/tally/bootstrap")
def bootstrap(q: str = Query(default="", max_length=100), ctx=Depends(context), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT id,name,service_id,holds,revision FROM tally_projects WHERE company_code=%s ORDER BY id DESC", (ctx["company"],))
        projects = cur.fetchall()
        cur.execute("SELECT plate,driver FROM tally_drivers WHERE company_code=%s ORDER BY plate", (ctx["company"],))
        drivers = cur.fetchall()
        cur.execute("SELECT name FROM tally_companies WHERE company_code=%s ORDER BY name", (ctx["company"],))
        companies = [r["name"] for r in cur.fetchall()]
        cur.execute("""SELECT consec,buque_contenedor,cliente FROM servicios WHERE company_code=%s
            AND (buque_contenedor ILIKE %s OR consec::text ILIKE %s) ORDER BY consec DESC LIMIT 100""",
                    (ctx["company"], "%" + q + "%", "%" + q + "%"))
        services = cur.fetchall()
    return dict(projects=projects, drivers=drivers, companies=companies, services=services, editable=ctx["editable"])


def validate_service(cur, sid, ctx):
    if sid is not None:
        cur.execute("SELECT consec FROM servicios WHERE consec=%s AND company_code=%s", (sid, ctx["company"]))
        if not cur.fetchone():
            raise HTTPException(422, "Servicio no pertenece a esta empresa")


@router.post("/tally/projects")
def create_project(body: ProjectInput, ctx=Depends(writer), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        validate_service(cur, body.service_id, ctx)
        cur.execute("""INSERT INTO tally_projects(company_code,name,service_id,holds,created_by)
            VALUES(%s,%s,%s,%s,%s) RETURNING *""", (ctx["company"], body.name, body.service_id, Json(body.holds), ctx["user"]))
        row = cur.fetchone()
        audit(cur, row["id"], ctx, "PROJECT_CREATED", body.model_dump())
    conn.commit()
    return row


@router.put("/tally/projects/{pid}")
def update_project(pid: int, body: ProjectInput, ctx=Depends(writer), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        old = project(cur, pid, ctx, True)
        if old["revision"] != body.revision:
            raise HTTPException(409, "El proyecto cambio. Recargue antes de editar.")
        removed = list(set(old["holds"]) - set(body.holds))
        if removed:
            cur.execute("SELECT hold FROM tally_sheets WHERE project_id=%s AND hold=ANY(%s) AND jsonb_array_length(rows)>0", (pid, removed))
            if cur.fetchone():
                raise HTTPException(409, "No se puede quitar una bodega con registros")
        validate_service(cur, body.service_id, ctx)
        cur.execute("UPDATE tally_projects SET name=%s,service_id=%s,holds=%s,revision=revision+1 WHERE id=%s RETURNING *", (body.name, body.service_id, Json(body.holds), pid))
        row = cur.fetchone()
        audit(cur, pid, ctx, "PROJECT_UPDATED", body.model_dump())
    conn.commit()
    return row


@router.get("/tally/projects/{pid}/sheets/{hold}")
def get_sheet(pid: int, hold: int, ctx=Depends(context), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        p = project(cur, pid, ctx)
        if hold not in p["holds"]:
            raise HTTPException(404, "Bodega no configurada")
        cur.execute("SELECT * FROM tally_sheets WHERE project_id=%s AND hold=%s", (pid, hold))
        return cur.fetchone() or dict(rows=[], revision=0, hold=hold)


@router.put("/tally/projects/{pid}/sheets/{hold}")
def save_sheet(pid: int, hold: int, body: SheetInput, ctx=Depends(writer), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        p = project(cur, pid, ctx, True)
        if hold not in p["holds"]:
            raise HTTPException(404, "Bodega no configurada")
        cur.execute("INSERT INTO tally_sheets(project_id,hold) VALUES(%s,%s) ON CONFLICT DO NOTHING", (pid, hold))
        cur.execute("SELECT revision FROM tally_sheets WHERE project_id=%s AND hold=%s FOR UPDATE", (pid, hold))
        if cur.fetchone()["revision"] != body.revision:
            raise HTTPException(409, "Otra persona modifico esta bodega. Sus cambios no se sobrescribieron. Exporte su borrador y recargue.")
        cur.execute("SELECT plate,driver FROM tally_drivers WHERE company_code=%s", (ctx["company"],))
        drivers = {r["plate"]: r["driver"] for r in cur.fetchall()}
        cur.execute("SELECT name FROM tally_companies WHERE company_code=%s", (ctx["company"],))
        companies = {r["name"] for r in cur.fetchall()}
        rows = []
        for item in body.rows:
            row = item.model_dump(mode="json")
            row["plate"] = plate_key(row["plate"])
            if row["company"] and row["company"] not in companies:
                raise HTTPException(422, "Empresa no registrada en catalogo: " + row["company"])
            row["driver"] = drivers.get(row["plate"], "")
            rows.append(row)
        revision = body.revision + 1
        cur.execute("UPDATE tally_sheets SET rows=%s,revision=%s,updated_by=%s,updated_at=NOW() WHERE project_id=%s AND hold=%s", (Json(rows), revision, ctx["user"], pid, hold))
        audit(cur, pid, ctx, "SHEET_SAVED", {"rows": rows}, hold, revision)
    conn.commit()
    return {"revision": revision, "count": len(rows)}


@router.put("/tally/drivers")
def put_driver(body: DriverInput, ctx=Depends(writer), conn=Depends(get_db)):
    schema(conn)
    key = plate_key(body.plate)
    if not key or not body.driver.strip():
        raise HTTPException(422, "Placa y chofer requeridos")
    with conn.cursor() as cur:
        cur.execute("INSERT INTO tally_drivers VALUES(%s,%s,%s) ON CONFLICT(company_code,plate) DO UPDATE SET driver=EXCLUDED.driver", (ctx["company"], key, body.driver.strip()))
    conn.commit()
    return {"plate": key, "driver": body.driver.strip()}


@router.post("/tally/companies")
def put_company(body: CompanyInput, ctx=Depends(writer), conn=Depends(get_db)):
    schema(conn)
    if not body.name.strip():
        raise HTTPException(422, "Empresa requerida")
    with conn.cursor() as cur:
        cur.execute("INSERT INTO tally_companies VALUES(%s,%s) ON CONFLICT DO NOTHING", (ctx["company"], body.name.strip()))
    conn.commit()
    return {"name": body.name.strip()}


@router.post("/tally/catalog/import")
def import_catalog(file: UploadFile = File(...), ctx=Depends(writer), conn=Depends(get_db)):
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Archivo mayor a 5 MB")
    import zipfile
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError()
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if "Choferes" not in wb.sheetnames:
            raise ValueError()
        records, companies = {}, set()
        for values in wb["Choferes"].iter_rows(min_row=2, max_row=10001, max_col=3, values_only=True):
            plate, driver, company = values
            if plate is not None and driver:
                key = plate_key(plate)
                name = str(driver).strip()
                if len(key) > 30 or len(name) > 150:
                    raise ValueError()
                if key in records and records[key] != name:
                    raise ValueError()
                records[key] = name
            if company and str(company).strip():
                companies.add(str(company).strip()[:100])
        wb.close()
        if not records:
            raise ValueError()
    except Exception:
        raise HTTPException(422, "Excel invalido. Se requiere hoja Choferes: PLACA, CHOFER, EMPRESA; sin placas ambiguas.")
    schema(conn)
    with conn.cursor() as cur:
        for plate, driver in records.items():
            cur.execute("INSERT INTO tally_drivers VALUES(%s,%s,%s) ON CONFLICT DO NOTHING", (ctx["company"], plate, driver))
        for name in companies:
            cur.execute("INSERT INTO tally_companies VALUES(%s,%s) ON CONFLICT DO NOTHING", (ctx["company"], name))
    conn.commit()
    return {"processed": len(records), "companies": len(companies)}


@router.get("/tally/projects/{pid}/history")
def history(pid: int, ctx=Depends(context), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        project(cur, pid, ctx)
        cur.execute("SELECT hold,action,revision,performed_by,created_at FROM tally_audit WHERE project_id=%s ORDER BY id DESC LIMIT 50", (pid,))
        return cur.fetchall()


def export_workbook(p, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for hold in p["holds"]:
        ws = wb.create_sheet(f"Bodega {hold}")
        ws.append(HEADERS)
        for data in sheets.get(hold, []):
            values = [hold if field == "hold" else data.get(field) for field in FIELDS]
            ws.append(values)
            for index, field in enumerate(FIELDS, 1):
                cell = ws.cell(ws.max_row, index)
                value = data.get(field)
                if field == "date" and value:
                    cell.value = Date.fromisoformat(value)
                    cell.number_format = "dd/mm/yyyy"
                elif field in ("entry", "exit") and value:
                    cell.value = Time.fromisoformat(value)
                    cell.number_format = "hh:mm"
                elif isinstance(cell.value, str):
                    cell.data_type = "s"  # Never execute user-entered Excel formulas.
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="165B72")
            ws.column_dimensions[cell.column_letter].width = 20 if cell.column < 12 else 32
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@router.get("/tally/projects/{pid}/excel")
def excel(pid: int, ctx=Depends(context), conn=Depends(get_db)):
    schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        p = project(cur, pid, ctx)
        cur.execute("SELECT hold,rows FROM tally_sheets WHERE project_id=%s", (pid,))
        sheets = {r["hold"]: r["rows"] for r in cur.fetchall()}
    return StreamingResponse(export_workbook(p, sheets), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="Tally-{pid}.xlsx"'})
