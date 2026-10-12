"""Web assets and exports shared by the commercial workspace."""
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Literal
from xml.sax.saxutils import escape
import csv

from fastapi import APIRouter, Depends, Header, HTTPException
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
from psycopg2.extras import RealDictCursor
from database import get_db
from rbac_service import has_permission
from routers.comercial import require_permission
from services.tenanting import company_code
from services.commercial_access import BOARD_ONLY_USERS, require_commercial_detail

router = APIRouter(prefix="/comercial/web", tags=["Comercial Web"])
ASSETS = Path(__file__).resolve().parents[1] / "assets" / "comercial"


@router.get("/assets/{name}")
def asset(name: str):
    if name not in {"comercial.js", "comercial.css"}:
        raise HTTPException(404, "Archivo no encontrado")
    return FileResponse(ASSETS / name, media_type="text/javascript" if name.endswith(".js") else "text/css")


@router.get("/meta", dependencies=[Depends(require_permission("comercial", "view"))])
def metadata(x_user_role: str = Header(...), x_user: str = Header(""),
             x_company_code: str = Header("MSL-CR"), conn=Depends(get_db)):
    actions = ("board", "clients", "ports", "services", "quotations", "edit", "delete", "export")
    permissions = {action: has_permission(x_user_role, "comercial", action if action in {"edit", "delete", "export"} else "view") for action in actions}
    permissions["board"] = True
    if x_user.lower() in BOARD_ONLY_USERS:
        permissions = {key: key == "board" for key in permissions}
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("""SELECT DISTINCT cliente, continente, pais, puerto, surveyor, operacion,
                       EXTRACT(YEAR FROM fecha_inicio)::int AS year
                       FROM servicios
                       WHERE COALESCE(NULLIF(TRIM(company_code), ''), 'MSL-CR') = %s""",
                    (company_code(x_company_code),))
        rows = cur.fetchall()
    fields = ("cliente", "continente", "pais", "puerto", "surveyor", "operacion", "year")
    return {"permissions": permissions, "options": {key: sorted({r[key] for r in rows if r.get(key) is not None}) for key in fields}}


class TableExport(BaseModel):
    title: str = Field(default="Comercial", max_length=150)
    columns: list[str] = Field(max_length=50)
    labels: list[str] = Field(max_length=50)
    rows: list[dict[str, Any]] = Field(max_length=100000)


def safe_cell(value):
    value = "" if value is None else value
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@router.post("/export/{fmt}", dependencies=[Depends(require_permission("comercial", "export")), Depends(require_commercial_detail)])
def export_table(fmt: Literal["xlsx", "csv", "pdf"], payload: TableExport):
    if not payload.columns or len(payload.columns) != len(payload.labels):
        raise HTTPException(422, "Columnas invalidas")
    values = [[safe_cell(row.get(key)) for key in payload.columns] for row in payload.rows]
    output = BytesIO()
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = "Comercial"
        sheet.append(payload.labels)
        for row in values:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="006D77")
        for i, label in enumerate(payload.labels, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(40, max(16, len(label) + 3))
        book.save(output)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(payload.labels)
        writer.writerows(values)
        output.write(stream.getvalue().encode("utf-8-sig"))
        mime = "text/csv; charset=utf-8"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Spacer
        style = ParagraphStyle("cell", fontSize=7, leading=9, wordWrap="CJK")
        def para(value):
            return Paragraph(escape(str(value)), style)
        table = LongTable([[para(v) for v in payload.labels]] + [[para(v) for v in row] for row in values],
                          colWidths=[(landscape(A3)[0] - 48) / len(payload.columns)] * len(payload.columns), repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7f1f2")),
                                  ("GRID", (0, 0), (-1, -1), .25, colors.lightgrey),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        SimpleDocTemplate(output, pagesize=landscape(A3), leftMargin=24, rightMargin=24).build([
            Paragraph(escape(payload.title), getSampleStyleSheet()["Heading2"]), Spacer(1, 12), table])
        mime = "application/pdf"
    return Response(output.getvalue(), media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="comercial.{fmt}"'})
