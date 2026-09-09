from __future__ import annotations

import io
from dataclasses import dataclass

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse


router = APIRouter(prefix="/master-data/forms", tags=["Master Data - Forms"])


@dataclass(frozen=True)
class FormSpec:
    key: str
    label: str
    fields: tuple[str, ...]


SPECS = {
    "clientes": FormSpec("clientes", "Cliente", ("NombreJuridico", "NombreComercial", "Pais", "Correo", "Telefono", "CedulaJuridicaVAT", "Comentarios", "Provincia", "Canton", "Distrito", "DireccionExacta", "FechaDePago", "Prefijo", "ContactoPrincipal", "ContactoSecundario")),
    "proveedores": FormSpec("proveedores", "Proveedor", ("Nombre", "Apellidos", "NombreComercial", "Cedula", "Pais", "Provincia", "Canton", "Distrito", "DireccionExacta", "Prefijo", "Telefono", "Correo", "TerminosPago", "Banco", "CuentaIBAN", "SwiftCode", "UID", "DireccionBanco", "TipoProveeduria", "Comentarios", "Activo")),
    "empleados": FormSpec("empleados", "Empleado", ("nombre", "apellidos", "estado_civil", "genero", "nacionalidad", "prefijo", "telefono", "provincia", "canton", "distrito", "direccion", "jornada", "salario", "fecha_ingreso", "horas_contratadas", "horas_tope_ordinario", "horas_tope_maximo", "tarifa_hora_extra", "pago_minimo_garantizado", "pago", "banco", "cuenta_iban", "moneda", "enfermedades", "contacto_emergencia", "telefono_emergencia", "activo1", "marca1", "serial1", "activo2", "marca2", "serial2", "activo3", "marca3", "serial3", "activo")),
    "surveyores": FormSpec("surveyores", "Surveyor", ("nombre", "apellidos", "email", "estado_civil", "genero", "nacionalidad", "prefijo", "telefono", "provincia", "canton", "distrito", "direccion", "jornada", "operacion", "honorario", "pago", "banco", "direccion_banco", "cuenta_iban", "moneda", "swift", "uid", "enfermedades", "contacto_emergencia", "telefono_emergencia", "puerto", "activo")),
    "servicios-md": FormSpec("servicios-md", "Servicio", ("codigo_prod", "nombre", "costo")),
}


LABELS = {
    "nombre": "Nombre / First name",
    "apellidos": "Apellidos / Last name",
    "Nombre": "Nombre / First name",
    "Apellidos": "Apellidos / Last name",
    "NombreJuridico": "Nombre juridico / Legal name",
    "NombreComercial": "Nombre comercial / Trade name",
    "Pais": "Pais / Country",
    "Correo": "Correo electronico / Email",
    "Telefono": "Telefono / Phone",
    "CedulaJuridicaVAT": "Cedula juridica o VAT / Tax ID or VAT",
    "Comentarios": "Comentarios / Notes",
    "Provincia": "Provincia / Province",
    "Canton": "Canton / County",
    "Distrito": "Distrito / District",
    "DireccionExacta": "Direccion exacta / Full address",
    "FechaDePago": "Fecha de pago / Payment date",
    "Prefijo": "Prefijo telefonico / Phone prefix",
    "ContactoPrincipal": "Contacto principal / Main contact",
    "ContactoSecundario": "Contacto secundario / Secondary contact",
    "Cedula": "Cedula / ID",
    "TerminosPago": "Terminos de pago / Payment terms",
    "Banco": "Banco / Bank",
    "CuentaIBAN": "Cuenta IBAN / IBAN account",
    "SwiftCode": "Swift Code / Swift code",
    "UID": "UID",
    "DireccionBanco": "Direccion banco / Bank address",
    "TipoProveeduria": "Tipo de proveeduria / Supplier type",
    "Activo": "Activo / Active",
    "email": "Correo electronico / Email",
    "estado_civil": "Estado civil / Marital status",
    "genero": "Genero / Gender",
    "nacionalidad": "Nacionalidad / Nationality",
    "prefijo": "Prefijo telefonico / Phone prefix",
    "telefono": "Telefono / Phone",
    "provincia": "Provincia / Province",
    "canton": "Canton / County",
    "distrito": "Distrito / District",
    "direccion": "Direccion / Address",
    "jornada": "Jornada / Work schedule",
    "salario": "Salario / Salary",
    "fecha_ingreso": "Fecha ingreso / Start date",
    "horas_contratadas": "Horas pactadas / Contracted hours",
    "horas_tope_ordinario": "Primer aviso / Regular hours cap",
    "horas_tope_maximo": "Segundo aviso / Maximum hours cap",
    "tarifa_hora_extra": "Tarifa hora extra / Overtime rate",
    "pago_minimo_garantizado": "Pago minimo garantizado / Guaranteed minimum pay",
    "pago": "Forma de pago / Payment method",
    "banco": "Banco / Bank",
    "cuenta_iban": "Cuenta IBAN / IBAN account",
    "moneda": "Moneda / Currency",
    "enfermedades": "Enfermedades / Medical conditions",
    "contacto_emergencia": "Contacto emergencia / Emergency contact",
    "telefono_emergencia": "Telefono emergencia / Emergency phone",
    "activo1": "Activo asignado 1 / Assigned asset 1",
    "marca1": "Marca 1 / Brand 1",
    "serial1": "Serial 1",
    "activo2": "Activo asignado 2 / Assigned asset 2",
    "marca2": "Marca 2 / Brand 2",
    "serial2": "Serial 2",
    "activo3": "Activo asignado 3 / Assigned asset 3",
    "marca3": "Marca 3 / Brand 3",
    "serial3": "Serial 3",
    "activo": "Activo / Active",
    "operacion": "Operacion / Operation",
    "honorario": "Honorario / Fee",
    "direccion_banco": "Direccion banco / Bank address",
    "swift": "Swift Code / Swift code",
    "uid": "UID",
    "puerto": "Puerto / Port",
    "codigo_prod": "Codigo producto / Product code",
    "costo": "Costo / Cost",
}


def _spec(entity_key: str) -> FormSpec:
    spec = SPECS.get(entity_key)
    if not spec:
        raise HTTPException(status_code=400, detail="Tipo de formulario no soportado")
    return spec


def _filename(spec: FormSpec, suffix: str) -> str:
    return f"Formulario_MasterData_{spec.label}.{suffix}"


def _excel(spec: FormSpec) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = spec.label
    ws["A1"] = "Formulario Master Data / Master Data Form"
    ws["A2"] = spec.label
    ws.append([])
    ws.append(list(spec.fields))
    ws.append([LABELS.get(field, field) for field in spec.fields])
    for row in (4, 5):
        for cell in ws[row]:
            cell.font = Font(bold=row == 4)
            if row == 4:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column_cells in ws.columns:
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 38)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _word(spec: FormSpec) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Formulario Master Data / Master Data Form", level=1)
    doc.add_paragraph(spec.label)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Campo / Field"
    header[1].text = "Valor / Value"
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
    for field in spec.fields:
        cells = table.add_row().cells
        cells[0].text = LABELS.get(field, field)
        cells[1].text = ""
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


@router.get("/{entity_key}/{fmt}")
def download_masterdata_form(entity_key: str, fmt: str):
    spec = _spec(entity_key)
    fmt = fmt.lower().strip()
    if fmt in {"excel", "xlsx"}:
        content = _excel(spec)
        suffix = "xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt in {"word", "docx"}:
        content = _word(spec)
        suffix = "docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Use word o excel.")
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{_filename(spec, suffix)}"'},
    )
