from __future__ import annotations

import hashlib
import io
import os
import socket
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from psycopg2.extras import Json, RealDictCursor

from database import get_db
from routers.accounting import _ensure_accounting_professional_schema
from routers.accounting_tax import _ensure_schema as _ensure_tax_schema, tax_iva, obligations
from services.finance_audit import actor_from_headers, audit_event, ensure_finance_audit_schema
from services.tenanting import company_code as normalize_company_code


router = APIRouter(prefix="/accounting/advanced", tags=["Accounting Advanced Controls"])
MONEY = Decimal("0.01")


def _money(value) -> Decimal:
    try:
        return Decimal(str(value or 0)).quantize(MONEY)
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0.00")


def _to_float(value):
    return float(_money(value))


def _valid_period(period: str):
    if not period or len(period) != 7 or period[4] != "-":
        raise HTTPException(400, "period must use YYYY-MM")
    year = int(period[:4])
    month = int(period[5:7])
    if month < 1 or month > 12:
        raise HTTPException(400, "Invalid period month")
    return year, month


def _period_bounds(period: str):
    year, month = _valid_period(period)
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return start, end


def _ensure_schema(conn):
    _ensure_accounting_professional_schema(conn)
    _ensure_tax_schema(conn)
    with conn.cursor() as cur:
        ensure_finance_audit_schema(cur)
        for ddl in (
            "ALTER TABLE finance_audit_log ADD COLUMN IF NOT EXISTS ip_address TEXT",
            "ALTER TABLE finance_audit_log ADD COLUMN IF NOT EXISTS workstation TEXT",
            "ALTER TABLE accounting_entries ADD COLUMN IF NOT EXISTS prepared_by TEXT",
            "ALTER TABLE accounting_entries ADD COLUMN IF NOT EXISTS reviewed_by TEXT",
            "ALTER TABLE accounting_entries ADD COLUMN IF NOT EXISTS reviewed_at TIMESTAMP",
            "ALTER TABLE accounting_entries ADD COLUMN IF NOT EXISTS approval_required BOOLEAN NOT NULL DEFAULT TRUE",
            "ALTER TABLE accounting_entries ADD COLUMN IF NOT EXISTS approval_limit_amount NUMERIC(18,2)",
        ):
            cur.execute(ddl)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS accounting_approval_policies (
                id BIGSERIAL PRIMARY KEY,
                origin TEXT NOT NULL DEFAULT 'ALL',
                entry_type TEXT NOT NULL DEFAULT 'ALL',
                amount_from NUMERIC(18,2) NOT NULL DEFAULT 0,
                amount_to NUMERIC(18,2),
                required_reviewer_role TEXT,
                required_approver_role TEXT,
                active BOOLEAN NOT NULL DEFAULT TRUE,
                created_by TEXT,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMP NOT NULL DEFAULT NOW()
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS accounting_document_supports (
                id BIGSERIAL PRIMARY KEY,
                entry_id INTEGER REFERENCES accounting_entries(id) ON DELETE CASCADE,
                module TEXT NOT NULL DEFAULT 'accounting',
                entity_type TEXT,
                entity_id TEXT,
                document_type TEXT NOT NULL,
                filename TEXT NOT NULL,
                mime_type TEXT,
                size_bytes BIGINT,
                sha256 TEXT,
                stored_path TEXT NOT NULL,
                description TEXT,
                uploaded_by TEXT,
                uploaded_at TIMESTAMP NOT NULL DEFAULT NOW(),
                metadata JSONB NOT NULL DEFAULT '{}'::jsonb
            )
        """)
        cur.execute("CREATE INDEX IF NOT EXISTS idx_accounting_support_entry ON accounting_document_supports(entry_id, uploaded_at DESC)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_accounting_support_search ON accounting_document_supports(document_type, entity_type, entity_id)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS accounting_budgets (
                id BIGSERIAL PRIMARY KEY,
                company_code VARCHAR(30) NOT NULL DEFAULT 'MSL-CR',
                period VARCHAR(7) NOT NULL,
                account_code TEXT NOT NULL,
                cost_center_code TEXT NOT NULL DEFAULT '',
                currency_code VARCHAR(3) NOT NULL DEFAULT 'CRC',
                budget_amount NUMERIC(18,2) NOT NULL DEFAULT 0,
                purpose TEXT NOT NULL DEFAULT 'BUDGET',
                name TEXT,
                target_date DATE,
                monthly_contribution NUMERIC(18,2) NOT NULL DEFAULT 0,
                target_amount NUMERIC(18,2) NOT NULL DEFAULT 0,
                current_amount NUMERIC(18,2) NOT NULL DEFAULT 0,
                funding_bank_account_code TEXT,
                funding_bank_account_name TEXT,
                status TEXT NOT NULL DEFAULT 'ACTIVE',
                notes TEXT,
                created_by TEXT,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMP NOT NULL DEFAULT NOW(),
                UNIQUE(period, account_code, cost_center_code, currency_code)
            )
        """)
        for ddl in (
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS company_code VARCHAR(30) NOT NULL DEFAULT 'MSL-CR'",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS purpose TEXT NOT NULL DEFAULT 'BUDGET'",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS name TEXT",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS target_date DATE",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS monthly_contribution NUMERIC(18,2) NOT NULL DEFAULT 0",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS target_amount NUMERIC(18,2) NOT NULL DEFAULT 0",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS current_amount NUMERIC(18,2) NOT NULL DEFAULT 0",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS funding_bank_account_code TEXT",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS funding_bank_account_name TEXT",
            "ALTER TABLE accounting_budgets ADD COLUMN IF NOT EXISTS status TEXT NOT NULL DEFAULT 'ACTIVE'",
        ):
            cur.execute(ddl)
        cur.execute("ALTER TABLE accounting_budgets DROP CONSTRAINT IF EXISTS accounting_budgets_period_account_code_cost_center_code_currency_code_key")
        cur.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS ux_accounting_budgets_company_period_account
            ON accounting_budgets(company_code, period, account_code, cost_center_code, currency_code)
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS accounting_budget_contributions (
                id BIGSERIAL PRIMARY KEY,
                budget_id BIGINT REFERENCES accounting_budgets(id) ON DELETE CASCADE,
                company_code VARCHAR(30) NOT NULL DEFAULT 'MSL-CR',
                contribution_date DATE NOT NULL,
                amount NUMERIC(18,2) NOT NULL,
                currency_code VARCHAR(3) NOT NULL DEFAULT 'CRC',
                bank_account_code TEXT NOT NULL,
                bank_account_name TEXT,
                reference TEXT NOT NULL,
                notes TEXT,
                accounting_entry_id INTEGER REFERENCES accounting_entries(id),
                created_by TEXT,
                created_at TIMESTAMP NOT NULL DEFAULT NOW()
            )
        """)
        cur.execute("CREATE INDEX IF NOT EXISTS idx_accounting_budget_contrib_budget ON accounting_budget_contributions(budget_id, contribution_date)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS accounting_fx_revaluations (
                id BIGSERIAL PRIMARY KEY,
                period VARCHAR(7) NOT NULL,
                as_of_date DATE NOT NULL,
                currency_code VARCHAR(3) NOT NULL DEFAULT 'USD',
                exchange_rate NUMERIC(18,6) NOT NULL,
                total_open_currency NUMERIC(18,2) NOT NULL DEFAULT 0,
                total_crc_value NUMERIC(18,2) NOT NULL DEFAULT 0,
                previous_crc_value NUMERIC(18,2) NOT NULL DEFAULT 0,
                difference_crc NUMERIC(18,2) NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'PREVIEW',
                accounting_entry_id INTEGER REFERENCES accounting_entries(id),
                created_by TEXT,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                metadata JSONB NOT NULL DEFAULT '{}'::jsonb
            )
        """)
        cur.execute("""
            INSERT INTO accounting_approval_policies(
                origin, entry_type, amount_from, amount_to,
                required_reviewer_role, required_approver_role, created_by
            ) VALUES
              ('ALL','STANDARD',0,1000000,'finance','gerencia','SYSTEM_DEFAULT'),
              ('ALL','HIGH_VALUE',1000000,NULL,'finance','gerencia','SYSTEM_DEFAULT')
            ON CONFLICT DO NOTHING
        """)
    conn.commit()


def _latest_rate(cur, as_of: date):
    cur.execute("""
        SELECT rate, rate_date, source
        FROM exchange_rate
        WHERE rate_date <= %s
        ORDER BY rate_date DESC
        LIMIT 1
    """, (as_of,))
    row = cur.fetchone()
    if not row:
        raise HTTPException(409, "No historical exchange rate available for this date")
    return row


def _entry_total(cur, entry_id: int) -> Decimal:
    cur.execute("""
        SELECT GREATEST(COALESCE(SUM(debit),0), COALESCE(SUM(credit),0)) AS amount
        FROM accounting_lines
        WHERE entry_id=%s
    """, (entry_id,))
    return _money((cur.fetchone() or {}).get("amount"))


def _serialize(row):
    output = {}
    for key, value in dict(row or {}).items():
        if isinstance(value, Decimal):
            output[key] = _to_float(value)
        elif hasattr(value, "isoformat"):
            output[key] = value.isoformat()
        else:
            output[key] = value
    return output


def _company_from_header(x_company_code: str | None = None) -> str:
    return normalize_company_code(header_value=x_company_code)


def _budget_filters(
    company: str,
    period: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    purpose: str = "ALL",
    status: str = "ALL",
):
    filters = ["company_code = %s"]
    params: list = [company]
    if period:
        _valid_period(period)
        filters.append("period = %s")
        params.append(period)
    if date_from:
        filters.append("COALESCE(target_date, (period || '-01')::date) >= %s")
        params.append(date_from)
    if date_to:
        filters.append("COALESCE(target_date, (period || '-01')::date) <= %s")
        params.append(date_to)
    purpose = str(purpose or "ALL").strip().upper()
    if purpose != "ALL":
        filters.append("UPPER(COALESCE(purpose, 'BUDGET')) = %s")
        params.append(purpose)
    status = str(status or "ALL").strip().upper()
    if status != "ALL":
        filters.append("UPPER(COALESCE(status, 'ACTIVE')) = %s")
        params.append(status)
    return filters, params


def _get_postable_account(cur, code: str, label: str):
    code = str(code or "").strip()
    if not code:
        raise HTTPException(400, f"{label} is required")
    cur.execute(
        """
        SELECT account_code, account_name
        FROM accounting_accounts
        WHERE account_code=%s
          AND COALESCE(active, TRUE)=TRUE
          AND COALESCE(accepts_posting, FALSE)=TRUE
        LIMIT 1
        """,
        (code,),
    )
    account = cur.fetchone()
    if not account:
        raise HTTPException(400, f"{label} does not exist or is not postable: {code}")
    return account


@router.get("/fx/rate")
def historical_fx_rate(rate_date: date | None = None, conn=Depends(get_db)):
    _ensure_schema(conn)
    rate_date = rate_date or date.today()
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        row = _latest_rate(cur, rate_date)
    return {"rate": _to_float(row["rate"]), "date": row["rate_date"].isoformat(), "source": row.get("source")}


@router.get("/fx/revaluation-preview")
def fx_revaluation_preview(period: str, currency_code: str = "USD", conn=Depends(get_db)):
    _ensure_schema(conn)
    start, end = _period_bounds(period)
    as_of = min(end - timedelta(days=1), date.today())
    currency_code = (currency_code or "USD").upper()
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        rate = _latest_rate(cur, as_of)
        try:
            from routers.accounting_auxiliaries import _ensure_schema as _ensure_aux_schema, sync_auxiliaries
            _ensure_aux_schema(conn)
            sync_auxiliaries(conn)
        except Exception:
            conn.rollback()
        cur.execute("""
            SELECT e.entity_type, e.entity_code, e.entity_name,
                   COALESCE(e.control_account_code, s.control_account_code) AS control_account_code,
                   d.document_number, d.issue_date, d.due_date, d.currency_code,
                   d.open_amount, d.metadata
            FROM accounting_auxiliary_documents d
            JOIN accounting_auxiliary_entities e ON e.id=d.entity_id
            LEFT JOIN accounting_auxiliary_settings s ON s.entity_type=e.entity_type
            WHERE d.status='OPEN'
              AND UPPER(d.currency_code)=%s
              AND COALESCE(d.open_amount,0)<>0
            ORDER BY e.entity_type, e.entity_name, d.document_number
        """, (currency_code,))
        rows = []
        total_open = Decimal("0")
        total_crc = Decimal("0")
        previous_crc = Decimal("0")
        for row in cur.fetchall():
            open_amount = _money(row["open_amount"])
            current_crc = (open_amount * _money(rate["rate"])).quantize(MONEY)
            previous = _money((row.get("metadata") or {}).get("last_revalued_crc"))
            total_open += open_amount
            total_crc += current_crc
            previous_crc += previous
            rows.append(_serialize(row) | {
                "current_crc_value": _to_float(current_crc),
                "previous_crc_value": _to_float(previous),
                "difference_crc": _to_float(current_crc - previous),
            })
    return {
        "period": period,
        "as_of": as_of.isoformat(),
        "currency_code": currency_code,
        "exchange_rate": _to_float(rate["rate"]),
        "rate_date": rate["rate_date"].isoformat(),
        "rows": rows,
        "totals": {
            "open_currency": _to_float(total_open),
            "current_crc_value": _to_float(total_crc),
            "previous_crc_value": _to_float(previous_crc),
            "difference_crc": _to_float(total_crc - previous_crc),
        },
    }


@router.post("/fx/revaluation-post")
def post_fx_revaluation(
    payload: dict,
    conn=Depends(get_db),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    _ensure_schema(conn)
    period = payload.get("period")
    currency_code = (payload.get("currency_code") or "USD").upper()
    user, role = actor_from_headers(x_user, x_role, x_user_role)
    user = payload.get("user") or user
    preview = fx_revaluation_preview(period=period, currency_code=currency_code, conn=conn)
    diff = _money(preview["totals"]["difference_crc"])
    if diff == 0:
        return {"status": "ok", "message": "No FX difference to post", "preview": preview}
    gain_code = payload.get("gain_account_code") or "4.9.01"
    loss_code = payload.get("loss_account_code") or "5.9.01"
    suspense_code = payload.get("offset_account_code") or "2.9.99"
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        for code, name, typ, bal in (
            (gain_code, "Ganancia por diferencia cambiaria", "REVENUE", "CREDIT"),
            (loss_code, "Perdida por diferencia cambiaria", "EXPENSE", "DEBIT"),
            (suspense_code, "Revaluacion saldos abiertos USD", "LIABILITY", "CREDIT"),
        ):
            cur.execute("""
                INSERT INTO accounting_accounts(account_code,account_name,account_type,normal_balance,account_level,accepts_posting,active,created_by,updated_by)
                VALUES(%s,%s,%s,%s,1,TRUE,TRUE,%s,%s)
                ON CONFLICT(account_code) DO NOTHING
            """, (code, name, typ, bal, user, user))
        entry_date = date.fromisoformat(preview["as_of"])
        cur.execute("""
            INSERT INTO accounting_entries(entry_date,period,description,origin,created_by,workflow_status,company_code,currency_code,exchange_rate,posting_rule_code,posting_metadata,posted_by,posted_at)
            VALUES(%s,%s,%s,'FX_REVALUATION',%s,'POSTED','MSL-CR','CRC',1,'FX_MONTHLY_REVALUATION',%s,%s,NOW())
            RETURNING id
        """, (entry_date, period, f"Revaluacion mensual {currency_code} {period}", user, Json(preview), user))
        entry_id = cur.fetchone()["id"]
        amount = abs(diff)
        if diff > 0:
            lines = [(suspense_code, "Revaluacion saldos abiertos USD", amount, 0), (gain_code, "Ganancia por diferencia cambiaria", 0, amount)]
        else:
            lines = [(loss_code, "Perdida por diferencia cambiaria", amount, 0), (suspense_code, "Revaluacion saldos abiertos USD", 0, amount)]
        for code, name, debit, credit in lines:
            cur.execute("""
                INSERT INTO accounting_lines(entry_id,account_code,account_name,debit,credit,line_description)
                VALUES(%s,%s,%s,%s,%s,%s)
            """, (entry_id, code, name, debit, credit, f"FX revaluation {currency_code} {period}"))
        cur.execute("""
            INSERT INTO accounting_fx_revaluations(period,as_of_date,currency_code,exchange_rate,total_open_currency,total_crc_value,previous_crc_value,difference_crc,status,accounting_entry_id,created_by,metadata)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,'POSTED',%s,%s,%s)
            RETURNING *
        """, (
            period, entry_date, currency_code, preview["exchange_rate"],
            preview["totals"]["open_currency"], preview["totals"]["current_crc_value"],
            preview["totals"]["previous_crc_value"], preview["totals"]["difference_crc"],
            entry_id, user, Json(preview),
        ))
        row = cur.fetchone()
        audit_event(cur, "accounting", "FX_REVALUATION_POSTED", "accounting_entry", entry_id, user, role, payload.get("reason"), None, _serialize(row), {"period": period})
    conn.commit()
    return {"status": "ok", "entry_id": entry_id, "revaluation": row}


@router.get("/tax/deep-summary")
def tax_deep_summary(period: str, conn=Depends(get_db)):
    _ensure_schema(conn)
    start, end = _period_bounds(period)
    iva = tax_iva(period=period, conn=conn)
    calendar = obligations(year=int(period[:4]), period=period, pending_only=True, conn=conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("""
            SELECT direction, hacienda_status, COUNT(*) documents,
                   COALESCE(SUM(subtotal),0) subtotal,
                   COALESCE(SUM(tax_amount),0) tax,
                   COALESCE(SUM(total),0) total
            FROM tax_electronic_documents
            WHERE issue_datetime >= %s AND issue_datetime < %s
            GROUP BY direction, hacienda_status
            ORDER BY direction, hacienda_status
        """, (start, end))
        by_status = cur.fetchall()
        cur.execute("""
            SELECT COUNT(*) FILTER(WHERE LOWER(COALESCE(account_name,'')) LIKE '%%retenc%%' OR account_code LIKE '2.1.03%%') ret_line_count,
                   COALESCE(SUM(CASE WHEN LOWER(COALESCE(account_name,'')) LIKE '%%retenc%%' OR account_code LIKE '2.1.03%%' THEN credit-debit ELSE 0 END),0) retention_balance
            FROM accounting_lines l JOIN accounting_entries e ON e.id=l.entry_id
            WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s
        """, (start, end))
        retentions = cur.fetchone()
    return {
        "period": period,
        "iva": iva,
        "retentions": _serialize(retentions),
        "documents_by_status": [_serialize(row) for row in by_status],
        "fiscal_calendar": calendar,
        "controls": {
            "documental_vs_accounting_ready": bool(iva.get("ready_to_file")),
            "requires_review": not bool(iva.get("ready_to_file")),
        },
    }


@router.get("/approvals/policies")
def approval_policies(conn=Depends(get_db)):
    _ensure_schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM accounting_approval_policies WHERE active=TRUE ORDER BY amount_from, origin")
        return {"data": cur.fetchall()}


@router.get("/approvals/evaluate/{entry_id}")
def evaluate_entry_approval(entry_id: int, conn=Depends(get_db)):
    _ensure_schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM accounting_entries WHERE id=%s", (entry_id,))
        entry = cur.fetchone()
        if not entry:
            raise HTTPException(404, "Accounting entry not found")
        amount = _entry_total(cur, entry_id)
        cur.execute("""
            SELECT * FROM accounting_approval_policies
            WHERE active=TRUE
              AND (origin='ALL' OR origin=%s)
              AND amount_from <= %s
              AND (amount_to IS NULL OR amount_to >= %s)
            ORDER BY amount_from DESC
            LIMIT 1
        """, (entry.get("origin") or "ALL", amount, amount))
        policy = cur.fetchone()
    status = entry.get("workflow_status")
    return {
        "entry": _serialize(entry),
        "amount": _to_float(amount),
        "policy": _serialize(policy),
        "flow": {
            "prepared": bool(entry.get("created_by") or entry.get("prepared_by")),
            "reviewed": bool(entry.get("reviewed_by") or status in {"APPROVED", "POSTED"}),
            "approved": bool(entry.get("approved_by") or status == "POSTED"),
            "posted": status == "POSTED",
        },
        "requires_approval": bool(policy),
    }


@router.post("/documents/upload")
async def upload_accounting_support(
    entry_id: int | None = None,
    module: str = "accounting",
    entity_type: str | None = None,
    entity_id: str | None = None,
    document_type: str = "SUPPORT",
    description: str | None = None,
    file: UploadFile = File(...),
    conn=Depends(get_db),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    _ensure_schema(conn)
    user, role = actor_from_headers(x_user, x_role, x_user_role)
    content = await file.read()
    if not content:
        raise HTTPException(400, "Empty file")
    digest = hashlib.sha256(content).hexdigest()
    safe_name = "".join(ch for ch in (file.filename or "support") if ch.isalnum() or ch in "._- ").strip() or "support"
    folder = Path("storage/accounting/supports") / datetime.now().strftime("%Y/%m")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{digest[:12]}_{safe_name}"
    path.write_bytes(content)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("""
            INSERT INTO accounting_document_supports(entry_id,module,entity_type,entity_id,document_type,filename,mime_type,size_bytes,sha256,stored_path,description,uploaded_by)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            RETURNING *
        """, (entry_id, module, entity_type, entity_id, document_type, safe_name, file.content_type, len(content), digest, str(path), description, user))
        row = cur.fetchone()
        audit_event(cur, "accounting", "DOCUMENT_SUPPORT_UPLOADED", entity_type or "document_support", row["id"], user, role, description, None, _serialize(row), {"entry_id": entry_id})
    conn.commit()
    return {"status": "ok", "document": row}


@router.get("/documents")
def search_accounting_supports(
    entry_id: int | None = None,
    search: str | None = None,
    document_type: str | None = None,
    limit: int = Query(200, ge=1, le=1000),
    conn=Depends(get_db),
):
    _ensure_schema(conn)
    where, params = ["1=1"], []
    if entry_id:
        where.append("entry_id=%s"); params.append(entry_id)
    if document_type:
        where.append("document_type=%s"); params.append(document_type)
    if search:
        where.append("(filename ILIKE %s OR description ILIKE %s OR entity_id ILIKE %s)")
        params.extend([f"%{search}%", f"%{search}%", f"%{search}%"])
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(f"""
            SELECT * FROM accounting_document_supports
            WHERE {' AND '.join(where)}
            ORDER BY uploaded_at DESC, id DESC
            LIMIT %s
        """, [*params, limit])
        return {"data": cur.fetchall()}


@router.get("/documents/{support_id}/download")
def download_accounting_support(support_id: int, conn=Depends(get_db)):
    _ensure_schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM accounting_document_supports WHERE id=%s", (support_id,))
        row = cur.fetchone()
    if not row or not row.get("stored_path") or not os.path.exists(row["stored_path"]):
        raise HTTPException(404, "Document support not found")
    return FileResponse(row["stored_path"], filename=row["filename"], media_type=row.get("mime_type") or "application/octet-stream")


@router.get("/budgets")
def list_budgets(
    period: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    purpose: str = "ALL",
    status: str = "ALL",
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    _ensure_schema(conn)
    company = _company_from_header(x_company_code)
    filters, params = _budget_filters(company, period, date_from, date_to, purpose, status)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(
            f"""
            SELECT b.*,
                   COALESCE(a.account_name, b.account_code) AS account_name,
                   COALESCE(c.contrib_amount, 0) AS contributed_amount,
                   COALESCE(b.current_amount, 0) + COALESCE(c.contrib_amount, 0) AS progress_amount,
                   CASE
                     WHEN COALESCE(NULLIF(b.target_amount, 0), b.budget_amount, 0) = 0 THEN 0
                     ELSE ROUND(((COALESCE(b.current_amount, 0) + COALESCE(c.contrib_amount, 0))
                          / COALESCE(NULLIF(b.target_amount, 0), b.budget_amount, 1)) * 100, 2)
                   END AS progress_pct
            FROM accounting_budgets b
            LEFT JOIN accounting_accounts a ON a.account_code=b.account_code
            LEFT JOIN (
                SELECT budget_id, SUM(amount) AS contrib_amount
                FROM accounting_budget_contributions
                WHERE company_code=%s
                GROUP BY budget_id
            ) c ON c.budget_id=b.id
            WHERE {" AND ".join(filters)}
            ORDER BY COALESCE(b.target_date, (b.period || '-01')::date), b.purpose, b.account_code
            """,
            [company] + params,
        )
        return {"data": [_serialize(row) for row in cur.fetchall()]}


@router.put("/budget")
def upsert_budget(
    payload: dict,
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    _ensure_schema(conn)
    period = payload.get("period")
    _valid_period(period)
    account = str(payload.get("account_code") or "").strip()
    if not account:
        raise HTTPException(400, "account_code is required")
    company = _company_from_header(x_company_code)
    purpose = str(payload.get("purpose") or "BUDGET").strip().upper()
    if purpose not in {"BUDGET", "SAVINGS", "GOAL"}:
        raise HTTPException(400, "purpose must be BUDGET, SAVINGS or GOAL")
    status = str(payload.get("status") or "ACTIVE").strip().upper()
    if status not in {"ACTIVE", "PAUSED", "DONE", "CANCELLED"}:
        raise HTTPException(400, "status must be ACTIVE, PAUSED, DONE or CANCELLED")
    user, role = actor_from_headers(x_user, x_role, x_user_role)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        _get_postable_account(cur, account, "account_code")
        funding_code = str(payload.get("funding_bank_account_code") or "").strip()
        funding_account = _get_postable_account(cur, funding_code, "funding_bank_account_code") if funding_code else None
        cur.execute("SELECT * FROM accounting_budgets WHERE company_code=%s AND period=%s AND account_code=%s AND cost_center_code=%s AND currency_code=%s",
                    (company, period, account, payload.get("cost_center_code") or "", payload.get("currency_code") or "CRC"))
        before = cur.fetchone()
        cur.execute("""
            INSERT INTO accounting_budgets(
                company_code, period, account_code, cost_center_code, currency_code,
                budget_amount, purpose, name, target_date, monthly_contribution,
                target_amount, current_amount, funding_bank_account_code,
                funding_bank_account_name, status, notes, created_by
            )
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT(company_code,period,account_code,cost_center_code,currency_code) DO UPDATE SET
                company_code=EXCLUDED.company_code,
                budget_amount=EXCLUDED.budget_amount,
                purpose=EXCLUDED.purpose,
                name=EXCLUDED.name,
                target_date=EXCLUDED.target_date,
                monthly_contribution=EXCLUDED.monthly_contribution,
                target_amount=EXCLUDED.target_amount,
                current_amount=EXCLUDED.current_amount,
                funding_bank_account_code=EXCLUDED.funding_bank_account_code,
                funding_bank_account_name=EXCLUDED.funding_bank_account_name,
                status=EXCLUDED.status,
                notes=EXCLUDED.notes,
                updated_at=NOW()
            RETURNING *
        """, (
            company, period, account, payload.get("cost_center_code") or "", str(payload.get("currency_code") or "CRC").upper(),
            _money(payload.get("budget_amount")), purpose, payload.get("name") or None, payload.get("target_date") or None,
            _money(payload.get("monthly_contribution")), _money(payload.get("target_amount")), _money(payload.get("current_amount")),
            funding_account["account_code"] if funding_account else None,
            funding_account["account_name"] if funding_account else payload.get("funding_bank_account_name"),
            status, payload.get("notes"), user,
        ))
        after = cur.fetchone()
        audit_event(cur, "accounting", "BUDGET_UPSERTED", "accounting_budget", after["id"], user, role, payload.get("reason"), _serialize(before), _serialize(after))
    conn.commit()
    return {"status": "ok", "budget": after}


@router.delete("/budget/{budget_id}")
def delete_budget(
    budget_id: int,
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    _ensure_schema(conn)
    company = _company_from_header(x_company_code)
    user, role = actor_from_headers(x_user, x_role, x_user_role)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM accounting_budgets WHERE id=%s AND company_code=%s", (budget_id, company))
        before = cur.fetchone()
        if not before:
            raise HTTPException(404, "Budget not found")
        cur.execute("DELETE FROM accounting_budgets WHERE id=%s AND company_code=%s RETURNING *", (budget_id, company))
        after = cur.fetchone()
        audit_event(cur, "accounting", "BUDGET_DELETED", "accounting_budget", budget_id, user, role, None, _serialize(before), _serialize(after))
    conn.commit()
    return {"status": "ok", "deleted": budget_id}


@router.post("/budget/{budget_id}/contribution")
def post_budget_contribution(
    budget_id: int,
    payload: dict,
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    _ensure_schema(conn)
    company = _company_from_header(x_company_code)
    user, role = actor_from_headers(x_user, x_role, x_user_role)
    contribution_date = payload.get("contribution_date") or date.today().isoformat()
    try:
        contribution_date_obj = datetime.strptime(str(contribution_date), "%Y-%m-%d").date()
    except Exception:
        raise HTTPException(400, "contribution_date must use YYYY-MM-DD")
    amount = _money(payload.get("amount"))
    if amount <= 0:
        raise HTTPException(400, "amount must be greater than zero")
    reference = str(payload.get("reference") or "").strip()
    if not reference:
        raise HTTPException(400, "reference is required")
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM accounting_budgets WHERE id=%s AND company_code=%s FOR UPDATE", (budget_id, company))
        budget = cur.fetchone()
        if not budget:
            raise HTTPException(404, "Budget not found")
        debit_account = _get_postable_account(cur, budget["account_code"], "budget account")
        bank_code = str(payload.get("bank_account_code") or budget.get("funding_bank_account_code") or "").strip()
        bank_account = _get_postable_account(cur, bank_code, "bank_account_code")
        period = contribution_date_obj.strftime("%Y-%m")
        description = f"Aporte presupuesto/meta {budget.get('name') or budget.get('account_code')} - {reference}"
        cur.execute(
            """
            INSERT INTO accounting_entries(
                entry_date, period, description, origin, origin_id, created_by,
                workflow_status, company_code, currency_code, exchange_rate,
                posting_rule_code, posting_metadata, posted_by, posted_at
            )
            VALUES(%s,%s,%s,'BUDGET_CONTRIBUTION',%s,%s,'POSTED',%s,%s,1,%s,%s,%s,NOW())
            RETURNING id
            """,
            (
                contribution_date_obj, period, description, budget_id, user, company,
                str(payload.get("currency_code") or budget.get("currency_code") or "CRC").upper(),
                "BUDGET_CONTRIBUTION", Json({"budget_id": budget_id, "reference": reference}), user,
            ),
        )
        entry_id = cur.fetchone()["id"]
        cur.execute(
            """
            INSERT INTO accounting_lines(entry_id, account_code, account_name, debit, credit, line_description)
            VALUES(%s,%s,%s,%s,0,%s),(%s,%s,%s,0,%s,%s)
            """,
            (
                entry_id, debit_account["account_code"], debit_account["account_name"], amount, description,
                entry_id, bank_account["account_code"], bank_account["account_name"], amount, description,
            ),
        )
        cur.execute(
            """
            INSERT INTO accounting_budget_contributions(
                budget_id, company_code, contribution_date, amount, currency_code,
                bank_account_code, bank_account_name, reference, notes,
                accounting_entry_id, created_by
            )
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            RETURNING *
            """,
            (
                budget_id, company, contribution_date_obj, amount,
                str(payload.get("currency_code") or budget.get("currency_code") or "CRC").upper(),
                bank_account["account_code"], bank_account["account_name"], reference,
                payload.get("notes"), entry_id, user,
            ),
        )
        contribution = cur.fetchone()
        cur.execute("UPDATE accounting_budgets SET updated_at=NOW() WHERE id=%s", (budget_id,))
        audit_event(cur, "accounting", "BUDGET_CONTRIBUTION_POSTED", "accounting_budget", budget_id, user, role, None, None, _serialize(contribution), {"entry_id": entry_id})
    conn.commit()
    return {"status": "ok", "entry_id": entry_id, "contribution": _serialize(contribution)}


@router.get("/budget-vs-actual")
def budget_vs_actual(
    period: str,
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    _ensure_schema(conn)
    company = _company_from_header(x_company_code)
    start, end = _period_bounds(period)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("""
            SELECT b.period, b.account_code, COALESCE(a.account_name,b.account_code) account_name,
                   b.cost_center_code, b.currency_code, b.budget_amount,
                   COALESCE(SUM(l.debit-l.credit),0) actual_amount
            FROM accounting_budgets b
            LEFT JOIN accounting_accounts a ON a.account_code=b.account_code
            LEFT JOIN accounting_lines l ON l.account_code=b.account_code
            LEFT JOIN accounting_entries e ON e.id=l.entry_id AND e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s AND e.company_code=%s
            WHERE b.period=%s AND b.company_code=%s AND UPPER(COALESCE(b.purpose,'BUDGET'))='BUDGET'
            GROUP BY b.period,b.account_code,a.account_name,b.cost_center_code,b.currency_code,b.budget_amount
            ORDER BY b.account_code
        """, (start, end, company, period, company))
        rows = []
        for row in cur.fetchall():
            variance = _money(row["actual_amount"]) - _money(row["budget_amount"])
            rows.append(_serialize(row) | {"variance": _to_float(variance), "variance_pct": float((variance / _money(row["budget_amount"]) * 100).quantize(MONEY)) if _money(row["budget_amount"]) else 0.0})
    return {"period": period, "data": rows}


@router.get("/budget-report.xlsx")
def budget_report_excel(
    period: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    purpose: str = "ALL",
    status: str = "ALL",
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    _ensure_schema(conn)
    company = _company_from_header(x_company_code)
    filters, params = _budget_filters(company, period, date_from, date_to, purpose, status)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception as exc:
        raise HTTPException(500, f"No se pudo cargar Excel: {exc}")
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(
            f"""
            SELECT b.*, COALESCE(a.account_name,b.account_code) account_name,
                   COALESCE(c.contrib_amount,0) contributed_amount,
                   COALESCE(b.current_amount,0) + COALESCE(c.contrib_amount,0) progress_amount
            FROM accounting_budgets b
            LEFT JOIN accounting_accounts a ON a.account_code=b.account_code
            LEFT JOIN (
                SELECT budget_id, SUM(amount) contrib_amount
                FROM accounting_budget_contributions
                WHERE company_code=%s
                GROUP BY budget_id
            ) c ON c.budget_id=b.id
            WHERE {" AND ".join(filters)}
            ORDER BY COALESCE(b.target_date, (b.period || '-01')::date), b.purpose, b.account_code
            """,
            [company] + params,
        )
        rows = cur.fetchall()
        cur.execute(
            """
            SELECT c.*, b.name, b.purpose, b.account_code
            FROM accounting_budget_contributions c
            JOIN accounting_budgets b ON b.id=c.budget_id
            WHERE c.company_code=%s
              AND (%s::date IS NULL OR c.contribution_date >= %s::date)
              AND (%s::date IS NULL OR c.contribution_date <= %s::date)
            ORDER BY c.contribution_date DESC, c.id DESC
            """,
            [company, date_from, date_from, date_to, date_to],
        )
        contribs = cur.fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto Metas"
    ws.append([f"Reporte presupuesto, ahorro y metas | Empresa {company}"])
    headers = [
        "ID", "Periodo", "Tipo", "Nombre", "Cuenta", "Nombre cuenta", "Centro costo",
        "Moneda", "Presupuesto", "Meta", "Actual inicial", "Aportado", "Progreso",
        "Aporte mensual", "Fecha a cumplir", "Banco fondeo", "Nombre banco", "Estado", "Notas",
    ]
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003A75")
    for row in rows:
        ws.append([
            row.get("id"), row.get("period"), row.get("purpose"), row.get("name"), row.get("account_code"),
            row.get("account_name"), row.get("cost_center_code"), row.get("currency_code"),
            float(_money(row.get("budget_amount"))), float(_money(row.get("target_amount"))),
            float(_money(row.get("current_amount"))), float(_money(row.get("contributed_amount"))),
            float(_money(row.get("progress_amount"))), float(_money(row.get("monthly_contribution"))),
            row.get("target_date"), row.get("funding_bank_account_code"), row.get("funding_bank_account_name"),
            row.get("status"), row.get("notes"),
        ])
    ws2 = wb.create_sheet("Aportes")
    ws2.append(["ID", "Budget ID", "Fecha", "Tipo", "Nombre", "Cuenta", "Monto", "Moneda", "Banco", "Nombre banco", "Referencia", "Asiento", "Notas"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003A75")
    for row in contribs:
        ws2.append([
            row.get("id"), row.get("budget_id"), row.get("contribution_date"), row.get("purpose"), row.get("name"),
            row.get("account_code"), float(_money(row.get("amount"))), row.get("currency_code"),
            row.get("bank_account_code"), row.get("bank_account_name"), row.get("reference"),
            row.get("accounting_entry_id"), row.get("notes"),
        ])
    for sheet in wb.worksheets:
        for column_cells in sheet.columns:
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 42)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Presupuesto_Ahorro_Metas_{period or 'general'}_{purpose}_{status}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _smart_alerts(cur, period: str):
    start, end = _period_bounds(period)
    previous_start = (start.replace(day=1) - timedelta(days=1)).replace(day=1)
    alerts = []
    def add(severity, code, title, message, entity_type=None, entity_id=None, metadata=None):
        alerts.append({"severity": severity, "code": code, "title": title, "message": message, "entity_type": entity_type, "entity_id": str(entity_id) if entity_id is not None else None, "metadata": metadata or {}})
    cur.execute("""
        WITH current AS (
            SELECT l.account_code, MAX(l.account_name) account_name, COALESCE(SUM(l.debit),0) amount
            FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
            WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s AND l.account_code LIKE '5%%'
            GROUP BY l.account_code
        ), previous AS (
            SELECT l.account_code, COALESCE(SUM(l.debit),0) amount
            FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
            WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s AND l.account_code LIKE '5%%'
            GROUP BY l.account_code
        )
        SELECT c.account_code,c.account_name,c.amount current_amount,COALESCE(p.amount,0) previous_amount
        FROM current c LEFT JOIN previous p ON p.account_code=c.account_code
        WHERE c.amount > 0 AND (COALESCE(p.amount,0)=0 OR c.amount >= COALESCE(p.amount,0)*3)
        ORDER BY c.amount DESC
        LIMIT 25
    """, (start, end, previous_start, start))
    for row in cur.fetchall():
        add("warning", "EXPENSE_SPIKE", "Gasto con aumento inusual", f"{row['account_code']} {row['account_name']} subio a {_to_float(row['current_amount']):,.2f} vs {_to_float(row['previous_amount']):,.2f}.", "account", row["account_code"], _serialize(row))
    cur.execute("SELECT COUNT(*) count FROM tax_electronic_documents d WHERE d.issue_datetime >= %s AND d.issue_datetime < %s AND EXISTS(SELECT 1 FROM tax_document_lines l WHERE l.document_id=d.id AND COALESCE(l.cabys_code,'')='')", (start, end))
    missing_cabys = int((cur.fetchone() or {}).get("count") or 0)
    if missing_cabys:
        add("critical", "XML_MISSING_CABYS", "XML sin CAByS", f"{missing_cabys} documentos tienen lineas sin CAByS.", "tax", period)
    cur.execute("SELECT COUNT(*) count FROM bank_reconciliation_statement_lines l JOIN bank_reconciliation_statements s ON s.id=l.statement_id WHERE s.statement_period=%s AND l.match_status='OPEN'", (period,))
    open_bank = int((cur.fetchone() or {}).get("count") or 0)
    if open_bank:
        add("critical", "BANK_NOT_RECONCILED", "Banco no conciliado", f"{open_bank} partidas bancarias siguen abiertas.", "bank_reconciliation", period)
    cur.execute("""
        SELECT e.id, e.origin, l.account_code, l.account_name, COUNT(*) OVER(PARTITION BY l.account_code) frequency
        FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
        WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s
        ORDER BY frequency ASC, e.id DESC
        LIMIT 20
    """, (start, end))
    for row in cur.fetchall():
        if int(row["frequency"] or 0) <= 1:
            add("info", "UNCOMMON_ACCOUNT_USAGE", "Cuenta poco comun", f"Asiento {row['id']} usa cuenta poco frecuente {row['account_code']} {row['account_name']}.", "accounting_entry", row["id"], _serialize(row))
    return alerts


@router.get("/smart-alerts")
def smart_alerts(period: str, conn=Depends(get_db)):
    _ensure_schema(conn)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        alerts = _smart_alerts(cur, period)
    return {"period": period, "data": alerts, "counts": {level: sum(1 for a in alerts if a["severity"] == level) for level in ("critical", "warning", "info")}}


@router.get("/executive-dashboard")
def executive_dashboard(period: str, conn=Depends(get_db)):
    _ensure_schema(conn)
    start, end = _period_bounds(period)
    as_of = min(end - timedelta(days=1), date.today())
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("""
            SELECT COALESCE(SUM(CASE WHEN l.account_code LIKE '4%%' THEN l.credit-l.debit ELSE 0 END),0) revenue,
                   COALESCE(SUM(CASE WHEN l.account_code LIKE '5%%' THEN l.debit-l.credit ELSE 0 END),0) expenses
            FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
            WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s AND e.entry_date <= CURRENT_DATE
        """, (start, end))
        ledger = cur.fetchone()
        cur.execute("""
            SELECT COALESCE(SUM(l.debit-l.credit),0) banks
            FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
            WHERE e.workflow_status='POSTED'
              AND e.entry_date <= %s
              AND (l.account_code LIKE '1.1.02%%' OR LOWER(l.account_name) LIKE '%%banco%%')
        """, (as_of,))
        bank_row = cur.fetchone() or {"banks": 0}
        cur.execute("SELECT COALESCE(SUM(saldo_pendiente),0) total, COUNT(*) count FROM collections WHERE COALESCE(saldo_pendiente,0)>0 AND fecha_vencimiento<CURRENT_DATE")
        overdue_ar = cur.fetchone()
        cur.execute("SELECT COALESCE(SUM(balance),0) total, COUNT(*) count FROM payment_obligations WHERE active=TRUE AND record_type='OBLIGATION' AND COALESCE(balance,0)>0 AND due_date BETWEEN CURRENT_DATE AND CURRENT_DATE+INTERVAL '15 days'")
        upcoming_ap = cur.fetchone()
        iva = tax_iva(period=period, conn=conn)
        alerts = _smart_alerts(cur, period)
        cur.execute("""
            SELECT COALESCE(cliente,'SIN CLIENTE') client, COALESCE(SUM(honorarios),0) revenue
            FROM servicios
            WHERE COALESCE(fecha_inicio,fecha_fin,CURRENT_DATE) >= %s AND COALESCE(fecha_inicio,fecha_fin,CURRENT_DATE) < %s
            GROUP BY cliente ORDER BY revenue DESC LIMIT 10
        """, (start, end))
        top_clients = cur.fetchall()
        cur.execute("""
            SELECT l.account_code, MAX(l.account_name) account_name, COALESCE(SUM(l.debit),0) amount
            FROM accounting_entries e JOIN accounting_lines l ON l.entry_id=e.id
            WHERE e.workflow_status='POSTED' AND e.entry_date >= %s AND e.entry_date < %s AND l.account_code LIKE '5%%'
            GROUP BY l.account_code ORDER BY amount DESC LIMIT 10
        """, (start, end))
        top_expenses = cur.fetchall()
    revenue = _money(ledger["revenue"])
    expenses = _money(ledger["expenses"])
    return {
        "period": period,
        "liquidity": {"banks": _to_float(bank_row["banks"]), "as_of": as_of.isoformat()},
        "margin": {"revenue": _to_float(revenue), "expenses": _to_float(expenses), "profit": _to_float(revenue - expenses), "margin_pct": float(((revenue - expenses) / revenue * 100).quantize(MONEY)) if revenue else 0.0},
        "overdue_ar": _serialize(overdue_ar),
        "upcoming_payments": _serialize(upcoming_ap),
        "iva_estimated": iva.get("fiscal", {}),
        "cash_flow_projected": {"next_15_days_payments": _to_float(upcoming_ap.get("total")), "current_banks_less_upcoming": _to_float(_money(bank_row["banks"]) - _money(upcoming_ap.get("total")))},
        "top_clients": [_serialize(row) for row in top_clients],
        "top_expenses": [_serialize(row) for row in top_expenses],
        "alerts": {"critical": sum(1 for a in alerts if a["severity"] == "critical"), "warning": sum(1 for a in alerts if a["severity"] == "warning"), "info": sum(1 for a in alerts if a["severity"] == "info")},
    }


@router.post("/portia/review")
def portia_accounting_review(payload: dict, conn=Depends(get_db)):
    _ensure_schema(conn)
    period = payload.get("period") or date.today().strftime("%Y-%m")
    language = (payload.get("language") or "ES").upper()
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        dashboard = executive_dashboard(period=period, conn=conn)
        alerts = _smart_alerts(cur, period)
        budget = budget_vs_actual(period=period, conn=conn)
        tax = tax_deep_summary(period=period, conn=conn)
    lines = []
    if language.startswith("ES"):
        lines.append(f"PORTIA reviso el periodo {period}.")
        margin = dashboard.get("margin", {})
        lines.append(f"Resultado mensual estimado: ingresos {margin.get('revenue',0):,.2f}, gastos {margin.get('expenses',0):,.2f}, utilidad {margin.get('profit',0):,.2f}.")
        if alerts:
            lines.append(f"Se detectaron {len(alerts)} alertas inteligentes; priorizar criticas y advertencias antes del cierre.")
        iva = tax.get("iva", {}).get("differences", {})
        if any(abs(float(iva.get(k) or 0)) > 0.01 for k in iva):
            lines.append("El IVA documental y contable presenta diferencias; revisar XML, cuentas de IVA y documentos sin CAByS.")
        variances = [row for row in budget.get("data", []) if abs(float(row.get("variance_pct") or 0)) >= 20]
        if variances:
            lines.append(f"Hay {len(variances)} cuentas con desviacion presupuestaria superior al 20%.")
        lines.append("Sugerencia: documentar soportes en asientos relevantes, cerrar bancos, revisar auxiliares y ejecutar revaluacion USD antes del cierre mensual.")
    else:
        lines.append(f"PORTIA reviewed period {period}.")
        lines.append("Prioritize critical smart alerts, VAT reconciliation, bank reconciliation, supporting documents and USD revaluation before closing.")
    return {"period": period, "language": language, "commentary": "\n".join(lines), "context": {"dashboard": dashboard, "alerts": alerts, "budget": budget, "tax": tax}}
