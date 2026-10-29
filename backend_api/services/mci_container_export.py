"""MCI container exports using the user's retained layout, not the MSL form."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from services.mci_survey_schema import SurveyContent
from services.mci_container_schema import TITLES, GROUPS
from services.mci_survey_export import set_text
from services.web_docx_layout import fit_web_docx


def build_word(row):
    if row.get('company_code') != 'MCI-CR' or row.get('kind') != 'container':
        raise ValueError('MCI container context required')
    c = SurveyContent.model_validate(row['content'])
    values = c.model_dump()
    values.update({k: row.get(k, '') for k in ('vessel','requested_by','report_number')})
    values['survey_place'] = c.survey_place or row.get('location', c.location)
    doc = Document(Path(__file__).resolve().parents[1]/'templates/mci_container_original.docx')
    parents = {child: parent for parent, children in GROUPS.items() for child in children}
    def enabled(key):
        section = c.sections.get(key)
        return bool(section and section.enabled and (key not in parents or enabled(parents[key])))
    current = None
    extra = {'instructions','representatives','particulars','timesheet','attachments','introduction'}
    def append_extra(before, key):
        if key not in extra or not enabled(key):return
        for block in c.sections[key].blocks:
            if not block.text.strip():continue
            element=deepcopy(normal)
            set_text(element,block.text)
            before.addprevious(element)
    normal = deepcopy(next(p._p for p in doc.paragraphs if p.text=='{{blocks:damage}}'))
    for p in list(doc.paragraphs):
        heading = re.fullmatch(r'\{\{heading:(\w+)\}\}',p.text)
        if heading:
            append_extra(p._p,current)
            current = heading[1]
            if enabled(current):
                set_text(p,TITLES[current])
                p.paragraph_format.keep_with_next=True
            else:p._p.getparent().remove(p._p)
            continue
        disclaimer=p.text.startswith('THIS BONAFIDE')
        if disclaimer:append_extra(p._p,current)
        if current and not enabled(current) and not disclaimer:
            p._p.getparent().remove(p._p)
            continue
        block=re.fullmatch(r'\{\{blocks:(\w+)\}\}',p.text)
        if block:
            texts=[b.text for b in c.sections[block[1]].blocks if b.text.strip()]
            if block[1]=='enclosure':texts=['\n'.join(texts)] if texts else []
            for text in texts:
                element=deepcopy(p._p)
                set_text(element,text)
                para=Paragraph(element,p._parent)
                para.paragraph_format.keep_together=False
                para.paragraph_format.keep_with_next=False
                p._p.addprevious(element)
            p._p.getparent().remove(p._p)
        elif '{{' in p.text:
            set_text(p,re.sub(r'\{\{(\w+)\}\}',lambda m:str(values.get(m[1]) or ''),p.text))
    # Preserve field codes and artwork: replace only the report-number text nodes.
    for section in doc.sections:
        for part in (section.header,section.footer):
            for node in part._element.xpath('.//w:t'):
                if node.text:node.text=node.text.replace('{{report_number}}',row['report_number'])
    # Explicit section references keep the retained running furniture on forced
    # page breaks in both Word and the server's LibreOffice conversion.
    section_props=deepcopy(doc.sections[0]._sectPr)
    for p in list(doc.paragraphs):
        if not p.paragraph_format.page_break_before:continue
        p.paragraph_format.page_break_before=False
        boundary=OxmlElement('w:p');props=OxmlElement('w:pPr')
        section=deepcopy(section_props)
        typ=section.find(qn('w:type'))
        if typ is None:
            typ=OxmlElement('w:type');section.insert(2,typ)
        typ.set(qn('w:val'),'nextPage')
        props.append(section);boundary.append(props);p._p.addprevious(boundary)
    originals=(doc.sections[0].header,doc.sections[0].footer)
    for section in list(doc.sections)[1:]:
        for attr,original in zip(('header','footer'),originals):
            target=getattr(section,attr)
            # Other sections still use the original relationship. Removing it
            # through the property setter would drop that shared relationship.
            for ref in list(section._sectPr.findall(qn('w:'+attr+'Reference'))):
                section._sectPr.remove(ref)
            target.is_linked_to_previous=False
            for child in list(target._element):target._element.remove(child)
            for child in original._element:
                element=deepcopy(child)
                for node in element.iter():
                    for key,value in list(node.attrib.items()):
                        if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                            rel=original.part.rels[value]
                            node.set(key,target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,rel.is_external))
                target._element.append(element)
    setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true')
    doc.settings.element.append(setting)
    fit_web_docx(doc)
    output=BytesIO();doc.save(output)
    return output.getvalue()
