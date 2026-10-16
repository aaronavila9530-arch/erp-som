"""Bounded workbook readers; only explicitly named driver columns are consumed."""
from io import BytesIO
import re
import unicodedata
import zipfile
from openpyxl import load_workbook


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD', text(value).lower())
                           if unicodedata.category(c) != 'Mn').split())


def workbook_rows(content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError('Archivo mayor a 5 MB')
    if content.startswith(b'\xd0\xcf\x11\xe0'):
        import xlrd
        wb = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            for sheet in wb.sheets():
                if sheet.nrows > 10001 or sheet.ncols > 150:
                    raise ValueError('Hoja demasiado grande')
                yield sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            wb.release_resources()
    else:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Excel descomprimido demasiado grande')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in wb:
                if sheet.max_row > 10001 or sheet.max_column > 150:
                    raise ValueError('Hoja demasiado grande')
                yield sheet.title, list(sheet.iter_rows(values_only=True))
        finally:
            wb.close()


def driver_records(content, sheet_name=None):
    records = {}
    conflicts = set()
    for name, rows in workbook_rows(content):
        if sheet_name and name != sheet_name:
            continue
        columns = None
        for values in rows:
            names = [header(v) for v in values]
            if 'placa' in names and ('choferes' in names or 'chofer' in names):
                columns = (names.index('placa'), names.index('choferes') if 'choferes' in names else names.index('chofer'))
                continue
            if columns is None:
                continue
            plate, driver = (text(values[i]) if i < len(values) else '' for i in columns)
            if not plate and not driver:
                continue
            if not plate or not driver or driver.startswith('#'):
                raise ValueError(f'{name}: placa o chofer incompleto ({plate})')
            plate = re.sub(r'[\s-]+', '', plate).upper()
            driver = ' '.join(driver.split())
            if len(plate) > 30 or len(driver) > 150:
                raise ValueError('Placa o chofer demasiado largo')
            if plate in records and header(records[plate]) != header(driver):
                conflicts.add(plate)
            records[plate] = driver
    if conflicts:
        raise ValueError('Placas con varios choferes: ' + ', '.join(sorted(conflicts)))
    if not records:
        raise ValueError('No se encontraron columnas Placa y Choferes con datos')
    return records
