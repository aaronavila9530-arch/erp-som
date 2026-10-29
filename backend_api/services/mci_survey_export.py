"""Generate only from retained MCI references; never rewrite source templates."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from services.mci_survey_schema import SurveyContent, export_sections

TEMPLATES = Path(__file__).resolve().parents[1] / 'templates'
MONTHS = ('January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December')


def english_date(value, time=False):
    if not value:
        return ''
    result = f'{MONTHS[value.month-1]} {value.day}, {value.year}'
    return result + (value.strftime(' %H:%M') + ' LT' if time else '')


def set_text(p, value):
    """Keep paragraph/run formatting; replace text without carrying sample facts."""
    p = p if isinstance(p, Paragraph) else Paragraph(p, None)
    runs = list(p.runs)
    if runs:
        runs[0].text = str(value or '')
        for run in runs[1:]:
            run.text = ''
    else:
        p.add_run(str(value or ''))
    return p._p


def build_word(row):
    if row.get('company_code') != 'MCI-CR':
        raise ValueError('MCI document requires MCI context')
    content = SurveyContent.model_validate(row['content'])
    source = Document(TEMPLATES / 'mci_general_original.docx')
    doc = Document(TEMPLATES / 'mci_letterhead_original.docx')
    body = doc.element.body
    insertion = body[0]
    existing_styles = {s.style_id for s in doc.styles}
    for style in source.styles:
        if style.style_id not in existing_styles:
            doc.styles.element.append(deepcopy(style.element))
    def add(element):
        insertion.addprevious(element)
    def clone(text, value=None):
        p = next(p for p in source.paragraphs if p.text.strip() == text)
        element = deepcopy(p._p)
        for sect in list(element.xpath('.//w:sectPr')):
            sect.getparent().remove(sect)
        if value is not None:
            set_text(element, value)
        return element
    add(clone('PROFESSIONAL TECHNICAL REPORT'))
    add(clone('M/V LINA AKSOY', row['vessel']))
    report_type = {'vessel':'Vessel Survey', 'condition':'Vessel Condition Survey', 'cargo':'Cargo Condition Survey', 'logistics':'Logistics Supervision'}.get(row.get('kind'), 'Survey Report')
    add(clone('Multiple mooring-line failures during bulk cargo discharge\nBerth No. 4 - Puerto Caldera, Costa Rica', report_type+'\n'+row.get('location', content.location)))
    table = deepcopy(source.tables[0]._tbl)
    values = [row['requested_by'], f"{row['vessel']} / IMO {content.imo}", content.length_and_tonnage,
              content.cargo, ' - '.join(english_date(d, time=True) for d in (content.period_start,content.period_end) if d),
              english_date(content.issue_date), content.report_status or 'Draft']
    for tr, value in zip(table.xpath('./w:tr'), values):
        cell = tr.xpath('./w:tc')[1]
        paragraphs = cell.xpath('./w:p')
        set_text(paragraphs[0], value)
        for p in paragraphs[1:]:
            cell.remove(p)
    add(table)
    text='This report is issued exclusively for the client that instructed MSL and its advisers.\nIt does not constitute an admission of liability or a legal opinion.'
    add(clone(text, text.replace('instructed MSL', 'instructed MCI')))
    heading_sources = {
        'instructions':'INSTRUCTIONS, SCOPE AND METHODOLOGY', 'work':'1.1 Work performed',
        'limitations':'1.2 Main limitations of this report',
        'chronology':'1.3 Time and place of arrival and chronology of events:',
        'representatives':'Representatives', 'findings':'SURVEY FINDINGS', 'visual':'2.1 Visual inspection',
        'provided':'Documentation provided', 'missing':'Documentation not provided',
        'conclusion':'2. CONCLUSION', 'enclosure':'3. ENCLOSURE',
    }
    normal = next(p for p in source.paragraphs if p.text.strip() == 'Master:')
    headings = {'instructions':'1. INSTRUCTIONS, SCOPE AND METHODOLOGY',
                'work':'1.1 Work performed', 'limitations':'1.2 Main limitations of this report',
                'chronology':'1.3 Time and place of arrival and chronology of events:',
                'representatives':'Representatives', 'findings':'2. SURVEY FINDINGS',
                'visual':'2.1 Visual inspection', 'provided':'2.2 Documentation provided',
                'missing':'2.3 Documentation not provided', 'conclusion':'3. CONCLUSION', 'enclosure':'4. ENCLOSURE'}
    sections = export_sections(content)
    if row.get('kind') == 'logistics':
        from services.mci_logistics_schema import TITLES, GROUPS
        parents = {child:parent for parent,children in GROUPS.items() for child in children}
        sections = [(key,TITLES[key],[b.text for b in content.sections[key].blocks if b.text.strip()])
                    for key in TITLES if key in content.sections and content.sections[key].enabled
                    and (key not in parents or (content.sections.get(parents[key]) and content.sections[parents[key]].enabled))]
        heading_sources.update(particulars='SURVEY FINDINGS', process='SURVEY FINDINGS',
                               documentary='2.1 Visual inspection', operational='2.1 Visual inspection', incidents='2.1 Visual inspection')
        headings.update(particulars='2. VESSEL PARTICULARS', process='3. SUPERVISION PROCESS',
                        findings='4. SURVEY FINDINGS', documentary='4.1 Documentary findings',
                        operational='4.2 Operational control findings', incidents='4.3 Incidents',
                        conclusion='5. CONCLUSION', enclosure='6. ENCLOSURE')
    for key, title, blocks in sections:
        heading = clone(heading_sources[key], headings[key])
        # Explicit hierarchical labels must not inherit the template's flat list.
        ppr = heading.get_or_add_pPr()
        ppr.get_or_add_keepNext().val = True
        numbering = ppr.get_or_add_numPr()
        numbering.get_or_add_numId().val = 0
        add(heading)
        if key == 'representatives':
            blocks = [f'Master: {content.master}', f'Chief officer: {content.chief_officer}', f'Surveyor: {content.surveyor}'] + blocks
        elif key == 'particulars' and row.get('kind') == 'logistics':
            blocks = [f'Vessel: {row["vessel"]}', f'Flag / port of registry: {content.flag_port_registry}',
                      f'GRT: {content.grt}', f'NRT: {content.nrt}', f'IMO: {content.imo}',
                      f'Year built: {content.build_year}'] + blocks
        elif key == 'chronology' and row.get('kind') == 'logistics':
            blocks = [f'{label}: {english_date(value)}' for label,value in (
                ('Arrival',content.arrival_date), ('Inspection',content.inspection_date),
                ('Supervision completed',content.supervision_completed_date)) if value] + blocks
        elif key == 'enclosure':
            blocks = ['Picture link: ' + '\n'.join(blocks)]
        for text in blocks:
            p = deepcopy(normal._p)
            set_text(p, text)
            add(p)
    disclaimer = next(p.text for p in source.paragraphs if p.text.startswith('THIS BONA FIDE'))
    add(clone(disclaimer))
    # The first section has a clean footer; the separate signature page retains
    # the original footer-anchored seal and logo, and the original body spacers.
    break_p = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr')
    section = deepcopy(doc.sections[0]._sectPr)
    section_type = section.find(qn('w:type'))
    if section_type is None:
        section_type = OxmlElement('w:type'); section.append(section_type)
    section_type.set(qn('w:val'), 'nextPage')
    ppr.append(section); break_p.append(ppr); add(break_p)
    first, last = doc.sections[0], doc.sections[-1]
    original_footer = last.footer
    footer_elements = [deepcopy(e) for e in original_footer._element]
    # Remove only the first section's reference before creating its distinct footer.
    for ref in list(first._sectPr.findall(qn('w:footerReference'))):
        first._sectPr.remove(ref)
    first.footer.is_linked_to_previous = False
    clean_footer = first.footer
    for child in list(clean_footer._element):
        clean_footer._element.remove(child)
    for element in footer_elements:
        for drawing in list(element.xpath('.//w:drawing')):
            drawing.getparent().remove(drawing)
        # Hyperlinks and other referenced objects retain their part relationships.
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    rel = original_footer.part.rels[value]
                    node.set(attr, clean_footer.part.relate_to(rel.target_ref if rel.is_external else rel.target_part, rel.reltype, rel.is_external))
        clean_footer._element.append(element)
    for p in doc.paragraphs:
        if p.text.strip() == '{cert_num}':
            set_text(p, row['report_number'])
        elif 'Pabel' in p.text:
            set_text(p, content.surveyor)
        elif p.text.strip() == 'Marine Surveyors & Logistics Group':
            set_text(p, 'Marine Claims & Risk Intelligence')
    # Place the linked report number into the existing header, not the filename.
    doc.settings.odd_and_even_pages_header_footer = False
    seen_headers=set()
    for section in doc.sections:
        section.different_first_page_header_footer = False
        # Keep section children in OOXML order. Out-of-order references cause
        # converters to create spurious page styles without a running header.
        references=[]
        for tag in ('headerReference','footerReference'):
            refs=section._sectPr.findall(qn('w:'+tag))
            default=next((r for r in refs if r.get(qn('w:type'))=='default'),None)
            if default is not None:
                for ref in refs:
                    section._sectPr.remove(ref)
                references.append(deepcopy(default))
        for ref in reversed(references):
            section._sectPr.insert(0,ref)
        page_type=section._sectPr.find(qn('w:type'))
        if page_type is not None:
            section._sectPr.remove(page_type)
            section._sectPr.insert(len(references),page_type)
        header=section.header
        if header.part.partname in seen_headers:
            continue
        seen_headers.add(header.part.partname)
        for p in header._element.xpath('.//w:p'):
            if 'CONFIDENTIAL - WITHOUT PREJUDICE' in ''.join(p.xpath('.//w:t/text()')):
                number=deepcopy(p);set_text(number,f"Cert. No. {row['report_number']}");p.addprevious(number)
                break
    setting = OxmlElement('w:updateFields'); setting.set(qn('w:val'), 'true')
    doc.settings.element.append(setting)
    from services.web_docx_layout import fit_web_docx
    fit_web_docx(doc)
    output = BytesIO(); doc.save(output)
    return output.getvalue()
