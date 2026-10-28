"""Conservative flow layout for generated web documents, never source templates."""
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


def _paragraph(paragraph, width, height):
    pf = paragraph.paragraph_format
    if paragraph.text.strip():
        pf.widow_control = True
        if len(paragraph.text) > 600:
            pf.keep_together = False
    # Decorative floating artwork, header/footer anchors and textboxes remain
    # owned by the individual template. Only inline body pictures participate.
    available = max(1, width - max(0, pf.left_indent or 0) - max(0, pf.right_indent or 0))
    for inline in paragraph._p.xpath('./w:r/w:drawing/wp:inline'):
        extent = inline.find(qn('wp:extent'))
        if extent is None:
            continue
        cx, cy = int(extent.get('cx', 0)), int(extent.get('cy', 0))
        if not cx or not cy:
            continue
        scale = min(1, available / cx, height / cy)
        if scale == 1:
            continue
        target_width, target_height = round(cx * scale), round(cy * scale)
        extent.set('cx', str(target_width))
        extent.set('cy', str(target_height))
        for transform in inline.xpath('.//a:xfrm/a:ext'):
            transform.set('cx', str(target_width))
            transform.set('cy', str(target_height))


def _table(table, available, height):
    grid = list(table._tbl.tblGrid)
    widths = [int(col.get(qn('w:w'), 0)) for col in grid]
    total = sum(widths)
    if not total:
        return
    indent = table._tbl.tblPr.find(qn('w:tblInd'))
    offset = max(0, int(indent.get(qn('w:w'), 0))) if indent is not None else 0
    limit = max(1, available // 635 - offset)
    if total > limit:
        widths = [max(1, int(value * limit / total)) for value in widths]
        for col, value in zip(grid, widths):
            col.set(qn('w:w'), str(value))
    # Fixed column geometry lets content grow down, not beyond page margins.
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn('w:tblW'))
    if table_width is not None:
        table_width.set(qn('w:type'), 'dxa')
        table_width.set(qn('w:w'), str(sum(widths)))
    seen = set()
    for row in table.rows:
        if row._tr.trPr is not None:
            for row_height in row._tr.trPr.findall(qn('w:trHeight')):
                if row_height.get(qn('w:hRule')) == 'exact':
                    row_height.set(qn('w:hRule'), 'atLeast')
            if sum(len(cell.text) for cell in row.cells) > 600:
                for flag in list(row._tr.trPr.findall(qn('w:cantSplit'))):
                    row._tr.trPr.remove(flag)
        col = 0
        if row._tr.trPr is not None:
            before = row._tr.trPr.find(qn('w:gridBefore'))
            if before is not None:
                col = int(before.get(qn('w:val'), 0))
        for tc in row._tr.tc_lst:
            span = tc.grid_span
            cell_width = sum(widths[col:col + span])
            if cell_width:
                tc.width = cell_width * 635
            col += span
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            props = cell._tc.get_or_add_tcPr()
            for flag in list(props.findall(qn('w:noWrap'))):
                props.remove(flag)
            cell_width = int(cell.width or available)
            # Respect explicit padding; Word's default cell padding is 108 twips.
            margins = props.find(qn('w:tcMar'))
            padding = 0
            for side in ('left', 'right'):
                value = margins.find(qn('w:' + side)) if margins is not None else None
                padding += int(value.get(qn('w:w'), 108)) if value is not None else 108
            cell_width = max(1, cell_width - padding * 635)
            for p in cell.paragraphs:
                _paragraph(p, cell_width, height)
            for nested in cell.tables:
                _table(nested, cell_width, height)


def fit_web_docx(doc):
    """Keep fonts, margins, headings, headers, footers and artwork unchanged."""
    sections = list(doc.sections)
    index = 0
    for block in doc.element.body:
        section = sections[min(index, len(sections) - 1)]
        width = int(section.page_width - section.left_margin - section.right_margin
                    - (section.gutter or 0))
        height = max(1, int(section.page_height - section.top_margin - section.bottom_margin) - 152400)
        if block.tag == qn('w:tbl'):
            _table(Table(block, doc), width, height)
        elif block.tag == qn('w:p'):
            _paragraph(Paragraph(block, doc), width, height)
            if block.xpath('./w:pPr/w:sectPr'):
                index += 1


def fit_vessel_web_docx(doc):
    """Repair known Condition/Cargo template footer and signature anchors."""
    for paragraph in doc.paragraphs:
        for anchor in list(paragraph._p.xpath('./w:r/w:drawing/wp:anchor')):
            if not anchor.xpath('.//pic:pic'):
                continue
            extent = anchor.find(qn('wp:extent'))
            if extent is None:
                continue
            width, height = int(extent.get('cx')), int(extent.get('cy'))
            scale = min(1, int(Inches(.65)) / max(1, height))
            for target in [extent] + anchor.xpath('.//a:xfrm/a:ext'):
                target.set('cx', str(round(width * scale)))
                target.set('cy', str(round(height * scale)))
            inline = OxmlElement('wp:inline')
            for tag in ('wp:extent', 'wp:effectExtent', 'wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
                element = anchor.find(qn(tag))
                if element is not None:
                    inline.append(element)
            anchor.getparent().replace(anchor, inline)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.line_spacing = 1.0
    seen = set()
    for section in doc.sections:
        footer = section.footer
        if footer.part.partname in seen:
            continue
        seen.add(footer.part.partname)
        for p in footer.paragraphs:
            if 'Page ' not in p.text or not p._p.xpath('.//w:instrText'):
                continue
            label = ' '.join(p.text.split('Page ', 1)[0].split())
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pf = p.paragraph_format
            pf.tab_stops.clear_all()
            pf.left_indent = pf.right_indent = pf.first_line_indent = Pt(0)
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            if label:
                p.add_run(label).font.size = Pt(7)
                p.add_run().add_break()
            p.add_run('Page ').font.size = Pt(8)
            for instruction, suffix in [('PAGE', ' of '), ('NUMPAGES', '')]:
                field = OxmlElement('w:fldSimple')
                field.set(qn('w:instr'), instruction)
                p._p.append(field)
                if suffix:
                    p.add_run(suffix).font.size = Pt(8)
    fit_web_docx(doc)
