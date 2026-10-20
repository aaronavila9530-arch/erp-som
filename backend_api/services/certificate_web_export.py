"""Isolated copies of the audited desktop templates; no template file writes."""
from copy import copy
from pathlib import Path
import tempfile

from services.certificate_web_fields import TYPES, EXCEL_FIELDS
from services.container_web_export import convert

TEMPLATES = Path(__file__).resolve().parents[1] / "templates"
EXCEL_TEMPLATES = {
    "holds": ("holds_inspection_certificate.xlsx", "VESSEL HOLDS INSPECTION CERTIFI", "A1:H51"),
    "sampling": ("sampling_certificate_template.xlsx", "CERTIFICATE SAMPLING", "A1:I49"),
    "sealing": ("sealing_certificate_template.xlsx", "SEALING CERTIFICATE", "A1:I46"),
}
CELL_FIELDS = {
    "holds": {"C8":"vessel", "G8":"voyage", "C9":"load_port", "G9":"place", "C10":"installation",
              "C11":"date", "C12":"product", "E16":"inspection_time", "E19":"vessel_holds",
              "A24":"vessel_holds_status", "A27":"product", "D29":"cargo_holds", "C31":"accepted_time",
              "B33":"place_location", "B35":"place_date", "A41":"remarks", "E50":"master_chief_officer"},
    "sampling": {"H2":"certificate_no", "B11":"vessel", "G11":"date", "B13":"place", "G13":"cargo",
                 "A31":"observations", "H39":"closing_date", "H41":"closing_time", "F46":"master"},
    "sealing": {"B4":"vessel", "G4":"date", "B6":"location", "G6":"cargo", "A33":"remarks",
                "H35":"closing_date", "H37":"closing_time", "B41":"chief_officer"},
}


def text(value):
    return "" if value is None else str(value)


def _put(sheet, address, value):
    cell = sheet[address]
    cell.value = text(value)
    cell.data_type = "s"  # Narrative/identifiers never execute as spreadsheet formulas.
    cell.number_format = "@"


def _merge(sheet, address):
    if address not in sheet.merged_cells:
        sheet.merge_cells(address)


def excel(kind, data, target):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from math import ceil

    filename, sheet_name, area = EXCEL_TEMPLATES[kind]
    template = TEMPLATES / filename
    book = load_workbook(template, keep_links=False)
    try:
        sheet = book[sheet_name]
        cached = load_workbook(template, data_only=True, keep_links=False)
        try:
            # Freeze legacy formulas to their original cached value before
            # explicitly mapping every dynamic output below (no stale formulas).
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        _put(sheet, cell.coordinate, cached[sheet_name][cell.coordinate].value)
        finally:
            cached.close()
        for idx, key in enumerate(EXCEL_FIELDS[kind], 1):
            _put(book["data"], f"B{idx}", data.get(key))
        for address, key in CELL_FIELDS[kind].items():
            _put(sheet, address, data.get(key))
        if kind == "holds":
            _put(sheet, "B37", f"Water Tightness / Hose Test time: {text(data.get('hose_test_start'))} to {text(data.get('hose_test_end'))} hours local time")
            boxes = ["C8:E8", "G8:H8", "C9:E9", "G9:H9", "C10:H10", "C11:H11", "C12:H12",
                     "A24:H25", "B33:H33", "B35:H35", "B37:H38", "A41:H43"]
        elif kind == "sampling":
            _put(sheet, "H2", data.get("certificate_no") or data.get("report_no"))
            _put(sheet, "A16", f"SAMPLES WERE TAKEN FROM HOLDS {text(data.get('holds_inspected'))}, WITH SEALS NUMBERS:")
            for r in (21, 23, 25, 27):
                _put(sheet, f"A{r}", "")
                _put(sheet, f"C{r}", "")
            boxes = ["B11:E11", "B13:E13", "G13:I13", "A16:I17", "F46:I46"]
            _put(sheet, "F19", "COMPOSITE")
            _put(sheet, "H19", "SEALS")
            for i in range(1, 11):
                r = 20 + ((i - 1) % 5) * 2
                left, right = ("A", "C") if i <= 5 else ("F", "H")
                _put(sheet, f"{left}{r}", f"HOLD {i}")
                _put(sheet, f"{right}{r}", data.get(f"hold_{i}_seal"))
                boxes.append(f"{right}{r}:{'E' if i <= 5 else 'I'}{r}")
        else:
            boxes = ["B4:E4", "B6:E6", "G6:I6", "A33:I34"]
            for i in range(1, 7):
                r = 15 + ((i - 1) % 3) * 6
                col = "C" if i <= 3 else "H"
                for j, part in enumerate(("fwd_escape", "fwd_aft_hatch", "aft_escape")):
                    _put(sheet, f"{col}{r+j}", data.get(f"hold_{i}_{part}"))
                    boxes.append(f"{col}{r+j}:{'E' if i <= 3 else 'I'}{r+j}")
        for address in boxes:
            _merge(sheet, address)
        for row in sheet:
            for cell in row:
                if cell.value in (None, ""):
                    continue
                region = next((m for m in sheet.merged_cells.ranges if cell.coordinate in m), None)
                # Leave static sentence fragments in the template's flowing layout.
                dynamic = cell.coordinate in CELL_FIELDS[kind] or region or cell.coordinate in ("A16", "B37")
                if not dynamic:
                    continue
                first, last = (region.min_col, region.max_col) if region else (cell.column, cell.column)
                width = sum(((sheet.column_dimensions.get(get_column_letter(c)).width
                              if sheet.column_dimensions.get(get_column_letter(c)) else sheet.sheet_format.defaultColWidth or 8.43) * 7 + 5) * .75
                            for c in range(first, last + 1)) - 10
                width = max(15, width)
                size = float(cell.font.sz or 11)
                font = "Helvetica-Bold" if cell.font.bold else "Helvetica"
                lines = sum(sum(max(1, ceil(stringWidth(s, font, size)/width)) for s in simpleSplit(p, font, size, width)) or 1
                            for p in text(cell.value).split("\n"))
                rows = range(region.min_row, region.max_row + 1) if region else [cell.row]
                current = sum(sheet.row_dimensions[r].height or 15 for r in rows)
                extra = max(0, lines * size * 1.4 + 8 - current) / len(rows)
                for r in rows:
                    sheet.row_dimensions[r].height = (sheet.row_dimensions[r].height or 15) + extra
                a = copy(cell.alignment)
                a.wrap_text, a.shrink_to_fit, a.vertical = True, False, "top"
                cell.alignment = a
        for ws in book:
            ws.sheet_state = "visible" if ws == sheet else "hidden"
        book.active = book.index(sheet)
        sheet.print_area = area
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.page_setup.scale = None
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_view.showGridLines = False
        book.save(target)
    finally:
        book.close()


def word(kind, data, target):
    from docx import Document
    from services.draft_survey_report_formatting import replace_placeholders
    from services.document_branding import apply_mci_docx_branding
    doc = Document(TEMPLATES / f"{kind}_certificate_template.docx")
    values = {"{" + key + "}": text(data.get(key)) for key in set(TYPES[kind]["fields"]) | set(data)}
    def paragraphs(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)
    for container in [doc] + [part for section in doc.sections for part in (section.header, section.footer)]:
        for paragraph in paragraphs(container):
            if kind == "lashing":
                original = "".join(run.text for run in paragraph.runs)
                start = original.find("THIRTENN ")
                if start >= 0:
                    end, offset = start + len("THIRTENN "), 0
                    for run in paragraph.runs:
                        value = run.text
                        right = offset + len(value)
                        if right > start and offset < end:
                            run.text = value[:max(0, start-offset)] + value[max(0, end-offset):]
                        offset = right
            replace_placeholders(paragraph, values)
    apply_mci_docx_branding(doc, data)
    doc.save(target)


def build(kind, data, format):
    spec = TYPES[kind]
    with tempfile.TemporaryDirectory(prefix="som_certificate_web_") as directory:
        folder = Path(directory)
        source = folder / ("certificate.docx" if spec["document"] == "word" else "certificate.xlsx")
        (word if spec["document"] == "word" else excel)(kind, data, source)
        return (convert(source, folder) if format == "pdf" else source).read_bytes()
