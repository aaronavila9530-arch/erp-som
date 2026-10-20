"""Print geometry specific to the Container form; never resize shared templates."""
from copy import copy
from datetime import datetime
from math import ceil

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import RowBreak, ColBreak
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.lib.utils import simpleSplit


def restore_template_layout(source):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from services.container_report_excel_service import TEMPLATE_PATH

    book, template = load_workbook(source), load_workbook(TEMPLATE_PATH)
    try:
        sheet, original = book.active, template.active
        sheet.column_dimensions.clear()
        for key, dimension in original.column_dimensions.items():
            sheet.column_dimensions[key] = copy(dimension)
        for row in original:
            for cell in row:
                target = sheet[cell.coordinate]
                if target.__class__.__name__ != 'MergedCell':
                    target.alignment = copy(cell.alignment)
                    target.font = copy(cell.font)
        for key, dimension in original.row_dimensions.items():
            sheet.row_dimensions[key].height = dimension.height
        sheet['D8'] = sheet['C8'].value
        sheet['C8'] = None
        sheet.merge_cells('AD3:AJ3')
        sheet['AD3'].alignment = Alignment(horizontal='center')
        link, value = sheet['B42'].hyperlink, sheet['B42'].value
        sheet['B42'] = None
        sheet.merge_cells('A42:AJ42')
        sheet['A42'] = value
        if link:
            sheet['A42'].hyperlink = link.target
        sheet['AF53'] = sheet['AB53'].value
        sheet['AB53'] = None
        for region in ('AF53:AH53', 'D9:H9'):
            sheet.unmerge_cells(region)
        sheet.merge_cells('AF53:AJ53')
        sheet.merge_cells('D9:R9')
        fit_container_sheet(sheet)
        book.save(source)
    finally:
        book.close()
        template.close()


def fit_container_sheet(sheet):
    # Labels previously relied on Excel overflowing into empty cells. Explicit
    # boxes make wrapping predictable in Excel and LibreOffice alike.
    boxes = ['A5:B5', 'A6:B6', 'A7:B7', 'A8:C8', 'A9:C9',
             'Y6:AB6', 'S8:W8', 'S9:W9',
             'A21:AJ21', 'A24:AJ24', 'A26:AJ26']
    for row in (12, 13):
        boxes += [f'{a}{row}:{b}{row}' for a, b in
                  [('B', 'D'), ('F', 'H'), ('J', 'M'), ('O', 'P'),
                   ('R', 'V'), ('X', 'AA'), ('AC', 'AF'), ('AH', 'AJ')]]
    for row in range(44, 53):
        boxes.append(f'K{row}:R{row}')
    for row in range(17, 20):
        boxes += [f'V{row}:X{row}', f'Z{row}:AA{row}', f'AC{row}:AD{row}']
    for address in boxes:
        if address not in sheet.merged_cells:
            sheet.merge_cells(address)

    sheet.sheet_format.zeroHeight = False
    for row in sheet:
        for cell in row:
            if cell.value in (None, ''):
                continue
            if isinstance(cell.value, (int, float)) and cell.coordinate != 'A15':
                cell.number_format = 'General'
            elif isinstance(cell.value, str) and cell.data_type != 'f':
                cell.number_format = '@'
            region = next((r for r in sheet.merged_cells.ranges if cell.coordinate in r), None)
            c1, c2 = (region.min_col, region.max_col) if region else (cell.column, cell.column)
            r1, r2 = (region.min_row, region.max_row) if region else (cell.row, cell.row)
            width = 0
            for col in range(c1, c2 + 1):
                dim = sheet.column_dimensions.get(get_column_letter(col))
                width += ((dim.width if dim else sheet.sheet_format.defaultColWidth or 4) * 7 + 5) * .75
            width = max(10, width - 10)
            size = float(cell.font.sz or 10)
            font = 'Helvetica-Bold' if cell.font.bold else 'Helvetica'
            value = cell.value.strftime('%d/%m/%Y %H:%M') if isinstance(cell.value, datetime) else str(cell.value)
            count = 0
            for paragraph in value.split('\n'):
                lines = simpleSplit(paragraph, font, size, width)
                count += sum(max(1, ceil(stringWidth(line, font, size) / width)) for line in lines) or 1
            needed = count * size * 1.45 + 8
            rows = range(r1, r2 + 1)
            available = sum(sheet.row_dimensions[r].height or 15 for r in rows)
            if needed > available:
                extra = (needed - available) / len(rows)
                for r in rows:
                    sheet.row_dimensions[r].height = (sheet.row_dimensions[r].height or 15) + extra
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.shrink_to_fit = False
            alignment.vertical = 'top'
            cell.alignment = alignment

    sheet.print_area = 'A1:AJ60'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = None
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.usePrinterDefaults = False
    sheet.page_margins.left = sheet.page_margins.right = .3
    sheet.page_margins.top = sheet.page_margins.bottom = .4
    sheet.page_margins.header = sheet.page_margins.footer = .15
    sheet.print_options.horizontalCentered = True
    sheet.sheet_view.showGridLines = False
    sheet.row_breaks = RowBreak()
    sheet.col_breaks = ColBreak()
    sheet.oddFooter.right.text = 'Page &P of &N'
    sheet.oddFooter.right.size = 9
