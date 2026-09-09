from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

try:
    from branding import is_mci_context, logo_asset
except ModuleNotFoundError:
    from backend_api.branding import is_mci_context, logo_asset


def _remove_paragraph_drawings(paragraph) -> None:
    for drawing in list(paragraph._element.xpath(".//w:drawing")):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)


def _remove_header_images(header) -> None:
    for paragraph in header.paragraphs:
        _remove_paragraph_drawings(paragraph)
    for table in header.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _remove_paragraph_drawings(paragraph)


def apply_mci_docx_branding(doc, data: dict | None = None) -> None:
    if not is_mci_context(data or {}):
        return

    logo = logo_asset(data or {})
    if not logo:
        return

    for section in doc.sections:
        header = section.header
        _remove_header_images(header)
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run().add_picture(logo, width=Inches(1.45))

