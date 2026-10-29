"""MSL's original Logistics templates, filled in isolated per-request files."""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import tempfile
from docx import Document
from docx.shared import Pt
from services.draft_survey_report_formatting import replace_placeholders
from services.container_web_export import convert
from services.logistics_web_fields import SPEC
from services.template_autofit import apply_docx_autofit

TEMPLATES = Path(__file__).resolve().parents[1] / 'templates'
MONTHS = ('Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec')

def format_date(value):
    if not value:return ''
    if isinstance(value, datetime):value=value.date()
    if not isinstance(value, date):value=date.fromisoformat(str(value)[:10])
    return f'{MONTHS[value.month-1]} {value.day:02d} {value.year}'

def word(data, target, presentation=False):
    if data.get('company_code') != 'MSL-CR':
        raise ValueError('MSL Logistics template requires MSL context')
    doc=Document(TEMPLATES/('presentation_truck_supervision.docx' if presentation else 'vessel_truck_supervision_aligned_dates.docx'))
    values={k:str(data.get(k) or '') for k in SPEC['fields']}
    for key in ('report_date','arrival_date','inspection_date','supervision_completed_date'):
        values[key]=format_date(data.get(key))
    values.update(ship_grt=values['grt'],ship_nrt=values['nrt'],requested_by=values['customer'],
                  sampling_start_time=values['report_date'])
    replacements={'{'+k+'}':v for k,v in values.items()}
    def replace(container):
        for p in container.paragraphs:replace_placeholders(p,replacements)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:replace(cell)
    replace(doc)
    for section in doc.sections:
        replace(section.header);replace(section.footer)
    if presentation:
        for p in doc.paragraphs:
            if p.text.strip() in (values['vessel_name'], values['customer']) or p.text.strip().startswith('PUERTO'):
                p.paragraph_format.right_indent=Pt(70)
                for run in p.runs:
                    if run.font.size and run.font.size>Pt(18 if p.text.strip().startswith('PUERTO') else 28):
                        run.font.size=Pt(18 if p.text.strip().startswith('PUERTO') else 28)
    else:
        paragraphs=doc.paragraphs
        for i,p in enumerate(paragraphs):
            title=p.text.strip()
            if title and (title.isupper() and len(title)<90 or title.startswith(('5.1.1','5.1.2'))):
                p.paragraph_format.keep_with_next=True
                for following in paragraphs[i+1:]:
                    if following.text.strip():break
                    following.paragraph_format.keep_with_next=True
    apply_docx_autofit(doc)
    doc.save(target)

def build(data, format):
    with tempfile.TemporaryDirectory(prefix='som_logistics_web_') as folder:
        folder=Path(folder)
        if format in ('presentation','complete'):
            cover=folder/'presentation.docx';word(data,cover,True)
            cover_pdf=convert(cover,folder)
            if format=='presentation':return cover_pdf.read_bytes()
        report=folder/'report.docx';word(data,report)
        if format=='word':return report.read_bytes()
        report_pdf=convert(report,folder)
        if format=='pdf':return report_pdf.read_bytes()
        if format!='complete':raise ValueError('Unsupported format')
        from pypdf import PdfWriter
        with PdfWriter() as writer:
            writer.append(cover_pdf);writer.append(report_pdf)
            output=BytesIO();writer.write(output);return output.getvalue()
