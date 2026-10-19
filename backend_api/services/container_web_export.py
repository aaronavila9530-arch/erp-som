"""Web exports reuse desktop templates, with isolated conversion and cleanup."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile


def convert(source, folder):
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        raise RuntimeError("LibreOffice unavailable")
    profile = folder / (source.stem + "-profile")
    subprocess.run([binary, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo",
                    "--convert-to", "pdf", "--outdir", str(folder), str(source)],
                   check=True, timeout=75, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    result = folder / (source.stem + ".pdf")
    if not result.is_file():
        raise RuntimeError("PDF conversion produced no file")
    return result


def presentation(report, customer, folder):
    from docx import Document
    from services.presentation_doc_service import TEMPLATE_PATH, _replace_in_paragraphs, _replace_in_tables
    from services.template_autofit import apply_docx_autofit
    from services.document_branding import apply_mci_docx_branding
    raw = report.get("init_inspection_datetime")
    try:
        date = (raw if isinstance(raw, datetime) else datetime.fromisoformat(str(raw))).strftime("%d-%m-%Y")
    except (ValueError, TypeError):
        date = str(raw or "").split(" ")[0]
    data = {"cert_no": report.get("linked_report_number"), "container": report.get("report_no"),
            "to": customer, "place": report.get("inspection_place"), "date": date,
            "company_code": report.get("company_code")}
    placeholders = {"{{" + key.upper() + "}}": str(value or "") for key, value in data.items()}
    doc = Document(TEMPLATE_PATH)
    _replace_in_paragraphs(doc.paragraphs, placeholders)
    _replace_in_tables(doc.tables, placeholders)
    for section in doc.sections:
        for part in (section.header, section.footer):
            _replace_in_paragraphs(part.paragraphs, placeholders)
            _replace_in_tables(part.tables, placeholders)
    apply_mci_docx_branding(doc, data)
    apply_docx_autofit(doc)
    source = folder / "presentation.docx"
    doc.save(source)
    return convert(source, folder)


def build(report, customer, kind):
    from services.container_report_excel_service import generate_container_report_excel
    # Treat user-entered narrative as text, never as an executable Excel formula.
    safe = {key: "'" + value if isinstance(value, str) and value.startswith("=") else value
            for key, value in report.items()}
    with tempfile.TemporaryDirectory(prefix="som_container_web_") as directory:
        folder = Path(directory)
        report_pdf = None
        if kind in ("excel", "pdf", "unified"):
            generated = Path(generate_container_report_excel(safe))
            try:
                source = folder / "report.xlsx"
                shutil.copyfile(generated, source)
            finally:
                generated.unlink(missing_ok=True)
            if kind == "excel":
                return source.read_bytes()
            report_pdf = convert(source, folder)
            if kind == "pdf":
                return report_pdf.read_bytes()
        cover = presentation(report, customer, folder)
        if kind == "presentation":
            return cover.read_bytes()
        from pypdf import PdfWriter
        output = BytesIO()
        with PdfWriter() as writer:
            writer.append(str(cover))
            writer.append(str(report_pdf))
            writer.write(output)
        return output.getvalue()
