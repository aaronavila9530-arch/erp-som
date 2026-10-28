"""Continuation pages for text that cannot safely fit a printed Excel cell."""
from copy import copy
from math import ceil

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase.pdfmetrics import stringWidth


def _height(value, width, size):
    lines = sum(sum(max(1, ceil(stringWidth(line, 'Helvetica', size) / width))
                    for line in simpleSplit(part, 'Helvetica', size, width)) or 1
                for part in str(value).split('\n'))
    return lines * size * 1.5 + 12


def continue_long_cells(book, sheet, addresses, reference=''):
    """Keep the original form; link oversized values to readable extra sheets."""
    moved = []
    for address, label in addresses.items():
        cell = sheet[address]
        if not isinstance(cell.value, str) or not cell.value or cell.data_type == 'f':
            continue
        region = next((r for r in sheet.merged_cells.ranges if address in r), None)
        first, last = (region.min_col, region.max_col) if region else (cell.column, cell.column)
        width = 0
        for col in range(first, last + 1):
            dim = sheet.column_dimensions.get(get_column_letter(col))
            width += ((dim.width if dim else sheet.sheet_format.defaultColWidth or 8.43) * 7 + 5) * .75
        width = max(10, width - 12)
        size = float(cell.font.sz or 11)
        if _height(cell.value, width, size) <= 240:
            continue
        moved.append((address, label, cell.value, copy(cell.font)))
        cell.value = f'See continuation: {label} ({address}).'
        cell.data_type = 's'
        for row in range(region.min_row, region.max_row + 1) if region else [cell.row]:
            if (sheet.row_dimensions[row].height or 15) > 120:
                sheet.row_dimensions[row].height = 120
    if not moved:
        return None
    extra = book.create_sheet('Continuation')
    extra.sheet_view.showGridLines = False
    extra.column_dimensions['A'].width = 95
    extra.page_setup.orientation = sheet.page_setup.orientation or 'portrait'
    extra.page_setup.paperSize = sheet.page_setup.paperSize or extra.PAPERSIZE_A4
    extra.page_setup.fitToWidth, extra.page_setup.fitToHeight = 1, 0
    extra.page_setup.scale = None
    extra.sheet_properties.pageSetUpPr.fitToPage = True
    extra.page_margins = copy(sheet.page_margins)
    extra['A1'] = 'CONTINUATION - ' + str(reference or sheet.title)
    extra['A1'].data_type = 's'
    extra['A1'].font = Font(name='Arial', size=11, bold=True)
    extra['A1'].alignment = Alignment(wrap_text=True)
    extra.row_dimensions[1].height = _height(extra['A1'].value, 490, 11)
    extra.print_title_rows = '1:1'
    extra.oddHeader.center.text = ('Continuation - ' + str(reference or sheet.title)).replace('&', '&&')
    extra.oddHeader.center.size = 9
    extra.oddFooter.right.text = 'Page &P of &N'
    row = 3
    for address, label, value, font in moved:
        extra.cell(row, 1, f'{label} ({address})').font = Font(name=font.name or 'Arial', size=11, bold=True)
        row += 1
        # Exact slices preserve every character, including long URLs and newlines.
        # Bounded chunks stay below Excel's row-height limit and a printed page.
        start = 0
        while start < len(value):
            end = min(len(value), start + 700)
            if end < len(value):
                space = value.rfind(' ', start, end)
                if space > start:
                    end = space + 1
            chunk = value[start:end]
            while _height(chunk, 490, 11) > 240 and end > start + 1:
                end = start + max(1, (end - start) // 2)
                chunk = value[start:end]
            target = extra.cell(row, 1, chunk)
            target.data_type = 's'
            target.font = Font(name=font.name or 'Arial', size=11)
            target.alignment = Alignment(wrap_text=True, vertical='top', shrink_to_fit=False)
            extra.row_dimensions[row].height = _height(chunk, 490, 11)
            row += 1
            start = end
        row += 1
    extra.print_area = f'A1:A{row-1}'
    return extra
