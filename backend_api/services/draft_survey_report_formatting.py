"""Presentation dates and field-safe replacements for Draft Survey reports."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def normalize_draft_layout(doc, presentation=False):
    """Fix only Draft templates' floating website and space-aligned footer."""
    roots = [doc._element]
    seen = set()
    for section in doc.sections:
        for part in (section.header, section.footer):
            if id(part._element) not in seen:
                roots.append(part._element)
                seen.add(id(part._element))
    # The website's fixed-height textbox clips in LibreOffice. Use normal flow,
    # replacing the entire AlternateContent so its VML fallback cannot reappear.
    from docx.text.paragraph import Paragraph
    for root in roots:
        for node in list(root.xpath('.//w:p')):
            boxes = node.xpath('./w:r/*[local-name()="AlternateContent"]')
            for box in boxes:
                texts = [t.text or '' for t in box.iter(qn('w:t'))]
                website = next((t.strip() for t in texts
                                if t.strip().startswith('www.mslogisticsgroup.')), None)
                if not website:
                    continue
                box.getparent().remove(box)
                p = Paragraph(node, doc)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.right_indent = Pt(0)
                p.paragraph_format.first_line_indent = Pt(0)
                run = p.add_run(website)
                run.font.size = Pt(12 if presentation else 10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 255)

    for section in doc.sections:
        for p in list(section.footer.paragraphs):
            if not p._p.xpath('.//w:instrText') or 'Page ' not in p.text:
                continue
            label = p.text.split('Page ', 1)[0].strip()
            label = re.sub(r'\s+', ' ', label)
            label = re.sub(r'\bMV\s+MV\b', 'MV', label, flags=re.I)
            p.clear()
            pf = p.paragraph_format
            pf.tab_stops.clear_all()
            pf.left_indent = pf.right_indent = pf.first_line_indent = Pt(0)
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(label).font.size = Pt(7)
            p.add_run().add_break()
            p.add_run('Page\u00a0').font.size = Pt(8)
            for field, suffix in (('PAGE', '\u00a0of\u00a0'), ('NUMPAGES', '')):
                node = OxmlElement('w:fldSimple')
                node.set(qn('w:instr'), field)
                run = OxmlElement('w:r')
                text = OxmlElement('w:t')
                text.text = '1'
                run.append(text)
                node.append(run)
                p._p.append(node)
                if suffix:
                    p.add_run(suffix).font.size = Pt(8)

        if presentation:
            # The template background exceeded the page height and was anchored
            # relative to a paragraph. Fit it proportionally to the actual page.
            for anchor in section.header._element.xpath('.//wp:anchor'):
                if not list(anchor.iter(qn('pic:pic'))):
                    continue
                extent = anchor.find(qn('wp:extent'))
                width, height = int(extent.get('cx')), int(extent.get('cy'))
                if width < section.page_width * .8:
                    continue
                scale = min(section.page_width / width, section.page_height / height)
                width, height = int(width * scale), int(height * scale)
                for target in [extent] + list(anchor.iter(qn('a:ext'))):
                    if target is not extent and target.get('cx') is None:
                        continue
                    target.set('cx', str(width))
                    target.set('cy', str(height))
                for axis, offset in (('H', (section.page_width - width) // 2),
                                     ('V', (section.page_height - height) // 2)):
                    pos = anchor.find(qn('wp:position' + axis))
                    pos.set('relativeFrom', 'page')
                    for child in list(pos):
                        pos.remove(child)
                    child = OxmlElement('wp:posOffset')
                    child.text = str(offset)
                    pos.append(child)


def refresh_draft_difference(data):
    result = dict(data or {})
    def number(value):
        text = str(value or '').strip().replace(' ', '').replace('\u00a0', '')
        if ',' in text and '.' in text:
            text = text.replace('.', '').replace(',', '.') if text.rfind(',') > text.rfind('.') else text.replace(',', '')
        else:
            text = text.replace(',', '.')
        return Decimal(text)
    try:
        draft = number(result.get('word_draft_figures'))
        bl = number(result.get('word_bl_figures'))
        if draft.is_finite() and bl.is_finite():
            diff = draft - bl
            result['word_difference'] = f'{diff:,.2f}'
            result['word_percentage'] = f'{(diff / bl * 100) if bl else Decimal(0):.4f}'
    except (InvalidOperation, ValueError):
        pass
    return result


def report_date(data):
    for key in ('word_final_draft_date', 'word_final_draft',
                'word_completed_date', 'word_completed'):
        value = data.get(key)
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = re.sub(r'\s+LT\.?$', '', str(value or '').strip(), flags=re.I)
        if not text:
            continue
        try:
            return datetime.fromisoformat(text.replace('Z', '+00:00')).date()
        except ValueError:
            pass
        for fmt in ('%d/%m/%Y', '%d/%m/%Y %H:%M', '%d-%m-%Y',
                    '%d-%m-%Y %H:%M', '%b %d %Y', '%b %d %Y %H:%M',
                    '%B %d, %Y', '%B %d, %Y %H:%M'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def report_place_line(data):
    port = str(data.get('word_port') or '').strip()
    country = str(data.get('word_country') or '').strip()
    value = report_date(data)
    place = ' - '.join(part for part in (port, country) if part)
    if value:
        months = ('January', 'February', 'March', 'April', 'May', 'June',
                  'July', 'August', 'September', 'October', 'November', 'December')
        formatted = f'{months[value.month - 1]} {value.day:02d}, {value.year}'
        return f'{place}, {formatted}' if place else formatted
    return place


def replace_placeholders(paragraph, values):
    """Replace only placeholder runs; never redistribute cached PAGE field text."""
    runs = paragraph.runs
    text = ''.join(run.text for run in runs)
    for match in reversed(list(re.finditer(r'\{[^{}]+\}', text))):
        start, end = match.span()
        value = str(values.get(match.group(), '') or '')
        offset = 0
        inserted = False
        for run in runs:
            original = run.text
            right = offset + len(original)
            if right > start and offset < end:
                left_part = original[:max(0, start - offset)]
                right_part = original[max(0, end - offset):]
                run.text = left_part + (value if not inserted else '') + right_part
                inserted = True
            offset = right
