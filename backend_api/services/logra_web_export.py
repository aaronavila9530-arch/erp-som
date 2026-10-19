"""Memory-only LOGRA exports; never imports desktop UI or mutates reports."""
from io import BytesIO
from xml.sax.saxutils import escape


def sections(payload, executive=None):
    report = payload['report']
    yield 'title', report.get('title') or 'LOGRA'
    if executive is not None:
        yield 'text', 'Borrador generado con IA. Requiere revision profesional.'
        for line in executive.splitlines():
            if line.strip(): yield 'text', line
        return
    yield 'text', f"Estado: {report.get('status') or '-'} | Version: {report.get('version') or 1} | Autor: {report.get('created_by') or '-'}"
    if report.get('agenda_items'):
        yield 'heading', 'Agenda'
        for item in report['agenda_items']:
            yield 'subheading', f"{item.get('date_iso') or item.get('date') or item.get('date_long') or ''}  {item.get('start_time') or ''} - {item.get('end_time') or ''}"
            for label, key in [('Tema','topic'),('Lugar','place'),('Persona','person'),('Empresa','company'),('Empresa / cargo','company_role'),('Telefono','phone'),('Prioridad','priority'),('Estado','status')]:
                if item.get(key): yield 'text', f'{label}: {item[key]}'
    if report.get('agenda_notes'):
        yield 'heading', 'Notas de agenda'
        yield 'text', report['agenda_notes']
    current = None
    for answer in payload.get('answers', []):
        group = (answer['form_slug'], answer['section'])
        if group != current:
            yield 'heading', answer.get('form_title') or answer['form_slug']
            yield 'subheading', 'Preguntas criticas' if answer['section'] == 'critical_questions' else 'Preguntas detalladas'
            current = group
        yield 'question', f"{answer['item_key']}. {answer.get('question_text') or ''}"
        for bullet in answer.get('bullets') or ['Sin respuesta registrada.']:
            yield 'text', str(bullet)
        names = [a['original_filename'] for a in payload.get('attachments', [])
                 if (a['form_slug'], a['section'], a['item_key']) == (*group, answer['item_key'])]
        if names: yield 'text', 'Adjuntos: ' + ', '.join(names)


def build(payload, kind, executive=None):
    output = BytesIO()
    if kind == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        book = Workbook()
        sheet = book.active
        sheet.title = 'Respuestas'
        def safe(value):
            value = str(value or '')
            return "'" + value if value.startswith(('=', '+', '-', '@')) else value
        sheet.append(['Cuestionario', 'Seccion', 'Pregunta', 'Respuesta'])
        for a in payload.get('answers', []):
            sheet.append([safe(a.get('form_title')), safe(a['section']), safe(a.get('question_text')), safe('\n'.join(a.get('bullets') or []))])
        agenda = book.create_sheet('Agenda')
        keys = ['date_iso','start_time','end_time','place','person','company','phone','topic','priority','status']
        agenda.append(['Fecha','Inicio','Fin','Lugar','Persona','Empresa','Telefono','Tema','Prioridad','Estado'])
        for item in payload['report'].get('agenda_items') or []:
            agenda.append([safe(item.get(key) or (item.get('date') if key == 'date_iso' else '')) for key in keys])
        for ws in book:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True,color='FFFFFF')
                cell.fill = PatternFill('solid',fgColor='126E74')
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment = Alignment(wrap_text=True,vertical='top')
            for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 35
        book.save(output)
        return output.getvalue()
    if kind == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt
        doc = Document()
        doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(.65)
        doc.styles['Normal'].font.name = 'Arial'
        doc.styles['Normal'].font.size = Pt(10)
        for style, value in sections(payload, executive):
            if style == 'title': doc.add_heading(value, 0)
            elif style == 'heading': doc.add_heading(value, 1)
            elif style == 'subheading': doc.add_heading(value, 2)
            else:
                paragraph = doc.add_paragraph(value)
                if style == 'question':
                    paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs: run.bold = True
        doc.save(output)
        return output.getvalue()
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    styles = getSampleStyleSheet()
    styles['Normal'].spaceAfter = 6
    styles.add(ParagraphStyle('Question',parent=styles['Normal'],fontName='Helvetica-Bold',spaceBefore=10,spaceAfter=5,keepWithNext=True))
    mapping = {'title':'Title','heading':'Heading1','subheading':'Heading2','question':'Question','text':'Normal'}
    story = []
    for style, value in sections(payload, executive):
        story.append(Paragraph(escape(str(value)).replace('\n','<br/>'), styles[mapping[style]]))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica',8)
        canvas.drawString(42,24,'SOM | LOGRA')
        canvas.drawRightString(A4[0]-42,24,str(doc.page))
        canvas.restoreState()
    SimpleDocTemplate(output,pagesize=A4,leftMargin=42,rightMargin=42,topMargin=42,bottomMargin=42).build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()
