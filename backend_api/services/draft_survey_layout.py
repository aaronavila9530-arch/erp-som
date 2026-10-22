"""Flow-based Draft Survey layouts shared by Word and PDF exports."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _paragraph(p, size=10):
    f = p.paragraph_format
    f.left_indent = f.right_indent = f.first_line_indent = Pt(0)
    f.line_spacing = 1.12
    f.space_before = Pt(0)
    f.space_after = Pt(6)
    f.widow_control = True
    f.keep_together = False
    f.keep_with_next = False
    f.tab_stops.clear_all()
    for node in list(p._p.xpath('./w:pPr/w:numPr')):
        node.getparent().remove(node)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.underline = False


def _field(p, instruction):
    node = OxmlElement('w:fldSimple')
    node.set(qn('w:instr'), instruction)
    p._p.append(node)


def _frame(doc, data):
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    doc.settings.odd_and_even_pages_header_footer = False
    for section in doc.sections:
        section.different_first_page_header_footer = False
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(.8)
        section.top_margin, section.bottom_margin = Inches(1), Inches(.95)
        section.header_distance = section.footer_distance = Inches(.3)
        for part in (section.header, section.footer):
            for child in list(part._element):
                part._element.remove(child)
        header = section.header.add_paragraph()
        try:
            from branding import logo_asset, footer_text
        except ModuleNotFoundError:
            from backend_api.branding import logo_asset, footer_text
        logo = logo_asset(data)
        if logo and Path(logo).exists():
            header.add_run().add_picture(str(logo), height=Inches(.42))
        header.add_run('   DRAFT SURVEY | ' + str(data.get('draft_report_number', '')))
        _paragraph(header, 8)
        footer = section.footer.add_paragraph()
        footer.add_run(footer_text(data) + '\n')
        footer.add_run(str(data.get('draft_report_number', '')) + ' | ' +
                       str(data.get('word_vessel', '')))
        footer.add_run('\nPage\u00a0')
        _field(footer, 'PAGE')
        footer.add_run('\u00a0of\u00a0')
        _field(footer, 'NUMPAGES')
        _paragraph(footer, 8)
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    settings = doc.settings.element
    if settings.find(qn('w:updateFields')) is None:
        update = OxmlElement('w:updateFields')
        update.set(qn('w:val'), 'true')
        settings.append(update)


def _table(doc, rows, before=None):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    widths = (Inches(3.1), Inches(3.8))
    for col, width in zip(table.columns, widths):
        col.width = width
    for label, value in rows:
        row = table.add_row()
        for cell, text, width in zip(row.cells, (label, value), widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            _paragraph(cell.paragraphs[0])
            cell.paragraphs[0].paragraph_format.space_after = Pt(4)
            cell.paragraphs[0].paragraph_format.space_before = Pt(4)
        row.cells[0].paragraphs[0].runs[0].bold = True
        cant_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant_split)
    if before is not None:
        before.addprevious(table._tbl)
    return table


def harden_written_report(doc, data):
    """Preserve template prose; replace its space/tab aligned data blocks."""
    _frame(doc, data)
    groups = []
    current = []
    headings = {'INTRODUCTION', 'VESSEL PARTICULARS', 'EXTRACT TIME SHEET',
                'THE GOODS', 'WEIGTH CONTROL', 'CARGO QUANTITY DISCHARGED BY DRAFT SURVEY',
                'ENCLOSURES'}
    labels = {'Survey requested by:', 'On behalf of:', 'Master of the ship:',
              'Chief Officer:', 'Name', 'Port Of Registry / Flag', 'GRT / NET',
              'Year Built', 'IMO Number', 'Vessel Arrived at Sea buoy',
              'N.O.R Tendered', 'All Fast', 'Initial Draft Survey',
              'Commenced Discharge', 'Completed Discharge', 'Final Draft Survey',
              'Draft Survey Figure', 'B/L Figures', 'Difference', 'Percentage',
              'Shore Scale Figures'}
    for p in list(doc.paragraphs):
        text = p.text.strip()
        # The template supplies a vessel prefix even when the saved name has it.
        cleaned = re.sub(r'\bM/?V\s+[\u201c\"]?M/?V\s+', 'MV ', text)
        if cleaned != text:
            cleaned = cleaned.replace('\u201d', '')
            p.text = cleaned
            text = cleaned
        if not text:
            if current:
                groups.append(current)
                current = []
            p._p.getparent().remove(p._p)
            continue
        # The source template encodes record columns as tabs or repeated spaces.
        parts = re.split(r'\t+| {2,}', text, maxsplit=1)
        is_data = len(parts) == 2 and parts[0].strip() in labels
        if is_data:
            current.append((p, parts[0].strip(), re.sub(r'\s+', ' ', parts[1]).strip()))
        else:
            if current:
                groups.append(current)
                current = []
            _paragraph(p)
            if text in headings or text in ('Purpose', 'Instructions', 'Representatives') or text.startswith('Total Discharged'):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(10)
                for run in p.runs:
                    run.bold = True
            if text in ('EXTRACT TIME SHEET', 'CARGO QUANTITY DISCHARGED BY DRAFT SURVEY'):
                p.paragraph_format.page_break_before = True
            if text == 'DRAFT SURVEY REPORT':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(16)
                    run.bold = True
    if current:
        groups.append(current)
    for group in groups:
        _table(doc, [(label, value) for _, label, value in group], group[0][0]._p)
        for p, _, _ in group:
            p._p.getparent().remove(p._p)


def build_presentation(data, place_line):
    doc = Document()
    _frame(doc, data)
    title = doc.add_paragraph('DRAFT SURVEY REPORT')
    _paragraph(title, 22)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    title.paragraph_format.space_after = Pt(24)
    title.runs[0].bold = True
    vessel = doc.add_paragraph(str(data.get('word_vessel') or ''))
    _paragraph(vessel, 20)
    vessel.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vessel.paragraph_format.space_after = Pt(36)
    vessel.runs[0].bold = True
    _table(doc, [('Certificate', str(data.get('draft_report_number') or '')),
                 ('Gross register tonnage', str(data.get('word_grt') or '')),
                 ('Net register tonnage', str(data.get('word_nrt') or '')),
                 ('Survey requested by', str(data.get('word_survey_requested_by') or '')),
                 ('Place and date', place_line)])
    return doc
