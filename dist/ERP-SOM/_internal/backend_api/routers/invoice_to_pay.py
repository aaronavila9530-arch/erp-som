from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    Header,
    UploadFile,
    File,
    Form
)
from fastapi.responses import StreamingResponse
from psycopg2.extras import RealDictCursor
from psycopg2.extras import Json
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import Optional
import io
import os
import shutil

from database import get_db
from rbac_service import has_permission
from services.finance_audit import actor_from_headers, audit_event, row_to_dict
from services.accounting_bank_rules import external_surveyor_settlement, resolve_itp_bank
from services.employee_payee_rules import (
    deactivate_employee_itp_obligations,
    is_employee_payee,
)
from services.tenanting import company_code as normalize_company_code


router = APIRouter(
    prefix="/invoice-to-pay",
    tags=["Finance - Invoice to Pay"]
)


def _ensure_company_column(cur):
    cur.execute("""
        ALTER TABLE payment_obligations
        ADD COLUMN IF NOT EXISTS company_code VARCHAR(30) NOT NULL DEFAULT 'MSL-CR'
    """)


def _money(value) -> Decimal:
    return Decimal(str(value or 0).replace(",", "")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _previous_period(period: str) -> str:
    year, month = [int(part) for part in str(period).split("-")[:2]]
    month -= 1
    if month == 0:
        year -= 1
        month = 12
    return f"{year:04d}-{month:02d}"


def _fortnight_due_date(period: str, fortnight: int) -> str:
    year, month = [int(part) for part in str(period).split("-")[:2]]
    return f"{year:04d}-{month:02d}-{'15' if int(fortnight or 1) == 1 else '30'}"


def _ensure_biweekly_schema(cur):
    cur.execute("""
        CREATE TABLE IF NOT EXISTS itp_biweekly_payment_batches (
            id BIGSERIAL PRIMARY KEY,
            company_code TEXT NOT NULL,
            period TEXT NOT NULL,
            fortnight INTEGER NOT NULL,
            created_by TEXT,
            created_at TIMESTAMP NOT NULL DEFAULT NOW(),
            status TEXT NOT NULL DEFAULT 'APPLIED'
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS itp_biweekly_payment_lines (
            id BIGSERIAL PRIMARY KEY,
            batch_id BIGINT REFERENCES itp_biweekly_payment_batches(id) ON DELETE CASCADE,
            company_code TEXT NOT NULL,
            category TEXT NOT NULL,
            beneficiary TEXT NOT NULL,
            amount NUMERIC(18,2) NOT NULL,
            currency TEXT NOT NULL DEFAULT 'CRC',
            amount_crc NUMERIC(18,2) NOT NULL DEFAULT 0,
            destination_account TEXT,
            bank_accounting_code TEXT NOT NULL,
            bank_accounting_name TEXT,
            bank_voucher TEXT NOT NULL,
            payment_date DATE NOT NULL,
            obligation_id BIGINT,
            reference TEXT,
            source TEXT,
            notes TEXT,
            accounting_entry_id INTEGER,
            created_at TIMESTAMP NOT NULL DEFAULT NOW()
        )
    """)
    cur.execute("ALTER TABLE payment_obligations ADD COLUMN IF NOT EXISTS last_payment_date DATE")
    cur.execute("ALTER TABLE payment_obligations ADD COLUMN IF NOT EXISTS payment_bank_account_code TEXT")
    cur.execute("ALTER TABLE payment_obligations ADD COLUMN IF NOT EXISTS payment_bank_account_name TEXT")


def _exchange_rate(cur, value_date: str) -> Decimal:
    cur.execute(
        """
        SELECT venta FROM tipo_cambio
        WHERE fecha <= %s
        ORDER BY fecha DESC
        LIMIT 1
        """,
        (value_date,),
    )
    row = cur.fetchone()
    return _money((row or {}).get("venta") if isinstance(row, dict) else (row[0] if row else 1))


def _debit_account_for(category: str):
    mapping = {
        "Planilla": ("2.1.02.07", "Salarios por pagar"),
        "CCSS": ("2.1.05.01", "Obligaciones patronales por pagar-CCSS"),
        "IVA": ("2.1.02.03", "Impuesto sobre valor agregado (IVA) por pagar"),
        "Tarjetas de credito": ("2.1.02.10", "Tarjeta corporativa BAC por pagar"),
        "Telefonia": ("500-001-001-023", "Telefonos"),
        "Viaticos": ("500-001-001-044", "Viaticos"),
        "Alquiler": ("500-001-001-045", "Alquileres"),
        "Internet": ("500-001-001-006", "Servicios Profesionales"),
        "Surveyors": ("2.1.01.01", "Cuentas por pagar-comerciales"),
    }
    return mapping.get(category, ("5.4", "Otros gastos"))

# ============================================================
# RBAC GUARD
# ============================================================
def require_permission(module: str, action: str):
    def checker(
        x_user_role: str = Header(..., alias="X-User-Role")
    ):
        if not has_permission(x_user_role, module, action):
            raise HTTPException(
                status_code=403,
                detail="No autorizado"
            )
    return checker

# ============================================================
# 🔁 SYNC SERVICIOS → PAYMENT OBLIGATIONS
# ============================================================
def _sync_servicios_to_itp(cur):
    """
    Sincroniza obligaciones desde servicios hacia Invoice To Pay.

    Reglas:
    - INSERTA si no existe
    - ACTUALIZA solo si:
        • origin = 'SERVICIOS'
        • status = 'PENDING' o 'PARTIAL'
        • el monto cambió
    - Respeta pagos parciales recalculando balance
    """

    # ============================================================
    # 1️⃣ INSERTAR HONORARIOS (SURVEYOR_FEE)
    # ============================================================
    cur.execute("""
        INSERT INTO payment_obligations (
            record_type,
            payee_type,
            payee_name,
            obligation_type,
            reference,
            vessel,
            country,
            operation,
            service_id,
            issue_date,
            due_date,
            currency,
            total,
            balance,
            status,
            origin,
            notes,
            created_at
        )
        SELECT
            'OBLIGATION',
            'SURVEYOR',
            s.surveyor,
            'SURVEYOR_FEE',
            s.consec,
            s.buque_contenedor,
            s.pais,
            s.operacion,
            s.consec,
            s.fecha_fin,
            (s.fecha_fin + INTERVAL '15 days'),
            'USD',
            s.honorarios,
            s.honorarios,
            'PENDING',
            'SERVICIOS',
            s.detalle,
            NOW()
        FROM servicios s
        WHERE
            s.surveyor IS NOT NULL
            AND s.honorarios IS NOT NULL
            AND s.honorarios > 0
            AND s.fecha_fin IS NOT NULL
            AND NOT EXISTS (
                SELECT 1
                FROM payment_obligations po
                WHERE po.service_id = s.consec
                  AND po.origin = 'SERVICIOS'
                  AND po.obligation_type = 'SURVEYOR_FEE'
            )
    """)

    # ============================================================
    # 2️⃣ INSERTAR COSTO TARJETAS (CARD_PROCESSING)
    # ============================================================
    cur.execute("""
        INSERT INTO payment_obligations (
            record_type,
            payee_type,
            payee_name,
            obligation_type,
            reference,
            vessel,
            country,
            operation,
            service_id,
            issue_date,
            due_date,
            currency,
            total,
            balance,
            status,
            origin,
            notes,
            created_at
        )
        SELECT
            'OBLIGATION',
            'SUPPLIER',
            'CARD PROCESSOR',
            'CARD_PROCESSING',
            s.consec,
            s.buque_contenedor,
            s.pais,
            s.operacion,
            s.consec,
            s.fecha_fin,
            (s.fecha_fin + INTERVAL '15 days'),
            'USD',
            s.costo_tarjetas,
            s.costo_tarjetas,
            'PENDING',
            'SERVICIOS',
            'Costo tarjetas - ' || COALESCE(s.detalle,''),
            NOW()
        FROM servicios s
        WHERE
            s.costo_tarjetas IS NOT NULL
            AND s.costo_tarjetas > 0
            AND s.fecha_fin IS NOT NULL
            AND NOT EXISTS (
                SELECT 1
                FROM payment_obligations po
                WHERE po.service_id = s.consec
                  AND po.origin = 'SERVICIOS'
                  AND po.obligation_type = 'CARD_PROCESSING'
            )
    """)

    # ============================================================
    # 3️⃣ ACTUALIZAR HONORARIOS MODIFICADOS
    # ============================================================
    cur.execute("""
        UPDATE payment_obligations po
        SET
            total = s.honorarios,
            balance = GREATEST(
                s.honorarios - (po.total - po.balance),
                0
            ),
            payee_name = COALESCE(s.surveyor, po.payee_name),
            vessel = COALESCE(s.buque_contenedor, po.vessel),
            country = COALESCE(s.pais, po.country),
            operation = COALESCE(s.operacion, po.operation),
            issue_date = COALESCE(s.fecha_fin, po.issue_date),
            due_date = COALESCE((s.fecha_fin + INTERVAL '15 days'), po.due_date),
            notes = COALESCE(s.detalle, po.notes),
            updated_at = NOW()
        FROM servicios s
        WHERE
            po.service_id = s.consec
            AND po.origin = 'SERVICIOS'
            AND po.obligation_type = 'SURVEYOR_FEE'
            AND s.honorarios IS NOT NULL
            AND s.honorarios > 0
            AND po.status IN ('PENDING', 'PARTIAL')
            AND po.total IS DISTINCT FROM s.honorarios
    """)

    # ============================================================
    # 4️⃣ ACTUALIZAR COSTO TARJETAS MODIFICADO
    # ============================================================
    cur.execute("""
        UPDATE payment_obligations po
        SET
            total = s.costo_tarjetas,
            balance = GREATEST(
                s.costo_tarjetas - (po.total - po.balance),
                0
            ),
            vessel = COALESCE(s.buque_contenedor, po.vessel),
            country = COALESCE(s.pais, po.country),
            operation = COALESCE(s.operacion, po.operation),
            issue_date = COALESCE(s.fecha_fin, po.issue_date),
            due_date = COALESCE((s.fecha_fin + INTERVAL '15 days'), po.due_date),
            notes = 'Costo tarjetas - ' || COALESCE(s.detalle,''),
            updated_at = NOW()
        FROM servicios s
        WHERE
            po.service_id = s.consec
            AND po.origin = 'SERVICIOS'
            AND po.obligation_type = 'CARD_PROCESSING'
            AND s.costo_tarjetas IS NOT NULL
            AND s.costo_tarjetas > 0
            AND po.status IN ('PENDING', 'PARTIAL')
            AND po.total IS DISTINCT FROM s.costo_tarjetas
    """)

    deactivate_employee_itp_obligations(cur)

@router.get("/search")
def search_invoice_to_pay(
    obligation_type: Optional[str] = Query(None),
    payee: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    issue_date_from: Optional[date] = Query(None),
    issue_date_to: Optional[date] = Query(None),
    payment_date_from: Optional[date] = Query(None),
    payment_date_to: Optional[date] = Query(None),
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    cur = conn.cursor(cursor_factory=RealDictCursor)
    company = normalize_company_code(header_value=x_company_code)
    _ensure_company_column(cur)

    # 🔁 Sync servicios → Invoice To Pay
    _sync_servicios_to_itp(cur)
    deactivate_employee_itp_obligations(cur)
    conn.commit()

    filters = ["COALESCE(active, TRUE) = TRUE", "company_code = %s"]
    params = [company]

    # =================
    # FILTRO POR ESTADO
    # =================
    if status:
        status = status.upper()

        if status == "ALL":
            # 🔥 Sin filtro: trae TODOS los registros
            pass
        else:
            filters.append("status = %s")
            params.append(status)

    # =================================
    # FILTRO POR TIPO DE OBLIGACIÓN
    # (basado en payee_type REAL)
    # =================================
    if obligation_type:
        if obligation_type.upper() == "SURVEYOR":
            filters.append("payee_type = 'SURVEYOR'")
        elif obligation_type.upper() == "SUPPLIER":
            filters.append("payee_type = 'SUPPLIER'")
        elif obligation_type.upper() == "MANUAL":
            filters.append("origin = 'MANUAL'")

    # =================
    # FILTRO BENEFICIARIO
    # =================
    if payee:
        filters.append("payee_name ILIKE %s")
        params.append(f"%{payee}%")

    # ================================
    # FILTROS POR RANGOS DE FECHA
    # ================================
    if issue_date_from:
        filters.append("issue_date >= %s")
        params.append(issue_date_from)

    if issue_date_to:
        filters.append("issue_date <= %s")
        params.append(issue_date_to)

    if payment_date_from:
        filters.append("last_payment_date >= %s")
        params.append(payment_date_from)

    if payment_date_to:
        filters.append("last_payment_date <= %s")
        params.append(payment_date_to)

    where_clause = ""
    if filters:
        where_clause = "WHERE " + " AND ".join(filters)

    # ========================================
    # SELECT FINAL CORRECTO
    # ========================================
    sql = f"""
        SELECT
            id,
            payee_name,

            -- ✅ COLUMNA Obligación = payee_type REAL
            payee_type AS obligation_type,

            -- ✅ COLUMNA Referencia (regla de negocio)
            CASE
                WHEN origin = 'SERVICIOS' THEN notes
                ELSE reference
            END AS referencia,

            vessel,
            country,
            operation,
            currency,
            total,
            balance,
            status,
            last_payment_date,
            issue_date,
            due_date,
            origin

        FROM payment_obligations
        {where_clause}
        ORDER BY issue_date DESC
    """

    try:
        cur.execute(sql, params)
        rows = cur.fetchall()
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"InvoiceToPay search error: {str(e)}"
        )

    return {"data": rows}


# ============================================================
# 2️⃣ KPIs — CONVERSIÓN CRC → USD (TC = 500)
# ============================================================
@router.get("/kpis")
def invoice_to_pay_kpis(conn=Depends(get_db)):
    cur = conn.cursor()
    deactivate_employee_itp_obligations(cur)
    conn.commit()

    cur.execute("""
        SELECT
            -- =====================================
            -- PENDING (TOTAL ACTUAL, NO HISTÓRICO)
            -- =====================================
            COALESCE(
                SUM(
                    CASE
                        WHEN status IN ('PENDING','PARTIAL') THEN
                            CASE
                                WHEN currency = 'CRC' THEN balance / 500.0
                                ELSE balance
                            END
                    END
                ), 0
            ) AS pending_usd,

            -- =====================================
            -- PAID (SOLO PAGOS DEL MES EN CURSO)
            -- =====================================
            COALESCE(
                SUM(
                    CASE
                        WHEN status IN ('PAID','PARTIAL')
                         AND last_payment_date IS NOT NULL
                         AND DATE_TRUNC('month', last_payment_date) = DATE_TRUNC('month', CURRENT_DATE)
                        THEN
                            CASE
                                WHEN currency = 'CRC' THEN (total - balance) / 500.0
                                ELSE (total - balance)
                            END
                    END
                ), 0
            ) AS paid_usd,

            -- =====================================
            -- DPO (SOLO FACTURAS PAGADAS ESTE MES)
            -- =====================================
            ROUND(
                AVG(
                    CASE
                        WHEN status = 'PAID'
                         AND last_payment_date IS NOT NULL
                         AND DATE_TRUNC('month', last_payment_date) = DATE_TRUNC('month', CURRENT_DATE)
                         AND issue_date IS NOT NULL
                        THEN (last_payment_date - issue_date)
                    END
                ), 2
            ) AS dpo,

            -- =====================================
            -- OVERDUE COUNT (ACTUAL)
            -- =====================================
            COUNT(
                CASE
                    WHEN balance > 0
                     AND due_date IS NOT NULL
                     AND due_date < CURRENT_DATE
                    THEN 1
                END
            ) AS overdue,

            -- =====================================
            -- OVERDUE AMOUNT (USD)
            -- =====================================
            COALESCE(
                SUM(
                    CASE
                        WHEN balance > 0
                         AND due_date IS NOT NULL
                         AND due_date < CURRENT_DATE
                        THEN
                            CASE
                                WHEN currency = 'CRC' THEN balance / 500.0
                                ELSE balance
                            END
                    END
                ), 0
            ) AS overdue_amount_usd

        FROM payment_obligations
        WHERE record_type = 'OBLIGATION'
          AND COALESCE(active, TRUE) = TRUE
    """)

    pending, paid, dpo, overdue, overdue_amount = cur.fetchone()

    return {
        "pending": round(pending, 2),
        "paid": round(paid, 2),
        "dpo": dpo,
        "overdue": overdue,
        "overdue_amount": round(overdue_amount, 2),
        "currency": "USD",
        "exchange_rate": 500,
        "scope": "CURRENT_MONTH"
    }


@router.get("/biweekly-obligations/preview")
def biweekly_obligations_preview(
    period: str = Query(...),
    fortnight: int = Query(1),
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    company = normalize_company_code(header_value=x_company_code)
    cur = conn.cursor(cursor_factory=RealDictCursor)
    _ensure_company_column(cur)
    rows = []
    default_crc_bank = "CR87010200009640180220"
    aaron_bank = "CR27010200009688657826"

    def suggested_bank(category, name, currency, current=""):
        if category == "Tarjetas de credito":
            return "BAC"
        if "aaron avila" in str(name or "").lower():
            return aaron_bank
        if str(currency or "CRC").upper() == "CRC":
            return default_crc_bank
        return current or ""

    def row(category, name, amount, currency="CRC", bank_account="", source="MANUAL", notes="", due_date=None, obligation_id=None, reference="", balance=None):
        amount = _money(amount)
        currency = currency or "CRC"
        return {
            "category": category,
            "name": str(name or "").strip(),
            "amount": float(amount),
            "currency": currency,
            "bank_account": suggested_bank(category, name, currency, bank_account),
            "due_date": due_date or _fortnight_due_date(period, fortnight),
            "source": source,
            "notes": notes or "",
            "obligation_id": obligation_id,
            "reference": reference or "",
            "balance": float(_money(balance if balance is not None else amount)),
            "bank_accounting_code": "1.1.02.02.01",
            "bank_accounting_name": "Banco BAC San Jose CRC CR87010200009640180220",
            "bank_voucher": "",
        }

    try:
        cur.execute(
            """
            SELECT nombre, apellidos, salario, pago, banco, cuenta_iban, moneda
            FROM empleados
            WHERE COALESCE(estado, 'Activo') = 'Activo'
              AND COALESCE(activo, TRUE) = TRUE
              AND company_code = %s
              AND COALESCE(salario, 0) > 0
            ORDER BY nombre, apellidos
            """,
            (company,),
        )
        for emp in cur.fetchall() or []:
            pago = str(emp.get("pago") or "").upper()
            amount = _money(emp.get("salario")) / (Decimal("2") if "QUINC" in pago else Decimal("1"))
            rows.append(row(
                "Planilla",
                f"{emp.get('nombre') or ''} {emp.get('apellidos') or ''}".strip(),
                amount,
                emp.get("moneda") or "CRC",
                emp.get("cuenta_iban") or emp.get("banco") or "",
                "EMPLEADOS",
                "Salario sugerido por quincena desde Master Data Empleados.",
            ))

        if int(fortnight or 1) == 1:
            prev_period = _previous_period(period)
            try:
                year, month = [int(part) for part in prev_period.split("-")]
                next_year, next_month = (year + 1, 1) if month == 12 else (year, month + 1)
                cur.execute(
                    """
                    WITH tax_ranked AS (
                        SELECT d.*,
                               ROW_NUMBER() OVER (
                                   PARTITION BY d.direction,
                                                COALESCE(NULLIF(d.document_number, ''), NULLIF(d.electronic_key, ''), d.source_table || ':' || d.source_id, d.id::text)
                                   ORDER BY
                                       CASE
                                           WHEN d.source_table IN ('hacienda_emitted_excel', 'hacienda_acceptance_excel', 'xml_upload') THEN 0
                                           WHEN d.xml_path IS NOT NULL THEN 1
                                           WHEN d.source_table IN ('invoicing', 'collections') THEN 2
                                           WHEN d.source_table = 'payment_obligations' THEN 3
                                           ELSE 4
                                       END,
                                       CASE WHEN COALESCE(d.tax_amount, 0) <> 0 THEN 0 ELSE 1 END,
                                       d.id DESC
                               ) AS tax_rank
                        FROM tax_electronic_documents d
                        WHERE d.issue_datetime >= %s
                          AND d.issue_datetime < %s
                          AND COALESCE(d.issue_datetime::date, CURRENT_DATE) <= CURRENT_DATE
                          AND company_code = %s
                    )
                    SELECT direction,
                           COALESCE(SUM(
                               CASE
                                   WHEN UPPER(COALESCE(currency_code,'CRC')) IN ('CRC','COLON','COLONES')
                                   THEN tax_amount
                                   ELSE tax_amount * COALESCE(NULLIF(exchange_rate,0),1)
                               END
                           ),0) AS tax_crc
                    FROM tax_ranked
                    WHERE tax_rank = 1
                    GROUP BY direction
                    """,
                    (f"{year:04d}-{month:02d}-01", f"{next_year:04d}-{next_month:02d}-01", company),
                )
                taxes = {r["direction"]: _money(r["tax_crc"]) for r in cur.fetchall() or []}
                iva_amount = _money(taxes.get("SALE", Decimal("0")) - taxes.get("PURCHASE", Decimal("0")))
                if iva_amount > 0:
                    rows.append(row("IVA", f"IVA por pagar {prev_period}", iva_amount, "CRC", default_crc_bank, "ACCOUNTING_TAX", "IVA venta menos IVA compra del mes anterior.", f"{period}-15"))
            except Exception as exc:
                conn.rollback()
                rows.append(row("IVA", f"Revisar IVA mes anterior {prev_period}", 0, "CRC", default_crc_bank, "REVISION", str(exc), f"{period}-15"))

            rows.append(row("CCSS", "CCSS por pagar", 0, "CRC", default_crc_bank, "MANUAL", "Completar monto confirmado por CCSS.", f"{period}-15"))
            rows.append(row("Telefonia", "Manfred Bolanos Barrantes", 7000, "CRC", default_crc_bank, "AUTO_FIXED", "Apoyo celular primera quincena.", f"{period}-15"))
            rows.append(row("Telefonia", "Erasmo Gomez Gomez", 7000, "CRC", default_crc_bank, "AUTO_FIXED", "Apoyo celular primera quincena.", f"{period}-15"))

            cur.execute(
                """
                SELECT DISTINCT ON (COALESCE(card_last4,''), COALESCE(NULLIF(TRIM(card_last4),''), source_filename, id::text))
                       card_last4, statement_period, payment_due_date, cash_payment_crc, cash_payment_usd
                FROM corporate_card_statements
                WHERE company_code=%s
                  AND COALESCE(status,'IMPORTED') <> 'VOID'
                ORDER BY COALESCE(card_last4,''), COALESCE(NULLIF(TRIM(card_last4),''), source_filename, id::text),
                         cutoff_date DESC NULLS LAST, id DESC
                """,
                (company,),
            )
            card_labels = {"3155": "Aaron", "1951": "Diana", "1936": "Diana", "1969": "Pabel", "1944": "Pabel", "3148": "ITP"}
            for st in cur.fetchall() or []:
                last4 = str(st.get("card_last4") or "").strip()
                label = card_labels.get(last4, f"Tarjeta {last4 or 'BAC'}")
                crc = _money(st.get("cash_payment_crc"))
                usd = _money(st.get("cash_payment_usd"))
                if crc > 0:
                    rows.append(row("Tarjetas de credito", f"BAC {label} contado CRC {st.get('statement_period') or ''}", crc, "CRC", "BAC", "CORP_CARD", f"Tarjeta {last4}", f"{period}-15"))
                if usd > 0:
                    rows.append(row("Tarjetas de credito", f"BAC {label} contado USD {st.get('statement_period') or ''}", usd, "USD", "BAC", "CORP_CARD", f"Tarjeta {last4}. Convertir/pagar segun banco.", f"{period}-15"))

        cur.execute(
            """
            SELECT id, payee_name, obligation_type, reference, currency, balance, issue_date, due_date,
                   payment_bank, payment_bank_account_code, payment_bank_account_name, notes
            FROM payment_obligations
            WHERE COALESCE(active, TRUE)=TRUE
              AND company_code=%s
              AND status IN ('PENDING','PARTIAL')
              AND COALESCE(balance,0) > 0
            ORDER BY due_date NULLS LAST, payee_name
            """,
            (company,),
        )
        for ob in cur.fetchall() or []:
            if int(fortnight or 1) != 1:
                continue
            issue_date = ob.get("issue_date")
            if issue_date and str(issue_date)[:7] != period:
                continue
            haystack = " ".join(str(ob.get(k) or "") for k in ("payee_name", "obligation_type", "notes", "reference")).lower()
            if "alquiler" in haystack or "rent" in haystack or "prime properties" in haystack:
                category = "Alquiler"
            elif "internet" in haystack or "american data" in haystack:
                category = "Internet"
            elif "surveyor" in haystack or str(ob.get("obligation_type") or "").upper() == "SURVEYOR_FEE":
                category = "Surveyors"
            else:
                continue
            rows.append(row(
                category,
                ob.get("payee_name") or ob.get("reference") or category,
                ob.get("balance"),
                ob.get("currency") or "CRC",
                ob.get("payment_bank_account_name") or ob.get("payment_bank_account_code") or ob.get("payment_bank") or "",
                "ITP",
                f"Aplicar pago a ITP #{ob.get('id')} | Ref: {ob.get('reference') or ''}".strip(),
                str(ob.get("due_date") or _fortnight_due_date(period, fortnight)),
                obligation_id=ob.get("id"),
                reference=ob.get("reference") or "",
                balance=ob.get("balance"),
            ))
        return {"period": period, "fortnight": int(fortnight or 1), "company_code": company, "rows": rows}
    except Exception as exc:
        conn.rollback()
        raise HTTPException(status_code=500, detail=f"No se pudo generar obligaciones quincenales: {exc}")


@router.post("/biweekly-obligations/apply")
def biweekly_obligations_apply(
    payload: dict,
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
    x_user: str | None = Header(None, alias="X-User"),
):
    company = normalize_company_code(header_value=x_company_code)
    user = x_user or "SYSTEM"
    period = str(payload.get("period") or "").strip()
    rows = payload.get("rows") or []
    if not period or not isinstance(rows, list):
        raise HTTPException(status_code=400, detail="Periodo y lineas son obligatorios")
    cur = conn.cursor(cursor_factory=RealDictCursor)
    _ensure_company_column(cur)
    _ensure_biweekly_schema(cur)
    errors = []
    saved = posted = applied = 0
    try:
        cur.execute(
            """
            INSERT INTO itp_biweekly_payment_batches(company_code, period, fortnight, created_by)
            VALUES(%s,%s,%s,%s)
            RETURNING id
            """,
            (company, period, int(payload.get("fortnight") or 1), user),
        )
        batch_id = cur.fetchone()["id"]
        for idx, item in enumerate(rows, start=1):
            try:
                category = str(item.get("category") or "").strip()
                beneficiary = str(item.get("name") or "").strip()
                bank_code = str(item.get("bank_accounting_code") or "").strip()
                voucher = str(item.get("bank_voucher") or "").strip()
                payment_date = str(item.get("due_date") or "").strip()
                currency = str(item.get("currency") or "CRC").upper()
                amount = _money(item.get("amount"))
                if amount <= 0:
                    continue
                missing = []
                if not category:
                    missing.append("rubro")
                if not beneficiary:
                    missing.append("beneficiario")
                if not bank_code:
                    missing.append("cuenta contable banco")
                if not voucher:
                    missing.append("comprobante bancario")
                if not payment_date:
                    missing.append("fecha pago")
                if missing:
                    raise ValueError("Faltan campos obligatorios: " + ", ".join(missing))
                datetime.strptime(payment_date, "%Y-%m-%d")
                cur.execute(
                    """
                    SELECT account_code, account_name
                    FROM accounting_accounts
                    WHERE account_code=%s
                      AND COALESCE(active, TRUE)=TRUE
                      AND COALESCE(accepts_posting, FALSE)=TRUE
                    LIMIT 1
                    """,
                    (bank_code,),
                )
                bank_row = cur.fetchone()
                if not bank_row:
                    raise ValueError(f"Cuenta contable banco invalida o inactiva: {bank_code}")
                rate = _exchange_rate(cur, payment_date) if currency == "USD" else Decimal("1.00")
                amount_crc = (amount * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                bank_name = bank_row["account_name"]
                description = f"Pago quincenal {category} - {beneficiary} - Comp {voucher}"
                obligation_id = item.get("obligation_id")
                if obligation_id:
                    cur.execute("SELECT id, balance FROM payment_obligations WHERE id=%s AND company_code=%s FOR UPDATE", (int(obligation_id), company))
                    ob = cur.fetchone()
                    if not ob:
                        raise ValueError(f"ITP {obligation_id} no existe")
                    balance = _money(ob.get("balance"))
                    if amount > balance:
                        raise ValueError(f"Pago excede saldo ITP {obligation_id}")
                    new_balance = (balance - amount).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                    cur.execute(
                        """
                        UPDATE payment_obligations
                           SET balance=%s,
                               status=%s,
                               last_payment_date=%s,
                               payment_bank_account_code=%s,
                               payment_bank_account_name=%s,
                               updated_at=NOW()
                         WHERE id=%s AND company_code=%s
                        """,
                        (new_balance, "PAID" if new_balance == 0 else "PARTIAL", payment_date, bank_code, bank_name, int(obligation_id), company),
                    )
                    applied += 1
                entry_origin = "ITP_PAYMENT" if obligation_id else "ITP_BIWEEKLY_PAYMENT"
                entry_origin_id = int(obligation_id) if obligation_id else batch_id * 10000 + idx
                cur.execute("SELECT id FROM accounting_entries WHERE origin=%s AND origin_id=%s AND company_code=%s LIMIT 1", (entry_origin, entry_origin_id, company))
                existing = cur.fetchone()
                if existing:
                    entry_id = existing["id"]
                    cur.execute("DELETE FROM accounting_lines WHERE entry_id=%s", (entry_id,))
                else:
                    cur.execute(
                        """
                        INSERT INTO accounting_entries(entry_date, period, description, origin, origin_id, created_by, workflow_status, company_code, currency_code, exchange_rate, posting_rule_code, posting_metadata, posted_by, posted_at)
                        VALUES(%s,%s,%s,%s,%s,%s,'POSTED',%s,'CRC',%s,%s,%s,%s,NOW())
                        RETURNING id
                        """,
                        (payment_date, period, description, entry_origin, entry_origin_id, user, company, rate, entry_origin, Json({"bank_voucher": voucher, "category": category, "source": item.get("source")}), user),
                    )
                    entry_id = cur.fetchone()["id"]
                debit_code, debit_name = ("2.1.01.01", "Cuentas por pagar-comerciales") if obligation_id else _debit_account_for(category)
                cur.execute(
                    """
                    INSERT INTO accounting_lines(entry_id, account_code, account_name, debit, credit, line_description)
                    VALUES(%s,%s,%s,%s,0,%s),(%s,%s,%s,0,%s,%s)
                    """,
                    (entry_id, debit_code, debit_name, amount_crc, description, entry_id, bank_code, bank_name, amount_crc, description),
                )
                posted += 1
                cur.execute(
                    """
                    INSERT INTO itp_biweekly_payment_lines(
                        batch_id, company_code, category, beneficiary, amount, currency, amount_crc,
                        destination_account, bank_accounting_code, bank_accounting_name, bank_voucher,
                        payment_date, obligation_id, reference, source, notes, accounting_entry_id
                    )
                    VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    """,
                    (
                        batch_id, company, category, beneficiary, amount, currency, amount_crc,
                        item.get("bank_account") or "", bank_code, bank_name, voucher, payment_date,
                        obligation_id, item.get("reference") or "", item.get("source") or "", item.get("notes") or "", entry_id,
                    ),
                )
                saved += 1
            except Exception as exc:
                errors.append(f"Linea {idx}: {exc}")
        if errors:
            conn.rollback()
            raise HTTPException(status_code=400, detail="\n".join(errors[:10]))
        conn.commit()
        return {"status": "ok", "batch_id": batch_id, "saved": saved, "posted": posted, "applied": applied}
    except HTTPException:
        raise
    except Exception as exc:
        conn.rollback()
        raise HTTPException(status_code=500, detail=f"No se pudo aplicar obligaciones quincenales: {exc}")


@router.post("/biweekly-obligations/export.xlsx")
def biweekly_obligations_export(payload: dict):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"No se pudo cargar motor Excel: {exc}")

    rows = payload.get("rows") or []
    if not isinstance(rows, list):
        raise HTTPException(status_code=400, detail="Lineas invalidas")

    columns = [
        ("category", "Rubro"),
        ("name", "Nombre / beneficiario"),
        ("amount", "Monto"),
        ("currency", "Moneda"),
        ("bank_account", "Cuenta bancaria destino"),
        ("bank_accounting_code", "Cuenta contable banco"),
        ("bank_voucher", "Comprobante"),
        ("due_date", "Fecha pago"),
        ("obligation_id", "ITP ID"),
        ("reference", "Referencia"),
        ("source", "Fuente"),
        ("notes", "Notas"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Obligaciones"
    period = str(payload.get("period") or "")
    fortnight = str(payload.get("fortnight") or "")
    ws["A1"] = f"Obligaciones quincenales {period} Q{fortnight}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append([label for _, label in columns])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003A75")

    for item in rows:
        if not isinstance(item, dict):
            continue
        ws.append([item.get(key) for key, _ in columns])

    for column_cells in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 46)
        ws.column_dimensions[column_cells[0].column_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Obligaciones_Quincenales_{period or 'periodo'}_Q{fortnight or '1'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

# ============================================================
# 3️⃣ APPLY PAYMENT — BLINDADO FINANCIERO
# ============================================================
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation

@router.post("/apply-payment")
def apply_payment(
    obligation_id: int,
    amount: float,
    payment_date: date,
    bank_account_code: Optional[str] = Query(None),
    bank_account_name: Optional[str] = Query(None),
    bank_name: Optional[str] = Query(None),
    payment_reference: Optional[str] = Query(None),
    conn=Depends(get_db),
    x_user: str | None = Header(None, alias="X-User"),
    x_role: str | None = Header(None, alias="X-Role"),
    x_user_role: str | None = Header(None, alias="X-User-Role"),
):
    cur = None

    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        bank_account_code = str(bank_account_code or "").strip()
        bank_account_name = str(bank_account_name or "").strip()
        bank_name = str(bank_name or bank_account_name or "").strip()
        performed_by, performed_role = actor_from_headers(x_user, x_role, x_user_role)

        cur.execute("""
            ALTER TABLE payment_obligations
            ADD COLUMN IF NOT EXISTS payment_bank TEXT
        """)
        cur.execute("""
            ALTER TABLE payment_obligations
            ADD COLUMN IF NOT EXISTS payment_bank_account_code TEXT
        """)
        cur.execute("""
            ALTER TABLE payment_obligations
            ADD COLUMN IF NOT EXISTS payment_bank_account_name TEXT
        """)
        cur.execute("""
            ALTER TABLE payment_obligations
            ADD COLUMN IF NOT EXISTS payment_reference TEXT
        """)
        payment_reference = str(payment_reference or "").strip()

        # =====================================================
        # 1️⃣ BLOQUEAR FILA (ANTI CONCURRENCIA)
        # =====================================================
        cur.execute("""
            SELECT *
            FROM payment_obligations
            WHERE id = %s
              AND record_type = 'OBLIGATION'
              AND COALESCE(active, TRUE) = TRUE
            FOR UPDATE
        """, (obligation_id,))

        obligation = cur.fetchone()

        if not obligation:
            raise HTTPException(
                status_code=404,
                detail="Obligation not found"
            )

        bank_row = resolve_itp_bank(
            cur,
            bank_account_code,
            bank_account_name,
            payee_name=obligation.get("payee_name"),
            payee_type=obligation.get("payee_type"),
            obligation_type=obligation.get("obligation_type"),
            country=obligation.get("country"),
            reference=obligation.get("reference"),
            notes=obligation.get("notes"),
        )
        if bank_account_code and not bank_row:
            raise HTTPException(
                status_code=400,
                detail="Selected bank account does not exist or is inactive"
            )
        if bank_row:
            bank_account_code = bank_row["account_code"]
            bank_account_name = bank_row["account_name"]
            bank_name = bank_account_name

        before_obligation = row_to_dict(obligation)

        # =====================================================
        # 2️⃣ CONVERTIR A DECIMAL (FINANCIERO CORRECTO)
        # =====================================================
        try:
            balance = Decimal(str(obligation["balance"])).quantize(
                Decimal("0.01"),
                rounding=ROUND_HALF_UP
            )

            amount_decimal = Decimal(str(amount)).quantize(
                Decimal("0.01"),
                rounding=ROUND_HALF_UP
            )

        except InvalidOperation:
            raise HTTPException(
                status_code=400,
                detail="Invalid monetary format"
            )

        # =====================================================
        # 3️⃣ VALIDACIONES DE NEGOCIO
        # =====================================================
        if amount_decimal <= Decimal("0.00"):
            raise HTTPException(
                status_code=400,
                detail="Payment amount must be greater than zero"
            )

        if obligation["status"] == "PAID":
            raise HTTPException(
                status_code=400,
                detail="Obligation is already fully paid"
            )

        if amount_decimal > balance:
            raise HTTPException(
                status_code=400,
                detail="Payment exceeds outstanding balance"
            )

        # =====================================================
        # 4️⃣ CÁLCULO SEGURO DE NUEVO SALDO
        # =====================================================
        settlement = external_surveyor_settlement(
            cur,
            amount_decimal,
            payee_name=obligation.get("payee_name"),
            fallback_country=obligation.get("country"),
            payee_type=obligation.get("payee_type"),
            obligation_type=obligation.get("obligation_type"),
        )

        new_balance = (balance - amount_decimal).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP
        )

        if new_balance < Decimal("0.00"):
            # Protección extrema (nunca debería pasar)
            raise HTTPException(
                status_code=400,
                detail="Resulting balance cannot be negative"
            )

        new_status = (
            "PAID"
            if new_balance == Decimal("0.00")
            else "PARTIAL"
        )

        # =====================================================
        # 5️⃣ UPDATE TRANSACCIONAL
        # =====================================================
        cur.execute("""
            UPDATE payment_obligations
            SET
                balance = %s,
                status = %s,
                last_payment_date = %s,
                payment_bank = %s,
                payment_bank_account_code = %s,
                payment_bank_account_name = %s,
                payment_reference = %s,
                updated_at = NOW()
            WHERE id = %s
        """, (
            new_balance,
            new_status,
            payment_date,
            bank_name or None,
            bank_account_code or None,
            bank_account_name or None,
            payment_reference or None,
            obligation_id
        ))

        cur.execute("SELECT * FROM payment_obligations WHERE id = %s", (obligation_id,))
        after_obligation = row_to_dict(cur.fetchone())

        audit_event(
            cur,
            module="itp",
            action="PAYMENT_APPLIED",
            entity_type="payment_obligation",
            entity_id=obligation_id,
            performed_by=performed_by,
            performed_role=performed_role,
            before=before_obligation,
            after=after_obligation,
            metadata={
                "applied_amount": str(amount_decimal),
                "external_surveyor_rule_applied": bool(settlement["applies"]),
                "deduction_usd": str(settlement["deduction"]),
                "withholding_usd": str(settlement["withholding"]),
                "net_payment_usd": str(settlement["net_payment"]),
                "payment_date": payment_date,
                "payment_reference": payment_reference or None,
                "bank_account_code": bank_account_code or None,
                "bank_account_name": bank_account_name or None,
                "new_balance": str(new_balance),
            },
        )

        conn.commit()

        accounting_warning = None
        try:
            from services.accounting_auto import sync_itp_to_accounting
            sync_itp_to_accounting(conn)
        except Exception as sync_error:
            accounting_warning = str(sync_error)

        return {
            "message": "Payment applied successfully",
            "obligation_id": obligation_id,
            "previous_balance": float(balance),
            "applied_amount": float(amount_decimal),
            "external_surveyor_rule_applied": bool(settlement["applies"]),
            "deduction_usd": float(settlement["deduction"]),
            "withholding_usd": float(settlement["withholding"]),
            "net_payment_usd": float(settlement["net_payment"]),
            "new_balance": float(new_balance),
            "status": new_status,
            "accounting_warning": accounting_warning
        }

    except HTTPException:
        if conn:
            conn.rollback()
        raise

    except Exception as e:
        if conn:
            conn.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Error applying payment: {repr(e)}"
        )

    finally:
        if cur:
            cur.close()

# ============================================================
# 4️⃣ MANUAL OBLIGATION
# ============================================================
@router.post("/manual")
def create_manual_obligation(
    payee_name: str,
    obligation_type: str,
    total: float,
    currency: str,
    reference: Optional[str] = None,
    notes: Optional[str] = None,
    payee_type: str = "OTHER",
    conn=Depends(get_db)
):
    if total <= 0:
        raise HTTPException(
            status_code=400,
            detail="Total must be greater than zero"
        )

    cur = conn.cursor(cursor_factory=RealDictCursor)

    try:
        if is_employee_payee(cur, payee_name):
            raise HTTPException(
                status_code=400,
                detail="El beneficiario existe como empleado en Master Data; no se registra en ITP."
            )

        cur.execute("""
            INSERT INTO payment_obligations (
                record_type,
                payee_type,
                payee_name,
                obligation_type,
                reference,
                currency,
                total,
                balance,
                status,
                origin,
                notes,
                active,
                created_at
            )
            VALUES (
                'OBLIGATION',
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                'PENDING',
                'MANUAL',
                %s,
                TRUE,
                NOW()
            )
            RETURNING id
        """, (
            payee_type,
            payee_name,
            obligation_type,
            reference,
            currency,
            total,
            total,
            notes
        ))

        new_id = cur.fetchone()["id"]
        conn.commit()

    except HTTPException:
        conn.rollback()
        raise

    except Exception as e:
        conn.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Error creating manual obligation: {str(e)}"
        )

    return {
        "message": "Manual obligation created successfully",
        "id": new_id
    }

# ============================================================
# 📥 UPLOAD XML (FACTURA / NC) — 100% BLINDADO
# ============================================================
@router.post("/upload/xml")
def upload_invoice_xml(
    file: UploadFile = File(...),
    conn=Depends(get_db),
    x_company_code: str | None = Header(None, alias="X-Company-Code"),
):
    from xml.etree import ElementTree as ET
    from datetime import datetime, date, timedelta
    import os
    import shutil

    cur = conn.cursor(cursor_factory=RealDictCursor)
    company = normalize_company_code(header_value=x_company_code)
    _ensure_company_column(cur)

    # ============================================================
    # GUARDAR ARCHIVO
    # ============================================================
    os.makedirs("storage/invoice_to_pay/xml", exist_ok=True)

    ts = int(datetime.now().timestamp())
    safe_name = file.filename.replace(" ", "_")
    filepath = f"storage/invoice_to_pay/xml/{ts}_{safe_name}"

    with open(filepath, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    # ============================================================
    # PARSE XML (TOLERANTE)
    # ============================================================
    try:
        tree = ET.parse(filepath)
        root = tree.getroot()

        # Namespace dinámico
        if "}" in root.tag:
            ns_uri = root.tag.split("}")[0].strip("{")
            ns = {"fe": ns_uri}
            p = "fe:"
        else:
            ns = {}
            p = ""

        def find_text(paths, default=None):
            for path in paths:
                node = root.find(path, ns)
                if node is not None and node.text:
                    return node.text.strip()
            return default

        # ------------------------------------------------------------
        # TIPO DOCUMENTO
        # ------------------------------------------------------------
        is_credit_note = root.find(f".//{p}NotaCreditoElectronica", ns) is not None \
                         or "NotaCredito" in root.tag

        obligation_type = (
            "SUPPLIER_CREDIT_NOTE" if is_credit_note else "SUPPLIER_INVOICE"
        )

        # ------------------------------------------------------------
        # CLAVE
        # ------------------------------------------------------------
        clave = find_text([
            f".//{p}Clave",
            ".//Clave"
        ])
        if not clave:
            raise ValueError("XML sin Clave")

        # ------------------------------------------------------------
        # FECHA EMISIÓN
        # ------------------------------------------------------------
        fecha_raw = find_text([
            f".//{p}FechaEmision",
            ".//FechaEmision"
        ])
        if not fecha_raw:
            raise ValueError("XML sin FechaEmision")

        issue_date = date.fromisoformat(fecha_raw.split("T")[0])

        # ------------------------------------------------------------
        # EMISOR
        # ------------------------------------------------------------
        emisor = find_text([
            f".//{p}Emisor/{p}Nombre",
            f".//{p}Nombre"
        ], "PROVEEDOR DESCONOCIDO")

        # ------------------------------------------------------------
        # MONEDA
        # ------------------------------------------------------------
        moneda = find_text([
            f".//{p}CodigoMoneda",
            ".//CodigoMoneda"
        ], "CRC")

        # ------------------------------------------------------------
        # TOTAL (FACTURA / NC)
        # ------------------------------------------------------------
        total_raw = find_text([
            f".//{p}TotalComprobante",
            f".//{p}MontoTotal",
            ".//TotalComprobante"
        ])

        if not total_raw:
            raise ValueError("XML sin TotalComprobante")

        total = float(total_raw)

        if is_credit_note:
            total = total * -1  # NC = negativo

        # ------------------------------------------------------------
        # PLAZO
        # ------------------------------------------------------------
        plazo_raw = find_text([
            f".//{p}PlazoCredito",
            ".//PlazoCredito"
        ])

        term_days = int(plazo_raw) if plazo_raw and plazo_raw.isdigit() else 30
        due_date = issue_date + timedelta(days=term_days)

    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Error parsing XML: {str(e)}"
        )

    # ============================================================
    # INSERTAR payment_obligations
    # ============================================================
    try:
        if is_employee_payee(cur, emisor):
            return {
                "message": "XML omitido para ITP: el emisor existe como empleado en Master Data.",
                "type": obligation_type,
                "reference": clave,
                "supplier": emisor,
                "total": total,
                "currency": moneda,
                "skipped": True,
            }

        cur.execute("""
            INSERT INTO payment_obligations (
                company_code,
                record_type,
                payee_type,
                payee_id,
                payee_name,
                obligation_type,
                reference,
                issue_date,
                due_date,
                country,
                currency,
                total,
                balance,
                status,
                origin,
                file_xml,
                active,
                notes,
                created_at,
                updated_at
            )
            VALUES (
                %s,
                'OBLIGATION',
                'SUPPLIER',
                NULL,
                %s,
                %s,
                %s,
                %s,
                %s,
                'Costa Rica',
                %s,
                %s,
                %s,
                'PENDING',
                'UPLOAD',
                %s,
                TRUE,
                %s,
                NOW(),
                NOW()
            )
        """, (
            company,
            emisor,
            obligation_type,
            clave,
            issue_date,
            due_date,
            moneda,
            total,
            total,
            filepath,
            f"Documento cargado por XML ({clave})"
        ))

        conn.commit()

    except Exception as e:
        conn.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"DB insert error: {str(e)}"
        )

    return {
        "message": "XML procesado correctamente",
        "type": obligation_type,
        "reference": clave,
        "supplier": emisor,
        "total": total,
        "currency": moneda
    }


# ============================================================
# 📥 UPLOAD PDF (ADJUNTO) con issue_date y due_date
# ============================================================
@router.post("/upload/pdf")
def upload_invoice_pdf(
    file: UploadFile = File(...),
    reference: str = Form(...),
    issue_date: Optional[date] = Form(None),
    due_date: Optional[date] = Form(None),
    conn=Depends(get_db)
):
    """
    Carga un PDF y crea una obligación en payment_obligations
    Usará la lógica:
      - issue_date: si provisto por UI, si no, fecha actual
      - due_date: si provisto por UI, si no, same day (contado)
    """
    cur = conn.cursor(cursor_factory=RealDictCursor)

    # Normalizar fechas
    issue_val = issue_date or date.today()
    due_val = due_date or issue_val

    os.makedirs("storage/invoice_to_pay/pdf", exist_ok=True)
    timestamp = int(datetime.now().timestamp())
    filename = f"pdf_{timestamp}_{file.filename}"
    filepath = os.path.join("storage/invoice_to_pay/pdf", filename)

    try:
        with open(filepath, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        cur.execute("""
            INSERT INTO payment_obligations (
                record_type,
                obligation_type,
                reference,
                issue_date,
                due_date,
                currency,
                total,
                balance,
                status,
                origin,
                file_pdf,
                notes,
                active,
                created_at
            )
            VALUES (
                'OBLIGATION',
                'PDF_ONLY',
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                'PENDING',
                'UPLOAD',
                %s,
                %s,
                TRUE,
                NOW()
            )
        """, (
            reference,
            issue_val,
            due_val,
            "USD",         # Por defecto USD (puedes ajustar luego)
            0.0,           # Sin monto explícito para PDF
            0.0,
            filepath,
            f"PDF adjunto para {reference}"
        ))

        conn.commit()

    except Exception as e:
        conn.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"PDF upload error: {str(e)}"
        )

    return {"message": "PDF uploaded and obligation created successfully"}


@router.delete("/{obligation_id}")
def delete_invoice_to_pay(obligation_id: int, conn=Depends(get_db)):
    cur = conn.cursor()

    cur.execute(
        "SELECT id FROM payment_obligations WHERE id = %s",
        (obligation_id,)
    )
    row = cur.fetchone()

    if not row:
        raise HTTPException(status_code=404, detail="Obligation not found")

    cur.execute(
        "DELETE FROM payment_obligations WHERE id = %s",
        (obligation_id,)
    )

    conn.commit()

    return {"status": "ok"}
