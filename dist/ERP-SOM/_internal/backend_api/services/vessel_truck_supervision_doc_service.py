import os
import re
import tempfile
from datetime import date, datetime
from docx import Document
try:
    from services.template_autofit import apply_docx_autofit
    from services.document_branding import apply_mci_docx_branding
except ModuleNotFoundError:
    from backend_api.services.template_autofit import apply_docx_autofit
    from backend_api.services.document_branding import apply_mci_docx_branding


# ============================================================
# GENERATE VESSEL TRUCK SUPERVISION WORD REPORT
# ============================================================

def generate_vessel_truck_supervision_doc(data: dict) -> str:
    data = dict(data or {})

    def _format_report_date(value):
        if value in (None, ""):
            return ""
        if isinstance(value, datetime):
            value = value.date()
        if isinstance(value, date):
            return value.strftime("%b %d %Y")
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%b %d %Y", "%B %d %Y"):
            try:
                return datetime.strptime(text.replace(",", " "), fmt).strftime("%b %d %Y")
            except ValueError:
                continue
        return text

    def _date_from_cert_no(value):
        match = re.search(r"-(\d{2})(\d{2})-(\d{4})\b", str(value or ""))
        if not match:
            return ""
        day, month, year = match.groups()
        try:
            return datetime(int(year), int(month), int(day)).strftime("%b %d %Y")
        except ValueError:
            return ""

    cert_date = _date_from_cert_no(data.get("cert_no"))
    for date_key in (
        "report_date",
        "arrival_date",
        "inspection_date",
        "supervision_completed_date",
    ):
        data[date_key] = _format_report_date(data.get(date_key)) or cert_date

    # ========================================================
    # LOAD TEMPLATE (RELATIVE PATH)
    # ========================================================

    base_dir = os.path.dirname(os.path.abspath(__file__))

    template_path = os.path.abspath(
        os.path.join(
            base_dir,
            "..",
            "templates",
            "vessel_truck_supervision_aligned_dates.docx"
        )
    )

    if not os.path.exists(template_path):
        raise Exception(f"Template not found at: {template_path}")

    doc = Document(template_path)

    # ========================================================
    # SAFE VALUE
    # ========================================================

    def safe(value):
        return "" if value is None else str(value)

    # ========================================================
    # SAFE REPLACEMENT (IGUAL QUE GRAIN SAMPLING)
    # ========================================================

    def replace_in_paragraph(paragraph):

        if not paragraph.runs:
            return

        full_text = "".join(run.text for run in paragraph.runs)

        # Reemplazar TODOS los placeholders del template
        for key, value in data.items():
            placeholder = f"{{{key}}}"
            if placeholder in full_text:
                full_text = full_text.replace(
                    placeholder,
                    safe(value)
                )

        # Reescribir todo el texto preservando formato
        index = 0

        for run in paragraph.runs:
            length = len(run.text)
            if length == 0:
                continue

            run.text = full_text[index:index + length]
            index += length

        # Si sobran caracteres
        if index < len(full_text):
            paragraph.runs[-1].text += full_text[index:]

    # ========================================================
    # BODY
    # ========================================================

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph)

    # ========================================================
    # TABLES
    # ========================================================

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)

    # ========================================================
    # HEADERS & FOOTERS
    # ========================================================

    for section in doc.sections:

        header = section.header
        for paragraph in header.paragraphs:
            replace_in_paragraph(paragraph)

        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_in_paragraph(paragraph)

        footer = section.footer
        for paragraph in footer.paragraphs:
            replace_in_paragraph(paragraph)

        for table in footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_in_paragraph(paragraph)

    # ========================================================
    # SAVE
    # ========================================================

    output_path = os.path.join(
        tempfile.gettempdir(),
        f"{data.get('cert_no', 'truck_supervision')}.docx"
    )

    apply_mci_docx_branding(doc, data)
    apply_docx_autofit(doc)
    doc.save(output_path)

    return output_path
