import os
import re
import tempfile
from datetime import date, datetime

from reportlab.lib.colors import black, blue, red
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas


BASE_DIR = os.path.join(tempfile.gettempdir(), "pdf")
os.makedirs(BASE_DIR, exist_ok=True)

COMPANY_LEGAL_NAME = "MSL SRL Marine Surveyors and Logistics Group"
COMPANY_TAX_ID = "3-102-920372"
COMPANY_ADDRESS = "San Jose, Costa Rica, C.A"
COMPANY_ADDRESS_2 = "Alajuela, Rio Segundo, Plaza Aeropuerto, Local G-14"
COMPANY_PHONE = "506-8814-07-84"

BANK_LINES = [
    ("Beneficiary Bank:", "BCR Banco de Costa Rica"),
    ("Direccion fisica:", "San Jose de Costa Rica"),
    ("SWIFT N°", "BCRICRSJ"),
    ("Account:", "308258-5"),
]
IBAN_CODE = "CR49015201308000025850"
BENEFICIARY = COMPANY_LEGAL_NAME


def _safe(value, default=""):
    return default if value in (None, "") else str(value).strip()


def _date_parts(value):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return f"{value.day:02d}", f"{value.month:02d}", f"{str(value.year)[-2:]}"
    text = _safe(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            parsed = datetime.strptime(text[:10], fmt).date()
            return f"{parsed.day:02d}", f"{parsed.month:02d}", f"{str(parsed.year)[-2:]}"
        except Exception:
            pass
    today = date.today()
    return f"{today.day:02d}", f"{today.month:02d}", f"{str(today.year)[-2:]}"


def _money_text(total, moneda="USD"):
    try:
        value = float(total or 0)
    except Exception:
        value = 0.0
    symbol = "$" if _safe(moneda, "USD").upper() == "USD" else _safe(moneda).upper()
    return f"{symbol} {value:,.2f}"


def _wrap(c, text, max_width, font="Times-Roman", size=11):
    words = _safe(text).split()
    lines, current = [], ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if c.stringWidth(candidate, font, size) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def _invoice_payload(data: dict) -> dict:
    place = _safe(data.get("place") or data.get("lugar"))
    if not place:
        place = ", ".join([p for p in [_safe(data.get("puerto")), _safe(data.get("pais"))] if p])
    survey = _safe(data.get("survey") or data.get("operacion") or "SURVEY")
    vessel = _safe(data.get("buque") or data.get("buque_contenedor") or data.get("container"))
    report = _safe(data.get("num_informe") or data.get("numero_informe"))
    client = _safe(data.get("cliente") or data.get("nombre_cliente"))
    description = _safe(data.get("descripcion") or data.get("descripcion_servicio"))
    if not description:
        description = " / ".join([p for p in [report, vessel, client] if p])
    payment_terms = _safe(data.get("payment_terms"))
    if not payment_terms:
        terms = _safe(data.get("termino_pago"), "0")
        payment_terms = f"CREDIT {terms} DAYS" if str(terms) not in ("", "0") else "DUE UPON RECEIPT"
    return {
        "numero_factura": _safe(data.get("numero_factura") or data.get("numero_documento"), "-"),
        "fecha_factura": data.get("fecha_factura") or data.get("fecha_emision") or date.today(),
        "cliente": client,
        "place": place,
        "description": description,
        "survey": survey,
        "total": data.get("total"),
        "moneda": _safe(data.get("moneda"), "USD"),
        "payment_terms": payment_terms,
    }


def generar_factura_manual_pdf(data: dict) -> str:
    invoice = _invoice_payload(data)
    numero = invoice["numero_factura"]
    filename = f"Factura_{re.sub(r'[^A-Za-z0-9_-]+', '_', str(numero))}.pdf"
    path = os.path.join(BASE_DIR, filename)
    day, month, year = _date_parts(invoice["fecha_factura"])

    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4

    left = 1.15 * cm
    top = height - 1.35 * cm
    main_w = width - 2.3 * cm
    main_h = 19.7 * cm
    bottom = top - main_h

    c.setStrokeColor(black)
    c.setLineWidth(1)
    c.rect(left, bottom, main_w, main_h, fill=0)

    # Header
    c.setFont("Times-Bold", 18)
    c.setFillColor(blue)
    c.drawString(left + 0.6 * cm, top - 1.0 * cm, "M.S.L S.R.L")
    c.setFont("Times-Bold", 15)
    c.setFillColor(red)
    c.drawString(left + 0.6 * cm, top - 1.65 * cm, "Marine Surveyors and Logistics Group")
    c.setFillColor(black)
    c.setFont("Times-Bold", 10)
    c.drawString(left + 0.6 * cm, top - 2.15 * cm, COMPANY_ADDRESS)
    c.drawString(left + 0.6 * cm, top - 2.6 * cm, COMPANY_ADDRESS_2)

    c.setFont("Times-Bold", 9)
    c.drawRightString(left + main_w - 0.55 * cm, top - 0.9 * cm, f"Ced. Jurídica {COMPANY_TAX_ID}")
    c.drawRightString(left + main_w - 0.55 * cm, top - 1.35 * cm, f"Phone {COMPANY_PHONE}")

    c.setFont("Times-Bold", 16)
    c.drawCentredString(left + main_w - 3.0 * cm, top - 3.8 * cm, "INVOICE")
    c.setFont("Times-Bold", 13)
    c.drawCentredString(left + main_w - 3.0 * cm, top - 4.35 * cm, "N°")
    c.setFillColor(red)
    c.drawString(left + main_w - 2.68 * cm, top - 4.35 * cm, str(numero))
    c.setFillColor(black)

    # Client and date boxes
    client_x = left + 0.35 * cm
    client_y = top - 8.55 * cm
    client_w = 13.45 * cm
    client_h = 3.75 * cm
    c.rect(client_x, client_y, client_w, client_h, fill=0)
    c.setFont("Times-Bold", 12)
    c.drawString(client_x + 0.28 * cm, client_y + client_h - 0.7 * cm, f"CLIENT: {invoice['cliente'].upper()}")
    c.drawString(client_x + 0.28 * cm, client_y + client_h - 2.0 * cm, f"PLACE: {invoice['place'].upper()}")

    date_x = client_x + client_w + 0.6 * cm
    date_y = client_y + 1.65 * cm
    date_w = 6.55 * cm
    date_h = 1.75 * cm
    c.rect(date_x, date_y, date_w, date_h, fill=0)
    for i in (1, 2):
        c.line(date_x + date_w * i / 3, date_y, date_x + date_w * i / 3, date_y + date_h)
    c.line(date_x, date_y + date_h / 2, date_x + date_w, date_y + date_h / 2)
    c.setFont("Times-Bold", 11)
    for i, label in enumerate(("DAY", "MONTH", "YEAR")):
        c.drawCentredString(date_x + date_w * (i + 0.5) / 3, date_y + date_h - 0.48 * cm, label)
    for i, value in enumerate((day, month, year)):
        c.drawCentredString(date_x + date_w * (i + 0.5) / 3, date_y + 0.23 * cm, value)

    c.setFont("Times-Bold", 8)
    c.setFillColor(red)
    c.drawString(date_x - 0.1 * cm, client_y + 0.42 * cm, f"TERM OF PAYMENT: {invoice['payment_terms'].upper()}")
    c.setFillColor(black)

    # Description
    desc_label_y = client_y - 1.2 * cm
    c.setFont("Times-Bold", 13)
    c.drawString(client_x + 0.3 * cm, desc_label_y, "DESCRIPTION")
    desc_x = client_x
    desc_y = bottom + 1.65 * cm
    desc_w = main_w - 0.7 * cm
    desc_h = desc_label_y - desc_y - 0.42 * cm
    c.rect(desc_x, desc_y, desc_w, desc_h, fill=0)

    y = desc_y + desc_h - 0.55 * cm
    c.setFont("Times-Roman", 11)
    for line in _wrap(c, invoice["description"].upper(), desc_w - 1.0 * cm, "Times-Roman", 11)[:3]:
        c.drawString(desc_x + 0.45 * cm, y, line)
        y -= 0.55 * cm
    y -= 0.35 * cm
    if invoice["place"]:
        c.drawString(desc_x + 0.45 * cm, y, invoice["place"].upper())
        y -= 1.1 * cm
    c.drawString(desc_x + 0.45 * cm, y, "SURVEY:")
    y -= 0.5 * cm
    c.drawString(desc_x + 0.45 * cm, y, f"-{invoice['survey'].upper()}")
    c.drawRightString(desc_x + desc_w - 0.85 * cm, y + 0.55 * cm, _money_text(invoice["total"], invoice["moneda"]))

    # Total box
    total_w = 5.25 * cm
    total_h = 1.0 * cm
    total_x = left + main_w - total_w - 0.75 * cm
    total_y = bottom + 0.35 * cm
    c.rect(total_x, total_y, total_w, total_h, fill=0)
    c.line(total_x + 2.55 * cm, total_y, total_x + 2.55 * cm, total_y + total_h)
    c.setFont("Times-Bold", 12)
    c.drawCentredString(total_x + 1.28 * cm, total_y + 0.34 * cm, "TOTAL")
    c.drawCentredString(total_x + 3.9 * cm, total_y + 0.34 * cm, _money_text(invoice["total"], invoice["moneda"]))

    # Bank block outside main box
    bank_x = left + 1.2 * cm
    bank_y = bottom - 1.2 * cm
    c.setFont("Times-Roman", 10)
    for label, value in BANK_LINES:
        c.drawString(bank_x, bank_y, f"{label} {value}")
        bank_y -= 0.43 * cm
    bank_y -= 0.45 * cm
    c.setFillColor(red)
    c.drawString(bank_x, bank_y, "IBAN CODE:")
    c.setFillColor(black)
    c.drawString(bank_x + 2.0 * cm, bank_y, IBAN_CODE)
    bank_y -= 0.43 * cm
    c.drawString(bank_x, bank_y, f"Beneficiary: {BENEFICIARY}")
    bank_y -= 0.43 * cm
    c.drawString(bank_x, bank_y, f"Address: {COMPANY_ADDRESS_2}")
    bank_y -= 0.43 * cm
    c.drawString(bank_x, bank_y, "Account: 308258-5 BCRICRSJ")

    bank_y -= 0.85 * cm
    c.setFont("Times-Bold", 8)
    c.drawString(bank_x, bank_y, "NOTE:")
    c.setFillColor(red)
    c.drawString(bank_x + 0.85 * cm, bank_y, "PAYMENTS TO BE DRAWN ON C.R BANK FREE OF")
    bank_y -= 0.35 * cm
    c.drawString(bank_x, bank_y, "ALL CHARGES / IN U.S DOLLARS")
    c.setFillColor(black)

    c.showPage()
    c.save()
    return path


def generar_factura_manual_word(data: dict) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    invoice = _invoice_payload(data)
    numero = invoice["numero_factura"]
    filename = f"Factura_{re.sub(r'[^A-Za-z0-9_-]+', '_', str(numero))}.docx"
    path = os.path.join(BASE_DIR, filename)
    day, month, year = _date_parts(invoice["fecha_factura"])

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    def run(paragraph, text, bold=False, color=None, size=10):
        r = paragraph.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return r

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    left_cell, right_cell = table.rows[0].cells
    p = left_cell.paragraphs[0]
    run(p, "M.S.L S.R.L\n", True, (0, 0, 255), 18)
    run(p, "Marine Surveyors and Logistics Group\n", True, (255, 0, 0), 14)
    run(p, f"{COMPANY_ADDRESS}\n{COMPANY_ADDRESS_2}", True, None, 9)
    p = right_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, f"Ced. Jurídica {COMPANY_TAX_ID}\nPhone {COMPANY_PHONE}\n\nINVOICE\nN°", True, None, 10)
    run(p, str(numero), True, (255, 0, 0), 11)

    doc.add_paragraph("")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    p = table.rows[0].cells[0].paragraphs[0]
    run(p, f"CLIENT: {invoice['cliente'].upper()}\n\nPLACE: {invoice['place'].upper()}", True, None, 11)
    p = table.rows[0].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "DAY     MONTH     YEAR\n", True, None, 10)
    run(p, f"{day}       {month}       {year}\n\n", True, None, 10)
    run(p, f"TERM OF PAYMENT: {invoice['payment_terms'].upper()}", True, (255, 0, 0), 8)

    p = doc.add_paragraph()
    run(p, "DESCRIPTION", True, None, 12)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    p = table.rows[0].cells[0].paragraphs[0]
    run(p, f"{invoice['description'].upper()}\n\n{invoice['place'].upper()}\n\nSURVEY:\n-{invoice['survey'].upper()}\n", False, None, 11)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = table.rows[0].cells[0].add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, _money_text(invoice["total"], invoice["moneda"]), False, None, 11)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    p = table.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "TOTAL", True, None, 12)
    p = table.rows[0].cells[1].paragraphs[0]
    run(p, _money_text(invoice["total"], invoice["moneda"]), True, None, 12)

    doc.add_paragraph("")
    for label, value in BANK_LINES:
        p = doc.add_paragraph()
        run(p, f"{label} {value}", False, None, 10)
    p = doc.add_paragraph()
    run(p, "IBAN CODE: ", False, (255, 0, 0), 10)
    run(p, IBAN_CODE, False, None, 10)
    for line in [
        f"Beneficiary: {BENEFICIARY}",
        f"Address: {COMPANY_ADDRESS_2}",
        "Account: 308258-5 BCRICRSJ",
    ]:
        p = doc.add_paragraph()
        run(p, line, False, None, 10)
    p = doc.add_paragraph()
    run(p, "NOTE: ", True, None, 8)
    run(p, "PAYMENTS TO BE DRAWN ON C.R BANK FREE OF\nALL CHARGES / IN U.S DOLLARS", True, (255, 0, 0), 8)

    doc.save(path)
    return path
